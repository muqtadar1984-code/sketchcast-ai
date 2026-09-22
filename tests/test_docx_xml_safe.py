"""docgen/docx_builder.xml_safe — a control character in the text never takes
a document down (platform issue 2cfc1585, 2026-09-21: "All strings must be
XML compatible" from a lesson plan)."""

from __future__ import annotations

import pytest

from docgen import docx_builder as B


def test_form_feed_becomes_a_newline_and_other_control_characters_go():
    assert B.xml_safe("Page one\x0cPage two") == "Page one\nPage two"
    assert B.xml_safe("a\x00b\x01c\x1fd") == "abcd"
    assert B.xml_safe("tabs\tand\nnewlines\r\nstay") == "tabs\tand\nnewlines\r\nstay"
    assert B.xml_safe("\ud800lone surrogate") == "lone surrogate"
    assert B.xml_safe(None) == ""
    assert B.xml_safe(42) == "42"


@pytest.mark.parametrize("bad", ["Nucleus\x0cMembrane", "Cell\x00 wall", "Photo\x1bsynthesis"])
def test_every_text_entry_point_survives_a_control_character(tmp_path, bad):
    doc = B.new_doc("Lesson plan " + bad, "Subtitle " + bad, kind="lesson_plan")
    B.heading(doc, bad)
    B.para(doc, bad)
    B.instructions(doc, bad)
    B.labelled(doc, bad, bad)
    B.bullets(doc, [bad, "fine"])
    B.numbered(doc, [bad])
    B.question(doc, "1. " + bad, first=True)
    B.table(doc, [bad, "b"], [[bad, bad]])
    B.answer_section(doc, bad, [bad])
    out = B.save(doc, tmp_path / "t.docx")
    assert out.exists() and out.stat().st_size > 0


def test_without_the_guard_python_docx_refuses_the_same_text():
    """Pins WHY the guard exists: the library's own check, unchanged."""
    from docx import Document

    with pytest.raises(ValueError, match="XML compatible"):
        Document().add_paragraph().add_run("a\x0cb")

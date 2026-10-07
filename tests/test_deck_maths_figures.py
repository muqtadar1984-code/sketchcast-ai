"""A maths chapter's deck teaches the lesson's verified worked examples;
a figure example's slide carries the picture the geometry engine drew."""

from __future__ import annotations

from pathlib import Path

from agent5_slides import deck_render
from agent5_slides.deck_generator import apply_maths_lesson
from agent5_slides.deck_storyboard import WORKED, storyboard
from maths import lesson as L
from shared.lesson_model import LessonModel
from tests.test_geometry_lesson import _b1_reply, _item
from tests.test_maths_lesson import GOOD_EX1


def _lesson() -> dict:
    fig = L.figure_example(_item(_b1_reply()))
    fig.label = "Example 1"
    return {"topic": "Angles on a straight line", "examples": [fig.model_dump(), dict(GOOD_EX1)]}


def test_the_lessons_examples_replace_the_articles_and_the_figure_is_drawn(tmp_path: Path):
    model = LessonModel(title="Angles", worked_examples=[("prose problem", "prose solution")])
    n = apply_maths_lesson(model, _lesson(), tmp_path, "en")
    assert n == 2 and [p for p, _ in model.worked_examples] == ["ABC is a straight line. Find x.", "3x + 5 = 20"]
    problem, solution = model.worked_examples[0]
    assert "70 + x = 180" in solution and "angles on a straight line add up to 180°" in solution
    assert solution.strip().endswith("x = 110")
    assert model.worked_figures == {0: "maths_fig_1"}
    fig = model.figures["maths_fig_1"]
    assert fig.png and fig.png.exists() and fig.w > 0 and fig.h > 0
    # the algebra example keeps its working, no picture
    assert "3x = 15" in model.worked_examples[1][1]


def test_the_worked_slide_carries_the_figure_and_the_deck_builds(tmp_path: Path):
    model = LessonModel(title="Angles")
    apply_maths_lesson(model, _lesson(), tmp_path, "en")
    slides = storyboard(model)
    worked = [s for s in slides if s.kind == WORKED]
    assert len(worked) == 2 and worked[0].figure is not None and worked[1].figure is None
    out = tmp_path / "deck.pptx"
    path, faults = deck_render.build(slides, out)
    assert not faults and Path(path).exists()
    from pptx import Presentation

    prs = Presentation(str(path))
    idx = slides.index(worked[0])
    kinds = [sh.shape_type for sh in prs.slides[idx].shapes]
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    assert MSO_SHAPE_TYPE.PICTURE in kinds, "the figure is on the worked slide"


def test_a_lesson_without_examples_leaves_the_article_examples(tmp_path: Path):
    model = LessonModel(title="t", worked_examples=[("p", "s")])
    assert apply_maths_lesson(model, {"examples": []}, tmp_path, "en") == 0
    assert model.worked_examples == [("p", "s")]

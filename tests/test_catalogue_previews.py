"""catalogue/previews.py — the preview cards of a kit (plan §4.4).

Pinned: text comes out of a real .pptx and .docx; a card is a 640×360 PNG
that visibly carries the watermark; the kit renderer reads ONLY the deck's
deck_pptx and the plan's and student worksheet's docx (never an answer key,
never questions_json, never the question bank), files the cards under
<slug>/<kit>/<kind>.png in the public bucket and records them on the kit;
the job finishes its own row and reports per kit.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image

from catalogue import previews as P
from tests.catalogue_fakes import FakeSB


def _pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text_frame.text = "Cells — the unit of life"
    s1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = "Stage 7 Biology"
    s2 = prs.slides.add_slide(blank)
    tf = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text_frame
    tf.text = "What every cell has"
    for line in ("A membrane that controls what enters", "Cytoplasm where reactions happen", "A nucleus with the genes"):
        tf.add_paragraph().text = line
    prs.slides.add_slide(blank)  # an empty third slide is never a card
    prs.save(str(path))
    return path


def _docx(path: Path, lines) -> Path:
    from docx import Document

    doc = Document()
    for ln in lines:
        doc.add_paragraph(ln)
    doc.save(str(path))
    return path


def test_text_comes_out_of_the_real_documents(tmp_path):
    slides = P.deck_slides(_pptx(tmp_path / "deck.pptx"))
    assert [s["title"] for s in slides] == ["Cells — the unit of life", "What every cell has"]
    assert slides[0]["lines"] == ["Stage 7 Biology"]
    assert slides[1]["lines"][:2] == ["A membrane that controls what enters", "Cytoplasm where reactions happen"]
    page = P.docx_page(_docx(tmp_path / "plan.docx", ["Cells · Lesson plan", "", "Cambridge 0893 · 7Bs.01", "Objectives", "Name the parts of a cell"]))
    assert page == {"title": "Cells · Lesson plan", "lines": ["Cambridge 0893 · 7Bs.01", "Objectives", "Name the parts of a cell"]}


def test_a_card_is_a_small_png_with_the_watermark_across_it(tmp_path):
    out = P.render_card("Cells — the unit of life", ["A membrane", "Cytoplasm", "A nucleus"], "Slide deck · title slide", tmp_path / "c.png")
    img = Image.open(out)
    assert img.size == (640, 360) and img.format == "PNG"
    # The watermark tints the centre band: a paper-white pixel there is gone.
    centre = img.convert("RGB").getpixel((320, 190))
    assert centre != (252, 252, 250)
    # Arabic renders too (shaped, right-aligned) without raising.
    P.render_card("الخلايا", ["الغشاء الخلوي", "السيتوبلازم"], "Slide deck", tmp_path / "ar.png")
    assert (tmp_path / "ar.png").exists()


def _kit_sb(tmp_path):
    sb = FakeSB()
    sb.tables["topic_kits"] = [{"id": "kit-1", "topic_id": "t-cell", "status": "approved", "previews": [],
                                "presentation_generation_id": "g-p",
                                "doc_generation_ids": {"deck": "g-d", "lesson_plan": "g-l", "worksheet": "g-w", "question_bank": "g-qb", "quiz": "g-q"}}]
    sb.tables["topics"] = [{"id": "t-cell", "title": "Cells", "slug": "cells"}]
    sb.tables["artifacts"] = [
        {"generation_id": "g-d", "kind": "deck_pptx", "storage_path": "sys/g-d/deck.pptx"},
        {"generation_id": "g-l", "kind": "docx", "storage_path": "sys/g-l/lesson_plan.docx"},
        {"generation_id": "g-w", "kind": "docx", "storage_path": "sys/g-w/worksheet.docx"},
        {"generation_id": "g-w", "kind": "answer_key_docx", "storage_path": "sys/g-w/answer_key.docx"},
        {"generation_id": "g-w", "kind": "questions_json", "storage_path": "sys/g-w/questions.json"},
        {"generation_id": "g-qb", "kind": "docx", "storage_path": "sys/g-qb/question_bank.docx"},
    ]
    sb.files[("artifacts", "sys/g-d/deck.pptx")] = _pptx(tmp_path / "deck.pptx").read_bytes()
    sb.files[("artifacts", "sys/g-l/lesson_plan.docx")] = _docx(tmp_path / "plan.docx", ["Cells · Lesson plan", "Objectives"]).read_bytes()
    sb.files[("artifacts", "sys/g-w/worksheet.docx")] = _docx(tmp_path / "ws.docx", ["Cells · Worksheet", "1. Name the organelle that…"]).read_bytes()
    sb.files[("artifacts", "sys/g-w/answer_key.docx")] = b"SECRET"
    sb.files[("artifacts", "sys/g-qb/question_bank.docx")] = b"BANK"
    return sb


def test_the_kit_renderer_reads_only_the_three_documents_and_files_the_cards_publicly(tmp_path):
    sb = _kit_sb(tmp_path)
    cards = P.render_kit_previews(sb, sb.tables["topic_kits"][0], "cells", tmp_path / "out")
    assert [c["kind"] for c in cards] == ["deck-1", "deck-2", "plan-1", "worksheet-1"]
    assert all(c["path"] == f"cells/kit-1/{c['kind']}.png" for c in cards)
    # what was read from storage — and what was NOT
    read = [p for b, p in sb.downloads if b == "artifacts"]
    assert read == ["sys/g-d/deck.pptx", "sys/g-l/lesson_plan.docx", "sys/g-w/worksheet.docx"]
    assert not any("answer_key" in p or "questions" in p or "question_bank" in p for p in read)
    # four PNGs in the PUBLIC bucket, recorded on the kit
    uploaded = sorted(p for b, p in sb.files if b == P.BUCKET)
    assert uploaded == sorted(c["path"] for c in cards)
    assert Image.open(__import__("io").BytesIO(sb.files[(P.BUCKET, "cells/kit-1/deck-1.png")])).size == (640, 360)
    assert sb.tables["topic_kits"][0]["previews"] == cards


def test_a_kit_without_a_deck_still_gets_its_document_cards(tmp_path):
    sb = _kit_sb(tmp_path)
    sb.tables["topic_kits"][0]["doc_generation_ids"].pop("deck")
    cards = P.render_kit_previews(sb, sb.tables["topic_kits"][0], "cells", tmp_path / "out")
    assert [c["kind"] for c in cards] == ["plan-1", "worksheet-1"]


def test_the_job_renders_one_kit_or_every_current_kit_without_previews_and_finishes_its_row(tmp_path, monkeypatch):
    sb = _kit_sb(tmp_path)
    sb.tables["jobs"] = [{"id": "job-1", "type": P.JOB_TYPE, "status": "processing", "params": {"kit_id": "kit-1"}}]
    summary = P.run_previews_job(sb, sb.tables["jobs"][0])
    assert summary["counts"] == {"rendered": 1, "skipped": 0, "errors": 0}
    assert sb.tables["jobs"][0]["status"] == "done"

    # The backfill shape: no kit_id → the current kits lacking previews.
    sb.tables["topic_kits"][0]["previews"] = []
    monkeypatch.setattr(P, "_kits_to_render", lambda sb_, kit_id: [] if kit_id else [sb_.tables["topic_kits"][0]])
    sb.tables["jobs"].append({"id": "job-2", "type": P.JOB_TYPE, "status": "processing", "params": {}})
    summary = P.run_previews_job(sb, sb.tables["jobs"][1])
    assert summary["counts"]["rendered"] == 1 and sb.tables["jobs"][1]["status"] == "done"


def test_a_topic_without_a_slug_is_skipped_not_failed(tmp_path):
    sb = _kit_sb(tmp_path)
    sb.tables["topics"][0]["slug"] = None
    sb.tables["jobs"] = [{"id": "job-1", "type": P.JOB_TYPE, "status": "processing", "params": {"kit_id": "kit-1"}}]
    summary = P.run_previews_job(sb, sb.tables["jobs"][0])
    assert summary["counts"] == {"rendered": 0, "skipped": 1, "errors": 0} and sb.tables["jobs"][0]["status"] == "done"


def test_enqueue_for_kit_writes_one_queued_job():
    sb = FakeSB()
    assert P.enqueue_for_kit(sb, "kit-9") is True
    assert sb.tables["jobs"] and sb.tables["jobs"][-1]["type"] == P.JOB_TYPE and sb.tables["jobs"][-1]["params"] == {"kit_id": "kit-9"}

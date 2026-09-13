"""The deck job on its new path: which source, which renderer, and when it refuses.

`worker.process._generate_deck` now has two renderers behind one variable.
These tests pin the decisions that choose between them and the one property
the founder set as a rule: a deck with a geometry fault is not uploaded.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation

from agent5_slides import deck_generator as dg
from tests.test_deck_generation import (ANALYSIS, BOOK, CHAPTER, _reply,
                                        _StubClient)
from tests.test_observer_job_guard import FakeSB

ARTICLE = {
    "id": "art-1",
    "title": "Cells",
    "objectives": [{"id": "o1", "text": "Describe the cell as the unit of life."}],
    "sections": [
        {"id": "s1", "heading": "What is a cell?",
         "body_md": "All living things are made of cells. They vary in size.",
         "figure_keys": []},
        {"id": "s2", "heading": "Inside the cell",
         "body_md": "The nucleus holds the DNA.\n\n- Mitochondria release energy.\n- Ribosomes make protein.",
         "figure_keys": []},
    ],
    "glossary": [{"term": "cell", "definition": "The basic unit of life."}],
    "misconceptions": [{"id": "m1", "misconception": "All cells are the same size.",
                        "correction": "They vary enormously."}],
    "worked_examples": [],
    "claims": [{"id": "c1", "text": "All living things are made of cells.", "section_id": "s1"},
               {"id": "c2", "text": "Cells vary in size.", "section_id": "s1"},
               {"id": "c3", "text": "The nucleus holds the DNA.", "section_id": "s2"}],
}


# What the BOOK path's authoring call returns now (2026-09-12): the chapter
# part as an article, the catalogue deck's own shape.
BOOK_ARTICLE_REPLY = {
    "title": "Materials and their structure",
    "objectives": [{"id": "o1", "text": "Sort materials into solids, liquids and gases."},
                   {"id": "o2", "text": "Use particle theory to explain their properties."}],
    "sections": [
        {"id": "s1", "heading": "Three states of matter",
         "body_md": "Matter exists as solids, liquids and gases. A solid keeps its shape and volume. "
                    "A liquid keeps its volume but takes the shape of its container. A gas fills any "
                    "container it is put in, spreading out to take both its shape and its volume.",
         "figure_keys": ["particle_states"], "covers": []},
        {"id": "s2", "heading": "Particle arrangement",
         "body_md": "In a solid the particles sit in a fixed pattern and vibrate in place. In a liquid "
                    "they touch but can move past one another. In a gas they are far apart and move "
                    "freely in every direction.",
         "figure_keys": ["particle_states"], "covers": []},
        {"id": "s3", "heading": "Compressing matter",
         "body_md": "A gas can be squeezed into a smaller space because its particles are far apart. "
                    "Liquids and solids cannot be compressed: their particles are already touching. "
                    "A syringe with its nozzle blocked shows this well. With air inside, the plunger "
                    "moves a long way before it stops, because the particles of the gas have empty "
                    "space between them and can be pushed closer together. With water inside, the "
                    "plunger will not move at all, however hard it is pushed, because the particles "
                    "of the liquid are already in contact and there is no empty space left to remove. "
                    "A block of wood behaves like the water. The same idea explains a vacuum, which is "
                    "not a very thin gas but a space with no particles in it whatsoever. A gas can "
                    "always be made thinner by removing particles, and a vacuum is what is left when "
                    "every particle has been removed. Diving cylinders and bicycle pumps rely on the "
                    "fact that a gas can be compressed, while hydraulic brakes rely on the fact that a "
                    "liquid cannot: pressing the pedal moves the fluid, and the fluid moves the brake.",
         "figure_keys": [], "covers": []},
    ],
    "glossary": [{"term": "particle", "definition": "A tiny portion of matter far too small to see."},
                 {"term": "vacuum", "definition": "A space that contains no particles at all."}],
    "misconceptions": [{"id": "m1", "misconception": "A vacuum is a gas with very few particles.",
                        "correction": "A vacuum has no particles at all; a gas, however thin, has some."}],
    "worked_examples": [{"id": "w1", "problem": "Classify tap water.",
                         "solution_md": "Tap water is a liquid: it pours and takes its container's shape."}],
    "claims": [
        {"id": "c1", "text": "Matter exists as solids, liquids and gases.", "section_id": "s1"},
        {"id": "c2", "text": "A solid keeps its shape and its volume.", "section_id": "s1"},
        {"id": "c3", "text": "Solid particles vibrate in a fixed pattern.", "section_id": "s2"},
        {"id": "c4", "text": "Gas particles are far apart and move freely.", "section_id": "s2"},
        {"id": "c5", "text": "Only gases can be compressed.", "section_id": "s3"},
    ],
    "figures": [{"figure_key": "particle_states", "caption": "Particles in the three states",
                 "spec": {"subject": "three containers of particles", "parts": ["solid", "liquid", "gas"],
                          "style": "whiteboard diagram", "notes": "side by side"}}],
    "depth_rationale": "Lower-secondary particle theory, as the chapter teaches it.",
}


def _env(monkeypatch):
    from worker import client as db
    from worker import process

    # The picture ladder searches the REAL library and may generate: a test of
    # the wiring must never reach a network or an image model.
    monkeypatch.setenv("DECK_IMAGES", "0")
    sb = FakeSB()
    sb.tables["generations"] = [{"id": "gen-1", "status": "processing", "kind": "deck",
                                 "owner_id": "u1", "book_id": "book-1", "params": {}}]
    sb.tables["jobs"] = [{"id": "job-1", "type": "deck", "status": "processing",
                          "generation_id": "gen-1", "progress": 45}]
    uploads: list[tuple[str, str]] = []

    def _upload(_sb, local_path, dest):
        assert Path(local_path).exists() and Path(local_path).stat().st_size > 0
        uploads.append((str(local_path), dest))
        return dest

    monkeypatch.setattr(db, "upload_artifact", _upload)
    return sb, uploads, process


class TestWhichRenderer:
    def test_it_ships_on(self, monkeypatch):
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        assert dg.storyboard_enabled() and dg.use_storyboard({})

    def test_one_variable_rolls_it_back(self, monkeypatch):
        monkeypatch.setenv("DECK_STORYBOARD", "0")
        assert not dg.use_storyboard({})

    def test_a_school_template_does_not_send_the_job_back_to_legacy(self, monkeypatch):
        """Four schools have uploaded a .pptx. They are users too: the
        storyboard builds ON their template rather than handing them the old
        PNG deck."""
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        assert dg.use_storyboard({"pptx_template": "/tmp/school.pptx"})
        assert dg.use_storyboard({"pptx_template": None, "accent_rgb": (1, 2, 3)})


def _template(path: Path, width_in: float, height_in: float) -> Path:
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width_in), Inches(height_in)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "School sample slide"
    prs.save(str(path))
    return path


class TestSchoolTemplates:
    def test_a_16_9_template_is_the_base_and_its_sample_slides_are_dropped(self, tmp_path):
        from agent5_slides import deck_render
        tpl = _template(tmp_path / "school.pptx", 13.333, 7.5)
        deck_render.set_branding({"pptx_template": str(tpl), "accent_rgb": (200, 30, 30)})
        prs = deck_render._base({"pptx_template": str(tpl)})
        assert len(prs.slides) == 0, "the school's sample slide is not this lesson"
        assert deck_render._ACCENT["templated"] is True
        assert deck_render._ACCENT["rgb"] == (200, 30, 30)

    def test_a_4_3_template_falls_back_to_the_blank_base_with_the_colours(self, tmp_path):
        from agent5_slides import deck_render
        tpl = _template(tmp_path / "old.pptx", 10, 7.5)
        deck_render.set_branding({"pptx_template": str(tpl), "accent_rgb": (10, 20, 30)})
        prs = deck_render._base({"pptx_template": str(tpl)})
        assert deck_render._ACCENT["templated"] is False
        assert prs.slide_width == deck_render.af.SLIDE_W
        assert deck_render._ACCENT["rgb"] == (10, 20, 30)

    def test_a_whole_deck_builds_on_a_template(self, tmp_path):
        from agent5_slides import deck_render
        from agent5_slides.deck_storyboard import storyboard
        tpl = _template(tmp_path / "school.pptx", 13.333, 7.5)
        model = dg.from_article(ARTICLE, [])
        path, faults = deck_render.build(storyboard(model), tmp_path / "deck.pptx",
                                         branding={"pptx_template": str(tpl), "accent_rgb": (200, 30, 30)})
        assert faults == []
        prs = Presentation(str(path))
        assert len(prs.slides) > 3
        assert not any(sh.has_text_frame and "School sample slide" in sh.text_frame.text
                       for s in prs.slides for sh in s.shapes)
        deck_render.set_branding(None)      # never leak a school's colour into the next test


class TestTheCataloguePath:
    def test_no_authoring_call_and_one_deck_uploaded(self, monkeypatch, tmp_path):
        """The article IS the lesson. A client that would fail if asked to
        write slides proves nobody asked it."""
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        sb, uploads, process = _env(monkeypatch)
        monkeypatch.setattr(dg, "model_from_article",
                            lambda _sb, article, _tmp: dg.from_article(article, []))
        client = _StubClient({})                      # would raise "0 slides" if used
        title = process._generate_deck(
            sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, client, {},
            {}, "en", "ltr", tmp_path, "u1/gen-1", "Cells",
            catalogue=SimpleNamespace(article=ARTICLE),
        )
        assert client.calls == []
        assert [d for _, d in uploads] == ["u1/gen-1/deck.pptx"]
        rows = sb.tables["artifacts"]
        assert [(r["kind"], r["storage_path"]) for r in rows] == [("deck_pptx", "u1/gen-1/deck.pptx")]
        assert title == "Living Things · Cells · Slide deck"
        assert sb.tables["jobs"][0]["progress"] == 96
        prs = Presentation(uploads[0][0])
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes
                 if getattr(sh, "has_text_frame", False) and sh.text_frame.text]
        assert any("Describe the cell" in t for t in texts), "objectives are on a slide"
        assert any("Cells vary in size" in t for t in texts), "claims are the points"
        assert not any(sh.shape_type == 13 for s in prs.slides for sh in s.shapes), \
            "no picture on a deck with no figures — every word is a text object"

    def test_the_prose_lands_in_the_notes_not_on_the_slide(self, monkeypatch, tmp_path):
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        sb, uploads, process = _env(monkeypatch)
        monkeypatch.setattr(dg, "model_from_article",
                            lambda _sb, article, _tmp: dg.from_article(article, []))
        process._generate_deck(
            sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, _StubClient({}), {},
            {}, "en", "ltr", tmp_path, "u1/gen-1", "Cells",
            catalogue=SimpleNamespace(article=ARTICLE),
        )
        prs = Presentation(uploads[0][0])
        notes = " ".join(s.notes_slide.notes_text_frame.text for s in prs.slides
                         if s.has_notes_slide)
        assert "They vary in size." in notes


class TestTheBookPath:
    def test_the_chapter_is_authored_as_an_article_and_becomes_native_slides(self, monkeypatch, tmp_path):
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        sb, uploads, process = _env(monkeypatch)
        client = _StubClient(BOOK_ARTICLE_REPLY)
        process._generate_deck(
            sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, client, {"num_slides": 8},
            {}, "en", "ltr", tmp_path, "u1/gen-1", "Cells",
        )
        assert len(client.calls) == 1, "the one authoring call the book path always made"
        assert "knowledge article" in client.calls[0]["prompt"], "authored as an ARTICLE, not as slides"
        assert [d for _, d in uploads] == ["u1/gen-1/deck.pptx"]
        prs = Presentation(uploads[0][0])
        assert len(list(prs.slides)) > 2
        assert not any(sh.shape_type == 13 for s in prs.slides for sh in s.shapes)
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes
                 if getattr(sh, "has_text_frame", False)]
        assert any("Words to know" in t for t in texts), "the article's glossary"
        assert any("By the end of this lesson" in t for t in texts), "the article's objectives"
        # Coverage is still measured and recorded on this path.
        cov = (sb.tables["generations"][0].get("params") or {}).get("coverage")
        assert cov and cov[0]["kind"] == "deck"

    def test_the_legacy_renderer_still_authors_slides(self, monkeypatch, tmp_path):
        monkeypatch.setenv("DECK_STORYBOARD", "0")
        sb, uploads, process = _env(monkeypatch)
        client = _StubClient(_reply())
        process._generate_deck(
            sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, client, {"num_slides": 8},
            {}, "en", "ltr", tmp_path, "u1/gen-1", "Cells",
        )
        assert "SLIDE DECK" in client.calls[0]["prompt"] and [d for _, d in uploads] == ["u1/gen-1/deck.pptx"]


class TestAFaultyDeckIsNotUploaded:
    def test_geometry_faults_fail_the_job_with_the_slide_named(self, monkeypatch, tmp_path):
        """A deck is the one artifact nobody watches fail: it opens, the
        shapes are valid, and the fault is on the projector. So a fault
        reported by the renderer is a failed job, with the slide in the
        error text for support."""
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        sb, uploads, process = _env(monkeypatch)
        from agent5_slides import deck_render

        real = deck_render.build

        def _faulty(slides, out_path, **kw):
            path, _ = real(slides, out_path, **kw)
            return path, ["slide 4 (diagram): labels overlap: nucleus / nucleolus"]

        monkeypatch.setattr(deck_render, "build", _faulty)
        monkeypatch.setattr(dg, "model_from_article",
                            lambda _sb, article, _tmp: dg.from_article(article, []))
        with pytest.raises(RuntimeError, match="slide 4 .diagram.: labels overlap"):
            process._generate_deck(
                sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, _StubClient({}), {},
                {}, "en", "ltr", tmp_path, "u1/gen-1", "Cells",
                catalogue=SimpleNamespace(article=ARTICLE),
            )
        assert uploads == [] and sb.tables.get("artifacts", []) == []


class _Q:
    """Just enough of the supabase query chain for `figure_art`."""

    def __init__(self, rows):
        self._rows = rows

    def select(self, *a, **k):
        return self

    def eq(self, *a, **k):
        return self

    def order(self, *a, **k):
        return self

    def limit(self, *a, **k):
        return self

    def execute(self):
        return SimpleNamespace(data=self._rows)


class _Storage:
    def __init__(self, blobs):
        self._blobs = blobs
        self.downloads: list[str] = []

    def from_(self, _bucket):
        return self

    def download(self, path):
        self.downloads.append(path)
        return self._blobs[path]


class _SB:
    def __init__(self, tables, blobs=None):
        self._tables = tables
        self.storage = _Storage(blobs or {})

    def table(self, name):
        return _Q(self._tables.get(name, []))


class TestFigureArtwork:
    PNG = (b"\x89PNG\r\n\x1a\n" + b"\x00" * 32)
    ROW = {"figure_key": "animal_cell", "visual_asset_id": "va-1"}

    def test_a_figure_with_no_asset_has_no_art(self, tmp_path):
        art = dg.figure_art(_SB({}), tmp_path)
        assert art({"figure_key": "x", "visual_asset_id": None}) is None

    def test_an_asset_without_a_measured_frame_is_not_used(self, tmp_path):
        """Regions without `w`/`h` are numbers with no scale; a label placed
        from them lands somewhere arbitrary. No picture beats a wrong one."""
        sb = _SB({"visual_assets": [{"asset_key": "animal_cell", "storage_path": "p.png",
                                     "vision": {"regions": {"nucleus": [[1, 1, 2, 2]]}}}]})
        assert dg.figure_art(sb, tmp_path)(self.ROW) is None
        assert sb.storage.downloads == []

    def test_a_measured_asset_is_downloaded_once(self, tmp_path):
        sb = _SB({"visual_assets": [{"asset_key": "animal_cell", "storage_path": "p.png",
                                     "vision": {"regions": {"nucleus": [[1, 1, 2, 2]]},
                                                "w": 704, "h": 613}}]},
                 blobs={"p.png": self.PNG})
        art = dg.figure_art(sb, tmp_path)
        got = art(self.ROW)
        assert got["w"] == 704 and got["png"].read_bytes() == self.PNG
        art(self.ROW)
        assert sb.storage.downloads == ["p.png"], "cached on disk after the first fetch"

    def test_a_storage_error_degrades_to_no_picture_not_a_failed_deck(self, tmp_path):
        sb = _SB({"visual_assets": [{"asset_key": "animal_cell", "storage_path": "missing.png",
                                     "vision": {"regions": {"n": [[1, 1, 2, 2]]}, "w": 10, "h": 10}}]})
        assert dg.figure_art(sb, tmp_path)(self.ROW) is None

    def test_model_from_article_wires_rows_to_art(self, tmp_path):
        sb = _SB({"article_figures": [dict(self.ROW, spec={"parts": ["nucleus"]}, caption="c")],
                  "visual_assets": [{"asset_key": "animal_cell", "storage_path": "p.png",
                                     "vision": {"regions": {"nucleus": [[1, 1, 2, 2]]},
                                                "w": 704, "h": 613}}]},
                 blobs={"p.png": self.PNG})
        m = dg.model_from_article(sb, ARTICLE, tmp_path)
        assert m.title == "Cells" and "animal_cell" in m.figures
        assert m.figures["animal_cell"].annotatable

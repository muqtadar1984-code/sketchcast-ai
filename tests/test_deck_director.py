"""The AI Presentation Director (agent5_slides/director): spec validation,
grounding, worked-example completeness, asset ladder, render editability and
QA — all offline, with a stub model and no network."""

from __future__ import annotations

import copy
import io
from pathlib import Path

import pytest
from PIL import Image

from agent5_slides.director import assets as A
from agent5_slides.director import context as C
from agent5_slides.director import director as D
from agent5_slides.director import maths_check as mc
from agent5_slides.director import pipeline as P
from agent5_slides.director import qa as Q
from agent5_slides.director import spec as S
from shared.lesson_model import from_article

ARTICLE = {
    "title": "Circumference and Area of a Circle", "language": "en",
    "objectives": [{"text": "Calculate the circumference of a circle."},
                   {"text": "Calculate the area of a circle."}],
    "sections": [
        {"id": "s1", "heading": "Anatomy of a Circle", "covers": ["9Gg.01"],
         "body_md": "The radius is the distance from the centre to the boundary. The diameter is twice the radius."},
        {"id": "s2", "heading": "Area", "covers": ["9Gg.01"],
         "body_md": "The area of a circle is A = pi r squared, measured in square units such as cm squared."},
    ],
    "claims": [{"id": "c1", "text": "The diameter is twice the radius.", "section_id": "s1"},
               {"id": "c2", "text": "Area is measured in square units.", "section_id": "s2"}],
    "glossary": [{"term": "Radius", "definition": "Centre to boundary."}],
    "misconceptions": [{"misconception": "Using the diameter as r.", "correction": "Halve d first."}],
    "worked_examples": [{
        "problem": "Calculate the area of a circle with a radius of 6 cm. Use 3.14 for pi.",
        "solution_md": ("Step 1: State the formula. A = pi times r squared.\n\n"
                        "Step 2: Square the radius. 6 squared is 36, so A = 36 pi.\n\n"
                        "Step 3: Multiply. 3.14 times 36 equals 113.04.\n\n"
                        "Step 4: State the final answer. The area is 113.04 square centimetres.")}],
}


def _png(path: Path, w=800, h=600, colour=(30, 120, 200)) -> Path:
    im = Image.new("RGB", (w, h), (250, 250, 250))
    for x in range(0, w, 40):
        for y in range(0, h, 40):
            if (x + y) // 40 % 2:
                im.paste(colour, (x, y, x + 40, y + 40))
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def spec_raw() -> dict:
    """A deck the way the Director writes one — every element type once."""
    return {
        "title": "Circumference & Area of a Circle", "subtitle": "Measure around. Measure inside.",
        "rationale": "Boundary vs surface, built up from anatomy to worked examples.",
        "theme": {"palette": "indigo", "cover_motif": "circle"},
        "assets": [{"id": "A1", "need": "photo", "subject": "bicycle wheel", "search_terms": ["bicycle wheel"],
                    "purpose": "circumference in real life"},
                   {"id": "A2", "need": "illustration", "subject": "circular garden from above"}],
        "slides": [
            {"id": "cover", "background": "dark", "kicker": "MATHEMATICS", "title": "Circumference & Area",
             "subtitle": "Measure the boundary. Measure the space inside.",
             "rows": [{"cols": [{"elements": [
                 {"type": "maths_diagram", "shape": "circle", "show": ["centre", "radius"]},
                 {"type": "chips", "items": [{"label": "Area", "value": "A = πr²"},
                                             {"label": "Diameter", "value": "d = 2r"}]}]}]}],
             "notes": "Welcome.", "evidence": ["o1"]},
            {"id": "roadmap", "kicker": "ROADMAP", "title": "Two ideas, one skill",
             "rows": [{"cols": [{"elements": [{"type": "cards", "numbered": True, "items": [
                 {"title": "Measure", "body": "Distance around"}, {"title": "Calculate", "body": "Space inside"}]}]}]}],
             "notes": "Here is the plan.", "evidence": ["o1", "o2"]},
            {"id": "life", "kicker": "CONNECT TO LIFE", "title": "Around or inside?",
             "rows": [{"cols": [{"elements": [{"type": "cards", "items": [
                 {"title": "Circumference", "body": "Around the wheel", "asset": "A1", "tone": "warm"},
                 {"title": "Area", "body": "Covering the garden", "asset": "A2", "tone": "good"}]}]}]}],
             "footer": {"type": "callout", "kind": "think", "text": "Which one uses square units?"},
             "notes": "Point at each picture.", "evidence": ["c2"]},
            {"id": "anatomy", "kicker": "CORE CONCEPT", "title": "Anatomy of a circle",
             "rows": [{"cols": [
                 {"weight": 1.1, "elements": [{"type": "maths_diagram", "shape": "circle",
                                               "show": ["centre", "radius", "diameter", "chord"],
                                               "labels": [{"part": "radius", "text": "r"}]}]},
                 {"weight": 1, "elements": [{"type": "formula", "expr": "d = 2r", "caption": "The diameter is twice the radius."},
                                            {"type": "bullets", "items": [{"body": "Radius: centre to boundary"}]}]}]}],
             "notes": "Show the radius.", "evidence": ["c1"]},
            {"id": "area", "kicker": "FORMULA", "title": "Area of a circle",
             "rows": [{"cols": [{"elements": [{"type": "formula", "expr": "A = πr²", "caption": "Square units"}]},
                                {"elements": [{"type": "maths_diagram", "shape": "semicircle",
                                               "show": ["radius", "arc", "area"], "highlight": "area"}]}]}],
             "notes": "Area is in square units.", "evidence": ["c2", "o2"]},
            {"id": "mistakes", "kicker": "WATCH OUT", "title": "A mistake worth catching",
             "rows": [{"cols": [{"elements": [{"type": "compare",
                                               "left": {"title": "Mistake", "tone": "bad", "items": ["Using d as r"]},
                                               "right": {"title": "Fix", "tone": "good", "items": ["Halve d first"]}}]}]}],
             "notes": "Common error.", "evidence": ["m1"]},
            {"id": "we1", "kicker": "WORKED EXAMPLE", "title": "Area when r = 6 cm",
             "rows": [{"cols": [{"elements": [{"type": "steps", "source": "W1", "from_step": 1, "to_step": 3,
                                               "items": [{"expr": "A = πr²", "body": "formula"},
                                                         {"expr": "A = π × 6² = 36π", "body": "square r"},
                                                         {"expr": "3.14 × 36 = 113.04", "body": "multiply"}]}]}]}],
             "notes": "Walk through.", "evidence": ["W1"]},
            {"id": "exit", "background": "dark", "kicker": "EXIT TICKET", "title": "One minute, two answers",
             "rows": [{"cols": [{"elements": [{"type": "questions", "items": [
                 {"q": "If d = 10 cm, what is r?", "answer": "5 cm"},
                 {"q": "Which formula gives area?", "answer": "A = πr²"}]}]}]}],
             "notes": "Collect answers.", "evidence": []},
        ],
    }


class StubClient:
    def __init__(self, reply):
        self.reply = reply
        self.calls = []
        self.session_usage = {"calls": 0, "input_tokens": 0, "output_tokens": 0, "cost_usd": 0.0}

    def analyze(self, prompt, **kw):
        self.calls.append(prompt)
        self.session_usage["calls"] += 1
        r = self.reply(prompt) if callable(self.reply) else copy.deepcopy(self.reply)
        return {"data": r}


@pytest.fixture
def model():
    return from_article(ARTICLE, [])


# ── spec ──────────────────────────────────────────────────────────────

def test_normalise_keeps_every_element_type_and_drops_unknown():
    raw = spec_raw()
    raw["slides"][1]["rows"][0]["cols"][0]["elements"].append({"type": "carousel", "items": ["x"]})
    deck = S.normalise(raw)
    types = {e["type"] for sl in deck.slides for e in sl.elements()}
    assert {"maths_diagram", "chips", "cards", "callout", "formula", "bullets", "compare", "steps",
            "questions"} <= types
    assert "carousel" not in types
    assert any("carousel" in r for r in deck.repairs)


def test_too_few_slides_is_refused():
    raw = spec_raw()
    raw["slides"] = raw["slides"][:3]
    with pytest.raises(S.SpecInvalid):
        S.normalise(raw)


def test_undeclared_asset_is_declared_and_unused_one_is_never_fetched():
    raw = spec_raw()
    raw["assets"].append({"id": "A9", "need": "photo", "subject": "unused"})
    raw["slides"][3]["rows"][0]["cols"][1]["elements"].append({"type": "image", "asset": "A7"})
    deck = S.normalise(raw)
    assert "A7" in deck.assets
    refs = S.asset_refs(deck)
    assert "A9" not in refs and "A7" in refs


# ── worked examples ───────────────────────────────────────────────────

def test_article_steps_split_on_step_markers(model):
    steps = C.article_steps(ARTICLE["worked_examples"][0]["solution_md"])
    assert len(steps) == 4 and steps[0]["line"].startswith("State the formula")


def test_steps_resolve_checks_lines_and_adds_missing_steps(model):
    ctx = C.build(model, article=ARTICLE)
    deck = S.normalise(spec_raw())
    shown = D.resolve_steps(deck, ctx)
    assert shown["W1"] == {1, 2, 3}
    added = D.ensure_examples(deck, ctx, shown)
    # the missing tail is the answer itself: it becomes the answer bar on the
    # same page, not a slide of its own
    assert added == ["W1 steps 4-4 (answer bar)"]
    shown = D.resolve_steps(deck, ctx)
    assert shown["W1"] == {1, 2, 3, 4}
    we = next(e for sl in deck.slides for e in sl.elements() if e["type"] == "steps")
    assert [r["line"] for r in we["resolved"]][1] == "A = π × 6² = 36π"
    assert "113.04 cm²" in we["answer_text"]
    # the full source working rides in the notes whatever the slide shows
    sl = next(s for s in deck.slides if s.id == "we1")
    assert "Full working (W1)" in sl.notes and "Step 3:" in sl.notes


def test_a_step_that_changes_the_maths_falls_back_to_the_source(model):
    ctx = C.build(model, article=ARTICLE)
    raw = spec_raw()
    raw["slides"][6]["rows"][0]["cols"][0]["elements"][0]["items"][2]["expr"] = "3.14 × 36 = 113.4"
    deck = S.normalise(raw)
    D.resolve_steps(deck, ctx)
    we = next(e for sl in deck.slides for e in sl.elements() if e["type"] == "steps")
    assert we["resolved"][2]["line"] != "3.14 × 36 = 113.4"
    assert any("refused" in r for r in deck.repairs)


def test_storyboard_worked_example_is_no_longer_truncated():
    """The storyboard kept `_paginate(...)[0]` — the first page of a long
    solution only. Every page now becomes a slide."""
    from agent5_slides import deck_storyboard as sbd
    # as the article stores it: sentences, whitespace collapsed by from_article
    long = " ".join(f"Step {i}: " + "This line carries one part of the working. " * 4 for i in range(1, 16))
    m = from_article({"title": "T", "worked_examples": [{"problem": "P?", "solution_md": long}]}, [])
    slides = [s for s in sbd.storyboard(m) if s.kind == sbd.WORKED]
    shown = " ".join(b.get("text", "") for s in slides for b in s.blocks)
    assert len(slides) >= 2
    assert all(f"Step {i}:" in shown for i in range(1, 16))     # nothing after page one is lost
    assert slides[1].continued and not slides[1].key_idea


# ── grounding ─────────────────────────────────────────────────────────

def test_new_numbers_and_exposed_answers_are_problems(model):
    ctx = C.build(model, article=ARTICLE)
    raw = spec_raw()
    raw["slides"][3]["rows"][0]["cols"][1]["elements"].append(
        {"type": "statement", "text": "A wheel is 622 mm across."})
    raw["slides"][7]["rows"][0]["cols"][0]["elements"].append(
        {"type": "callout", "kind": "tip", "text": "Hint: 5 cm"})
    deck = S.normalise(raw)
    D.resolve_steps(deck, ctx)
    probs = D.check(deck, ctx)
    assert any("622" in p for p in probs["anatomy"])
    assert any("visible" in p for p in probs["exit"])


def test_correction_replaces_only_the_failing_slide(model):
    ctx = C.build(model, article=ARTICLE)
    deck = S.normalise(spec_raw())
    before = [s.title for s in deck.slides]
    fixed = copy.deepcopy(spec_raw()["slides"][3])
    fixed["title"] = "Parts of a circle"
    client = StubClient({"slides": [fixed]})
    assert D.correct(deck, ctx, client, {"anatomy": ["too dense"]}) == ["anatomy"]
    after = [s.title for s in deck.slides]
    assert after[3] == "Parts of a circle" and after[:3] == before[:3] and after[4:] == before[4:]


def test_maths_check_rejects_wrong_operands():
    src = "A circle with a radius of 6 cm. Use 3.14."
    assert mc.line_ok("A = π × 6² = 36π", src, "3.14")[0]
    assert not mc.line_ok("A = π × 7² = 49π", src, "3.14")[0]
    assert not mc.chain_ok("6² = 12")[0]
    assert mc.chain_ok("C = (22 ÷ 7) × 14 = 44", "22/7")[0]


# ── assets ────────────────────────────────────────────────────────────

def test_licence_gate():
    assert A.licence_ok("CC BY 4.0") and A.licence_ok("Public domain") and A.licence_ok("CC0")
    assert not A.licence_ok("CC BY-NC 2.0") and not A.licence_ok("CC BY-ND 3.0")
    assert not A.licence_ok("CC BY-SA 4.0")            # share-alike only if the gate names it
    assert not A.licence_ok("")


def test_ladder_reuses_shelf_then_generates_under_budget(tmp_path, model):
    ctx = C.build(model, article=ARTICLE)
    deck = S.normalise(spec_raw())
    shelf = A.LocalShelf(tmp_path / "shelf")
    png = _png(tmp_path / "x.png")
    calls = []

    def gen(prompt, aspect):
        calls.append(prompt)
        buf = io.BytesIO()
        Image.open(_png(tmp_path / f"g{len(calls)}.png", colour=(200, 80, 40))).save(buf, "PNG")
        return buf.getvalue(), "stub-image-model", 0.04

    def no_photos(q, http=None):
        return []

    rep = A.resolve(deck, tmp=tmp_path / "t", shelf=shelf, ctx=ctx, budget=A.Budget(max_new_images=1),
                    spend=A.Spend(), generate=gen, photos=no_photos)
    assert len(calls) == 1                                   # budget: one new image
    assert rep.counts().get("generated") == 1 and rep.counts()["missing"] == 1
    assert "no text" in calls[0].lower() or "no captions" in calls[0].lower()
    assert all("whiteboard" not in c for c in calls)         # not the video's style
    # the next deck finds it on the shelf and spends nothing
    calls.clear()
    rep2 = A.resolve(deck, tmp=tmp_path / "t2", shelf=A.LocalShelf(tmp_path / "shelf"), ctx=ctx,
                     budget=A.Budget(max_new_images=0), spend=A.Spend(), generate=gen, photos=no_photos)
    assert rep2.counts().get("shelf") == 1 and not calls
    assert png.exists()


def test_ladder_yields_to_live_users(tmp_path, model):
    ctx = C.build(model, article=ARTICLE)
    deck = S.normalise(spec_raw())
    rep = A.resolve(deck, tmp=tmp_path, shelf=None, ctx=ctx, budget=A.Budget(), spend=A.Spend(),
                    generate=lambda p, a: pytest.fail("generated while a user was live"),
                    photos=lambda q, http=None: [], yield_to_users=lambda: True)
    assert all("never-starve" in v for v in rep.missing.values())


def test_deck_keys_are_refused_by_the_ink_library():
    from shared import visual_library as vl
    assert vl.is_shelf_key("bicycle_wheel__photo__deck")
    assert vl._is_colour_row({"asset_key": "bicycle_wheel__photo__deck"})
    assert vl.find("bicycle_wheel__photo__deck", "a bicycle wheel") is None


# ── render + QA ───────────────────────────────────────────────────────

def _built(tmp_path, model, raw=None):
    ctx = C.build(model, article=ARTICLE)
    deck = S.normalise(raw or spec_raw())
    D.ensure_examples(deck, ctx, D.resolve_steps(deck, ctx))
    pics = {"A1": A.Picture("A1", "k1", str(_png(tmp_path / "a1.png")), 800, 600, "licensed", "photo", provider="Wikimedia Commons",
                            licence="CC BY 4.0", author="X", source_url="https://commons.wikimedia.org/x",
                            subject="bicycle wheel"),
            "A2": A.Picture("A2", "k2", str(_png(tmp_path / "a2.png")), 800, 600, "generated", "illustration",
                            model="m", subject="garden")}
    from agent5_slides.director import render as R
    credits = sorted({p.credit for p in pics.values() if p.credit})
    path, reports = R.build(deck, pics, ctx.figures, tmp_path / "deck.pptx", credits=credits)
    return deck, path, reports


def test_render_is_editable_and_answers_live_in_notes(tmp_path, model):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    deck, path, reports = _built(tmp_path, model)
    prs = Presentation(str(path))
    assert len(prs.slides) == len(deck.slides) + 1            # + image credits
    pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert pics == 2                                         # the two card pictures; nothing else is a raster
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert "A = πr²" in texts and "d = 2r" in texts          # formulas are native text
    exit_slide = prs.slides[[s.id for s in deck.slides].index("exit")]
    visible = " ".join(sh.text_frame.text for sh in exit_slide.shapes if sh.has_text_frame)
    assert "5 cm" not in visible
    assert "5 cm" in exit_slide.notes_slide.notes_text_frame.text
    assert "Wikimedia Commons" in " ".join(texts)            # CC BY: attribution required
    assert "generated for SketchCast" not in " ".join(texts)  # our own art needs no credit
    # the circle is native shapes, not a picture
    anatomy = prs.slides[[s.id for s in deck.slides].index("anatomy")]
    assert not any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in anatomy.shapes)


def test_qa_flags_collisions_and_low_contrast():
    from agent5_slides.director.render import Box, SlideReport
    rep = SlideReport("x", 2, boxes=[Box("text", 1, 1, 4, 1, "hello world", 18, "body"),
                                     Box("text", 1.2, 1.2, 4, 1, "overlap me", 18, "body"),
                                     Box("text", 1, 4, 4, 0.5, "faint", 18, "body", fg=(200, 200, 200)),
                                     Box("text", 1, 5, 4, 0.5, "tiny text", 9, "body")])
    cats = {i["category"] for i in Q.check_slide(rep, "T")}
    assert {"overlap", "contrast", "type"} <= cats


def test_built_sample_has_no_high_deterministic_issue(tmp_path, model):
    deck, path, reports = _built(tmp_path, model)
    issues = Q.check_deck(deck, reports)
    high = {k: [i for i in v if i["severity"] == "high"] for k, v in issues.items()}
    assert not any(high.values()), high


def test_pipeline_end_to_end_offline(tmp_path, model):
    client = StubClient(spec_raw())
    shelf = A.LocalShelf(tmp_path / "shelf")

    def gen(prompt, aspect):
        buf = io.BytesIO()
        Image.open(_png(tmp_path / "g.png", colour=(10, 160, 90))).save(buf, "PNG")
        return buf.getvalue(), "stub", 0.04

    res = P.build_director_deck(model, client=client, out_path=tmp_path / "out.pptx", tmp=tmp_path / "w",
                                article=ARTICLE, shelf=shelf, generate=gen, photos=lambda q, http=None: [],
                                render_preview=False, budget=A.Budget(corrections=1))
    assert res.path.exists()
    r = res.report
    assert r["coverage"]["worked_examples"]["W1"] == "4/4"
    assert r["usage"]["images"]["generated"] == 2 and r["usage"]["total_usd"] == pytest.approx(0.08)
    assert set(r["timings_s"]) >= {"context", "design", "grounding", "assets", "render", "qa", "total"}


def test_flag_defaults_off(monkeypatch):
    from agent5_slides import director
    monkeypatch.delenv("DECK_AI_DIRECTOR", raising=False)
    assert director.enabled() is False
    monkeypatch.setenv("DECK_AI_DIRECTOR", "1")
    assert director.enabled() is True


# ── the worker seam ───────────────────────────────────────────────────

def test_build_deck_is_the_storyboard_when_the_flag_is_off(tmp_path, model, monkeypatch):
    from agent5_slides import deck_generator as dg
    monkeypatch.delenv("DECK_AI_DIRECTOR", raising=False)
    client = StubClient(spec_raw())
    p = dg.build_deck(model, tmp_path / "d.pptx", client=client, article=ARTICLE)
    assert p.exists() and client.calls == []                 # the Director was never asked


def test_build_deck_falls_back_to_the_storyboard_when_the_director_fails(tmp_path, model, monkeypatch):
    from agent5_slides import deck_generator as dg
    monkeypatch.setenv("DECK_AI_DIRECTOR", "1")
    client = StubClient({"title": "x", "slides": []})       # never a usable spec
    p = dg.build_deck(model, tmp_path / "d.pptx", client=client, article=ARTICLE)
    assert p.exists() and len(client.calls) == 2             # design + one re-ask, then storyboard


def test_verified_examples_replace_their_prose_copies(model):
    lesson = {"examples": [{"problem": ARTICLE["worked_examples"][0]["problem"], "task": "solve",
                            "steps": [{"after": ["A = 3.14 × 6²"]}, {"after": ["A = 113.04"]}],
                            "final_answer": ["A = 113.04"]}]}
    ctx = C.build(model, article=ARTICLE, maths_lesson=lesson)
    assert list(ctx.sources) == ["V1"] and ctx.sources["V1"].verified


def test_a_malformed_slide_is_salvaged_alone():
    import json
    raw = spec_raw()
    text = json.dumps(raw, ensure_ascii=False)
    # the live failure: an object where an array element belonged
    broken = text.replace('{"type": "chips", "items"', '"chips": {"items"', 1)
    assert broken != text
    deck = S.normalise({"raw_text": broken})
    assert len(deck.slides) == len(raw["slides"]) - 1
    assert any("salvage" in r and "cover" in r for r in deck.repairs)


def test_credits_only_where_a_licence_requires_them(tmp_path):
    def pic(origin, licence):
        return A.Picture("A", "k", "p.png", 10, 10, origin, licence=licence, subject="s", source_url="u")
    assert pic("licensed", "CC BY 4.0").credit
    assert not pic("licensed", "Public domain").credit
    assert not pic("licensed", "CC0").credit
    assert not pic("generated", "").credit and not pic("shelf", "").credit


def test_tidy_removes_a_restated_problem_and_splits_long_workings(model):
    ctx = C.build(model, article=ARTICLE)
    raw = spec_raw()
    we = raw["slides"][6]
    we["rows"][0]["cols"][0]["elements"].insert(0, {
        "type": "callout", "kind": "think",
        "text": "Problem: Calculate the area of a circle with a radius of 6 cm. Use 3.14 for pi."})
    we["rows"][0]["cols"][0]["elements"][1].update({"to_step": 4, "items": []})
    deck = S.normalise(raw)
    D.tidy(deck, ctx)
    D.resolve_steps(deck, ctx)
    sl = next(s for s in deck.slides if s.id == "we1")
    assert not any(e["type"] == "callout" for e in sl.elements())
    D.MAX_STEPS_PER_SLIDE, old = 2, D.MAX_STEPS_PER_SLIDE
    try:
        deck = S.normalise(raw)
        D.tidy(deck, ctx)
        shown = D.resolve_steps(deck, ctx)
    finally:
        D.MAX_STEPS_PER_SLIDE = old
    ids = [s.id for s in deck.slides]
    assert "we1_cont" in ids and shown["W1"] == {1, 2, 3, 4}
    first = next(e for e in deck.slides[ids.index("we1")].elements() if e["type"] == "steps")
    assert first["answer_text"] == ""                      # the answer waits for the last page


def test_every_element_type_renders_on_a_content_slide(tmp_path, model):
    """Each element through the content-slide path (the cover draws chips
    its own way, which once hid a crash in the content-slide chips)."""
    from agent5_slides.director import render as R
    raw = spec_raw()
    raw["slides"][1]["rows"][0]["cols"][0]["elements"].append(
        {"type": "chips", "items": [{"label": "Irrational approximation", "value": "≈ 3.14159"},
                                    {"label": "Fraction", "value": "22/7"}]})
    deck = S.normalise(raw)
    ctx = C.build(model, article=ARTICLE)
    D.resolve_steps(deck, ctx)
    _, reports = R.build(deck, {}, ctx.figures, tmp_path / "d.pptx")
    assert not [f for r in reports for f in r.faults if "FAILED" in f]


def test_elements_written_as_bare_names_are_refused():
    raw = spec_raw()
    for s in raw["slides"]:
        for r in s.get("rows") or []:
            for c in r["cols"]:
                c["elements"] = ["image", "bullets"]
    with pytest.raises(S.SpecInvalid, match="bare names"):
        S.normalise(raw)


def test_an_answer_only_slide_after_a_worked_example_is_merged(model):
    ctx = C.build(model, article=ARTICLE)
    raw = spec_raw()
    raw["slides"].insert(7, {"id": "answers", "title": "Final answers",
                             "rows": [{"cols": [{"elements": [{"type": "answer", "text": "113.04 cm²"}]}]}],
                             "notes": "Read the answer.", "evidence": ["W1"]})
    deck = S.normalise(raw)
    D.tidy(deck, ctx)
    assert "answers" not in [s.id for s in deck.slides]
    assert "Read the answer." in next(s for s in deck.slides if s.id == "we1").notes


def test_figure_with_caption_stays_on_the_slide(tmp_path, model):
    from agent5_slides.director import render as R
    png = _png(tmp_path / "cell.png", 704, 613)
    figs = {"cell": {"caption": "", "parts": [], "png": str(png), "regions": {}, "w": 704, "h": 613,
                     "labellable": False}}
    raw = spec_raw()
    raw["slides"][3]["rows"] = [{"cols": [{"elements": [{"type": "figure", "figure_key": "cell",
                                                        "caption": "Typical animal cell architecture " * 3}]}]}]
    deck = S.normalise(raw)
    _, reports = R.build(deck, {}, figs, tmp_path / "d.pptx")
    rep = next(r for r in reports if r.id == "anatomy")
    assert not [i for i in Q.check_slide(rep, "x") if i["category"] == "bounds"]


def test_an_overfull_column_degrades_instead_of_overflowing(tmp_path, model):
    from agent5_slides.director import render as R
    raw = spec_raw()
    long = "A definition that runs on for a good while so that four of them cannot stack. " * 2
    raw["slides"][3]["rows"][0]["cols"][1]["elements"] = [
        {"type": "cards", "direction": "column",
         "items": [{"title": t, "badge": "point", "body": long} for t in ("Centre", "Radius", "Diameter", "Chord")]},
        {"type": "callout", "kind": "rule", "text": "A chord through the centre is a diameter."}]
    deck = S.normalise(raw)
    _, reports = R.build(deck, {}, {}, tmp_path / "d.pptx")
    rep = next(r for r in reports if r.id == "anatomy")
    assert not [i for i in Q.check_slide(rep, "x") if i["severity"] == "high"]
    assert rep.degraded                                   # what moved is recorded…
    assert "Shortened on the slide" in next(s for s in deck.slides if s.id == "anatomy").notes  # …and kept


def test_labels_keep_maths_case_and_long_words_never_break():
    from agent5_slides.director import render as R
    assert R.caps("r = d ÷ 2") == "r = d ÷ 2" and R.caps("Boundary") == "BOUNDARY"
    assert R.caps("Pi ratio (C ÷ d)") == "Pi ratio (C ÷ d)"
    assert R.caps("Area of r") == "Area of r"
    R._UNFIT[0] = 0
    R._h("Circumference", 3.0, 26)
    assert R._UNFIT[0] == 0
    R._h("Circumference", 1.2, 26)                             # one word wider than its box
    assert R._UNFIT[0] == 1
    col = S.Column(1, [{"type": "cards", "direction": "row", "tone": "accent",
                        "items": [{"title": "Circumference", "tone": "accent"}] * 4}])
    env = R.Env(R.Theme.make("indigo"), {}, {}, dark=False)
    assert R._measure(col, 4.0, env)[1] == float("inf")      # that layout does not fit


def test_a_reused_cc_by_photo_keeps_its_credit(tmp_path, model):
    ctx = C.build(model, article=ARTICLE)
    deck = S.normalise(spec_raw())
    shelf = A.LocalShelf(tmp_path / "shelf")
    photo = _png(tmp_path / "wheel.png")

    def photos(q, http=None):
        return [{"title": "Bicycle wheel", "url": "https://x/w.jpg", "page": "https://flickr.com/p/1",
                 "w": 1024, "h": 768, "mime": "image/jpeg", "licence": "CC BY 2.0", "author": "Ann",
                 "provider": "flickr"}]

    class Http:
        def get(self, url, **kw):
            class R:
                content = photo.read_bytes()
                def raise_for_status(self):
                    pass
            return R()

    class Looks:                                             # a search hit is used only once looked at
        def analyze_image(self, path, prompt, max_tokens=0):
            return {"data": {"shows_subject": True, "relevance": 9, "quality": 8, "classroom_safe": True,
                             "has_text": False, "scientifically_wrong": False}}

    gen = lambda p, a: (None, "", 0.0)                       # noqa: E731 — nothing is generated
    A.resolve(deck, tmp=tmp_path / "t1", shelf=shelf, ctx=ctx, budget=A.Budget(), spend=A.Spend(),
              generate=gen, photos=photos, http=Http(), vision_client=Looks())
    rep = A.resolve(deck, tmp=tmp_path / "t2", shelf=A.LocalShelf(tmp_path / "shelf"), ctx=ctx,
                    budget=A.Budget(max_new_images=0), spend=A.Spend(), generate=gen,
                    photos=lambda q, http=None: [])
    a1 = rep.pictures["A1"]
    assert a1.origin == "shelf" and a1.credit and "Ann" in a1.credit and "CC BY 2.0" in a1.credit


def test_a_fallback_is_recorded_on_the_generation(tmp_path, model, monkeypatch):
    from agent5_slides import deck_generator as dg
    from worker import client as db
    seen = []
    monkeypatch.setattr(db, "merge_generation_params", lambda sb, gid, patch: seen.append((gid, patch)))
    monkeypatch.setenv("DECK_AI_DIRECTOR", "1")
    monkeypatch.setattr("agent5_slides.deck_art.user_builders_live", lambda sb, j: False)
    dg.build_deck(model, tmp_path / "d.pptx", client=StubClient({"title": "x", "slides": []}), article=ARTICLE,
                  sb=object(), generation_id="gen-1")
    assert seen and seen[0][0] == "gen-1"
    assert seen[0][1]["deck_director"]["route"] == "storyboard_fallback"


def test_the_total_budget_stops_corrections_and_vision(tmp_path, model):
    """DECK_DIRECTOR_TOTAL_USD is enforced: a deck whose design alone spent
    the ceiling gets no correction round and no paid vision review."""
    class Costly(StubClient):
        def analyze(self, prompt, **kw):
            self.session_usage["cost_usd"] += 2.0             # one call = the whole ceiling
            return super().analyze(prompt, **kw)
    raw = spec_raw()
    raw["slides"] = raw["slides"] + [dict(s, id=s["id"] + "_x") for s in raw["slides"][1:4]]   # a full deck
    raw["slides"][3]["rows"][0]["cols"][1]["elements"].append(
        {"type": "statement", "text": "A wheel is 622 mm across."})   # a grounding problem to correct
    client = Costly(raw)
    res = P.build_director_deck(model, client=client, out_path=tmp_path / "o.pptx", tmp=tmp_path / "w",
                                article=ARTICLE, generate=lambda p, a: (None, "", 0.0),
                                photos=lambda q, http=None: [], render_preview=False,
                                budget=A.Budget(max_total_usd=1.5, corrections=2))
    assert len(client.calls) == 1                             # the design call only
    assert any("grounding correction skipped" in s for s in res.report["budget_stops"])


def test_rtl_pass_tags_arabic_runs_and_pins_subscripts(tmp_path):
    """PowerPoint lays out an rtl paragraph with an LTR base unless its runs
    carry an Arabic language — every colon and bracket on the wrong side."""
    from pptx import Presentation
    from pptx.util import Inches
    from agent5_slides.slide_builder import _mirror_deck_rtl
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    tb.text_frame.text = "النباتات: تصنع سكريات من CO₂ والماء"
    tb2 = s.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    tb2.text_frame.text = "A = πr²"                         # a Latin-only box keeps its language
    p = tmp_path / "rtl.pptx"
    prs.save(p)
    _mirror_deck_rtl(p)
    prs = Presentation(p)
    r = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert r.font._element.get("lang") == "ar-SA"
    assert "CO₂\u200e" in r.text
    r2 = prs.slides[0].shapes[1].text_frame.paragraphs[0].runs[0]
    assert r2.font._element.get("lang") != "ar-SA"


def test_two_steps_blocks_of_one_example_on_a_slide_become_one(model):
    ctx = C.build(model, article=ARTICLE)
    raw = spec_raw()
    col = raw["slides"][6]["rows"][0]["cols"][0]
    col["elements"] = [{"type": "steps", "source": "W1", "from_step": 1, "to_step": 2},
                       {"type": "steps", "source": "W1", "from_step": 3, "to_step": 4}]
    deck = S.normalise(raw)
    D.tidy(deck, ctx)
    sl = next(s for s in deck.slides if s.id == "we1")
    blocks = [e for e in sl.elements() if e["type"] == "steps"]
    assert len(blocks) == 1 and (blocks[0]["from_step"], blocks[0]["to_step"]) == (1, 4)


def test_a_licensed_hit_that_is_not_the_thing_itself_is_refused():
    base = {"shows_subject": True, "relevance": 8, "quality": 8, "classroom_safe": True}
    assert A.verdict_ok(base | {"is_the_thing_itself": True}, "licensed")
    assert not A.verdict_ok(base | {"is_the_thing_itself": False}, "licensed")
    assert not A.verdict_ok(base | {"relevance": 6}, "licensed")          # 6 is not enough for a search hit
    assert A.verdict_ok(base | {"relevance": 6}, "generated")


def test_english_figure_labels_become_numbers_on_a_jawi_deck(tmp_path, model):
    from pptx import Presentation
    from agent5_slides.director import render as R
    png = _png(tmp_path / "cell.png", 800, 600)
    figs = {"cell": {"caption": "", "parts": ["nucleus", "cell membrane"], "png": str(png),
                     "regions": {"nucleus": [[300, 200, 420, 320]], "cell membrane": [[100, 80, 700, 540]]},
                     "w": 800, "h": 600, "labellable": True}}
    raw = spec_raw()
    raw["slides"][3]["rows"] = [{"cols": [{"elements": [{"type": "figure", "figure_key": "cell",
                                                        "parts": ["nucleus", "cell membrane"],
                                                        "labels": [{"part": "nucleus", "text": "نوکليوس"}]}]}]}]
    deck = S.normalise(raw)
    path, reports = R.build(deck, {}, figs, tmp_path / "d.pptx", language="ms-arab", direction="rtl")
    slide = Presentation(path).slides[3]
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert "نوکليوس" in texts and "cell membrane" not in texts and "1" in texts
    assert "1 = cell membrane" in slide.notes_slide.notes_text_frame.text
    rep = next(r for r in reports if r.id == "anatomy")
    assert not [i for i in Q.check_slide(rep, "x") if i["category"] == "bounds"]


def test_a_maths_diagram_outside_maths_is_corrected_then_removed(model):
    ctx = C.build(model, article=ARTICLE, subject="chemistry")
    raw = spec_raw()
    raw["slides"][4]["rows"][0]["cols"][1]["elements"].append({"type": "callout", "kind": "tip", "text": "Square units."})
    deck = S.normalise(raw)
    probs = D.check(deck, ctx)
    assert any("maths_diagram" in p for p in probs.get("area", []))
    D.drop_unfixable(deck, {k: v for k, v in probs.items() if k == "area"}, ctx)
    sl = next(s for s in deck.slides if s.id == "area")
    assert not any(e["type"] == "maths_diagram" for e in sl.elements())


def test_rows_rebalance_before_anything_shrinks():
    from agent5_slides.director import render as R
    env = R.Env(R.Theme.make("indigo"), {}, {}, dark=False)
    long = {"type": "bullets", "items": [{"body": "A point that needs a full line of room " * 6}] * 5}
    rows = [S.Row(1, [S.Column(1, [long])]), S.Row(1, [S.Column(1, [{"type": "statement", "text": "Short."}])])]
    h = R.row_heights(rows, 5.0, env)
    assert h[0] > h[1] and abs(sum(h) - 5.0) < 1e-6


def test_a_missing_picture_is_pruned_before_layout(tmp_path, model):
    from agent5_slides.director import render as R
    deck = S.normalise(spec_raw())
    sl = next(s for s in deck.slides if s.id == "anatomy")
    sl.rows[0].cols.append(S.Column(1, [{"type": "image", "asset": "A_missing", "tone": "accent"}]))
    _, reports = R.build(deck, {}, {}, tmp_path / "d.pptx")
    assert not any(c.elements and c.elements[0].get("asset") == "A_missing" for r in sl.rows for c in r.cols)
    rep = next(r for r in reports if r.id == "anatomy")
    assert not [f for f in rep.faults if "unavailable" in f]


def test_a_whole_lesson_in_six_slides_is_asked_again(model):
    ctx = C.build(model, article=ARTICLE)
    short, full = spec_raw(), spec_raw()
    full["slides"] = full["slides"] + [dict(s, id=s["id"] + "_x") for s in full["slides"][1:4]]
    replies = iter([short, full])
    client = StubClient(lambda prompt: next(replies))
    deck = D.design(ctx, client)
    assert len(client.calls) == 2 and len(deck.slides) >= 10


def test_the_answer_bar_drops_the_whole_instruction_sentence():
    assert D.notation_answer("State the final answer clearly with units. The value of y is 100 degrees.") \
        == "The value of y is 100 degrees."
    assert D.notation_answer("State the final answers with correct units. The area is 77 square centimetres.") \
        == "The area is 77 cm²."


def test_a_worked_example_owns_its_column(tmp_path, model):
    """Live deck, Real-Life Graphs slide 10: chips under the steps hid step 3."""
    from agent5_slides.director import render as R
    ctx = C.build(model, article=ARTICLE)
    raw = spec_raw()
    col = raw["slides"][6]["rows"][0]["cols"][0]
    col["elements"][0]["items"] = []
    col["elements"][0]["to_step"] = 4
    col["elements"].append({"type": "chips", "items": [{"label": f"Part {i}", "value": "a value " * 4}
                                                       for i in range(4)]})
    col["elements"].append({"type": "paragraph", "text": "A long paragraph that will not fit here. " * 6})
    deck = S.normalise(raw)
    D.resolve_steps(deck, ctx)
    path, reports = R.build(deck, {}, {}, tmp_path / "d.pptx")
    sl = next(s for s in deck.slides if s.id == "we1")
    rep = next(r for r in reports if r.id == "we1")
    assert [e["type"] for e in sl.elements()] == ["steps"]
    assert "Part 0" in sl.notes
    assert not [i for i in Q.check_slide(rep, "x") if i["severity"] == "high"]


def test_a_deck_with_layout_faults_falls_back_to_the_storyboard(tmp_path, model, monkeypatch):
    from agent5_slides import deck_generator as dg
    from agent5_slides.director import pipeline as Pl
    from worker import client as db
    seen = []
    monkeypatch.setattr(db, "merge_generation_params", lambda sb, gid, patch: seen.append(patch))
    monkeypatch.setenv("DECK_AI_DIRECTOR", "1")
    monkeypatch.setattr("agent5_slides.deck_art.user_builders_live", lambda sb, j: False)

    class Res:
        path = tmp_path / "director.pptx"
        report = {"blocking_layout": {"s10": ["'a' overlaps 'b'"]}, "usage": {"total_usd": 0.4},
                  "timings_s": {"total": 300}}
    monkeypatch.setattr(Pl, "build_director_deck", lambda *a, **k: Res())
    out = dg.build_deck(model, tmp_path / "d.pptx", client=StubClient(spec_raw()), article=ARTICLE,
                        sb=object(), generation_id="g")
    assert out == tmp_path / "d.pptx" and out.exists()                # the storyboard deck
    assert seen[0]["deck_director"]["route"] == "storyboard_fallback"

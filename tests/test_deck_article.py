"""The book deck is built from an ARTICLE authored off the chapter — the
catalogue deck's own shape — so a class reads claims, never narration.

Founder, 2026-09-12, on the Materials deck: "it looks like now we are
printing the narration itself back on the slide. we don't need to do that…
use the text from the book as the underlying article and build the deck the
same way we did for the internal library."
"""
from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from agent5_slides import deck_article as da
from agent5_slides import deck_art
from agent5_slides.deck_storyboard import SECTION, distil_points, section_content, storyboard, summarise
from shared.lesson_model import Figure, LessonModel, Section, from_article
from tests.test_deck_generation import ANALYSIS, BOOK, CHAPTER, _StubClient
from tests.test_deck_generator import BOOK_ARTICLE_REPLY as ARTICLE_REPLY, _env
from catalogue.key import canonical_key

KEY = canonical_key("particle_states")     # the validator canonicalises figure keys


def _all_text(prs) -> list[str]:
    """Every string on the slides — text frames AND table cells (the glossary
    is a table)."""
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if getattr(sh, "has_text_frame", False):
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False):
                out.extend(c.text for c in sh.table.iter_cells())
    return out



class TestTheAuthoringCall:
    def test_the_reply_is_validated_into_the_catalogue_shape(self):
        client = _StubClient(ARTICLE_REPLY)
        art = da.author_article(BOOK, CHAPTER, ANALYSIS, client, {}, "en", title="Materials · Part 1")
        assert [s["id"] for s in art["sections"]] == ["s1", "s2", "s3"]
        assert art["figures"][0]["figure_key"] == KEY
        assert art["figures"][0]["spec"]["parts"] == ["solid", "liquid", "gas"]
        assert art["sections"][0]["figure_keys"] == [KEY]
        assert len(art["claims"]) == 5 and art["language"] == "en"
        assert len(client.calls) == 1
        call = client.calls[0]
        assert "chapter above" in call["prompt"] and call["cache_prefix"].startswith(
            da.chapter_grounding(BOOK, CHAPTER, ANALYSIS)[:40]), "the documents' grounding block, cached"
        assert call["response_schema"] is da.RESPONSE_SCHEMA

    def test_a_thin_reply_fails_the_job_before_anything_is_built(self):
        with pytest.raises(RuntimeError, match="deck article authoring"):
            da.author_article(BOOK, CHAPTER, ANALYSIS, _StubClient({}), {}, "en")
        thin = {**ARTICLE_REPLY, "sections": ARTICLE_REPLY["sections"][:1]}
        with pytest.raises(RuntimeError, match="usable section"):
            da.author_article(BOOK, CHAPTER, ANALYSIS, _StubClient(thin), {}, "en")

    def test_the_prompt_asks_for_claims_as_short_bullet_points(self):
        p = da.build_article_prompt("en")
        assert "at most 14 words" in p and "never a narration line" in p
        assert "never shown on a slide" in p

    def test_the_language_reaches_the_prompt_and_jawi_is_two_scripts(self):
        assert da.build_article_prompt("ms") != da.build_article_prompt("en")
        assert "JAWI" in da.build_article_prompt("ms-arab")

    def test_coverage_reads_headings_claims_and_prose(self):
        segs = da.coverage_segments({**ARTICLE_REPLY})
        assert segs[0]["slide_heading"] == "Three states of matter"
        assert segs[0]["slide_points"] == ["Matter exists as solids, liquids and gases.",
                                           "A solid keeps its shape and its volume."]
        assert "Matter exists" in segs[0]["text"]


class TestTheBookPathIsTheCataloguePath:
    def test_the_deck_shows_claims_never_the_prose(self, monkeypatch, tmp_path):
        monkeypatch.delenv("DECK_STORYBOARD", raising=False)
        sb, uploads, process = _env(monkeypatch)
        client = _StubClient(ARTICLE_REPLY)
        title = process._generate_deck(
            sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, client, {},
            {}, "en", "ltr", tmp_path, "u1/gen-1", "Materials · Part 1",
        )
        assert len(client.calls) == 1, "one authoring call, as before"
        assert [d for _, d in uploads] == ["u1/gen-1/deck.pptx"]
        assert title == "Living Things · Materials · Part 1 · Slide deck"
        prs = Presentation(uploads[0][0])
        slide_text = _all_text(prs)
        notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
        assert any("A solid keeps its shape and its volume." in t for t in slide_text), "a claim is a point"
        assert not any("spreading out to take both its shape" in t for t in slide_text), "prose never on a slide"
        assert any("spreading out to take both its shape" in n for n in notes), "prose is the teacher's notes"
        assert any("Words to know" in t for t in slide_text) and any("vacuum" in t for t in slide_text)
        assert any("A vacuum is a gas with very few particles." in t for t in slide_text), "misconceptions slide"
        assert any("Classify tap water." in t for t in slide_text), "worked example"
        cov = (sb.tables["generations"][0].get("params") or {}).get("coverage")
        assert cov and cov[0]["kind"] == "deck" and cov[0]["model"] == "stub-model"

    def test_a_declared_figure_without_artwork_is_not_a_picture(self):
        model = from_article({**ARTICLE_REPLY, "language": "en"}, ARTICLE_REPLY["figures"], art=None)
        s1 = model.sections[0]
        assert s1.figure_keys == ["particle_states"] and not model.figures["particle_states"].annotatable, \
            "from_article takes the keys as written; the validator canonicalised them upstream"
        assert not deck_art.pictured(model, s1)
        assert deck_art.declared_figure(model, s1).parts == ["solid", "liquid", "gas"]
        # The storyboard draws no diagram slide and no illustration for it.
        kinds = summarise(storyboard(model))
        assert "diagram" not in kinds and kinds[SECTION] == 3

    def test_the_declared_spec_names_the_parts_the_generator_must_draw(self, monkeypatch, tmp_path):
        """The generate rung asks for the ARTICLE's figure — its key, caption
        and parts — not a spec invented from the heading."""
        model = from_article({**ARTICLE_REPLY, "language": "en"}, ARTICLE_REPLY["figures"], art=None)
        prompts: list[str] = []
        png = tmp_path / "art" / "particle_states.png"
        png.parent.mkdir(parents=True)
        png.write_bytes(b"png")

        class _Backend:
            set_yield = None

            def set_context(self, **kw):
                pass

            def budget_exhausted(self):
                return False

            def generate(self, key, prompt):
                prompts.append(f"{key}: {prompt}")
                return object()

            def publish(self, *a):
                pass

        monkeypatch.setattr(deck_art, "_backend", lambda: _Backend())
        monkeypatch.setattr(deck_art, "user_builders_live", lambda sb, ex: False)
        import catalogue.figures as cf
        monkeypatch.setattr(cf, "lookup_asset", lambda sb, r: {"asset_key": "particle_states", "storage_path": "x"})
        monkeypatch.setattr(deck_art, "figure_from_row", lambda sb, row, tmp, caption="": Figure(
            key="particle_states", caption=caption, png=png, regions={"solid": [[0, 0, 1, 1]]}, w=10, h=10,
            parts=["solid"]))
        n = deck_art.generate_figures(model, object(), tmp_path, {}, "job-1", 2, "job-1")
        assert n == 2, "the declared figure once (two sections share it), then one from a heading"
        assert [p.split(":")[0] for p in prompts] == ["particle_states", "compressing_matter"]
        assert "solid, liquid, gas" in prompts[0], prompts
        assert deck_art.pictured(model, model.sections[0]) and deck_art.pictured(model, model.sections[1])
        assert model.sections[0].figure_keys == ["particle_states"]


class TestNeverAParagraphOnASlide:
    def test_a_section_with_no_claims_gets_short_sentences_not_prose(self):
        sec = Section(id="s1", heading="Inside a liquid",
                      body_md="In a liquid, the particles still touch each other, but the forces holding "
                              "them are weaker. This lets them slide past one another. It is like a "
                              "crowded room where people are close together but can still move around "
                              "freely between one another without ever leaving the room itself.")
        blocks, notes = section_content(sec)
        assert blocks and blocks[0]["kind"] == "list"
        items = blocks[0]["items"]
        assert 1 <= len(items) <= 3 and all(len(i) <= 90 for i in items), items
        assert "This lets them slide past one another." in items
        assert "crowded room" in notes, "the prose is the notes"

    def test_a_long_sentence_is_cut_at_its_first_clause(self):
        pts = distil_points([{"kind": "para", "text": (
            "Solids keep a fixed shape because their particles are locked tightly in a regular pattern "
            "that only lets them vibrate about fixed positions without ever changing places with each other.")}])
        assert pts == ["Solids keep a fixed shape"]

    def test_narration_on_the_video_route_is_distilled_too(self):
        m = LessonModel(title="L", sections=[Section(
            id="s1", heading="H", body_md="", narration="Now, what happens if we warm things up? "
            "The particles break free from that neat grid. Spot on! In a liquid, the particles still touch.")])
        m.sections[0].body_md = m.sections[0].narration
        blocks, notes = section_content(m.sections[0])
        assert blocks[0]["kind"] == "list" and len(blocks[0]["items"]) == 3
        assert notes == m.sections[0].narration


class TestTheVideosPicturesRespectTheArticlesPlans:
    def test_a_section_that_planned_its_own_diagram_keeps_the_slot(self, tmp_path):
        """The cycle diagram the video drew landed on "Particle arrangement"
        (which had planned its own figure) instead of "State changes", which
        had nothing planned. A section with a plan keeps its slot for it."""
        png = tmp_path / "state_cycles.png"
        png.write_bytes(b"png")
        model = from_article({**ARTICLE_REPLY, "language": "en"}, ARTICLE_REPLY["figures"], art=None)
        model.sections.append(Section(id="s4", heading="State changes through heating and cooling",
                                      body_md="Ice melts to liquid water; water boils to steam.",
                                      points=["Heating melts ice to water and boils water to steam."]))
        fig = Figure(key="state_cycles", caption="ice, liquid water and steam with arrows", png=png)
        n = deck_art.place_video_figures(model, [(1, "state_cycles")], {"state_cycles": fig}, 3, False, 4)
        assert n == 1
        assert "state_cycles" in model.sections[3].figure_keys, "the bare section, not the planned one"
        assert model.sections[1].figure_keys == ["particle_states"]

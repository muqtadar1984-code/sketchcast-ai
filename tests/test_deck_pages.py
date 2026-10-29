"""Deck pages for the Classroom Board (founder, 2026-10-28): every deck slide
as a PNG plus ONE text map (titles, notes, words with boxes) so a teacher can
highlight part of a page and ask SK about it.

No network and no real renderer: PDFs are made here with PyMuPDF, decks with
python-pptx, PNGs with Pillow, and the renderer is injected."""

from __future__ import annotations

import json
import logging
import os
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import fitz
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from agent5_slides import deck_pages as DP
from worker import client as db

# ── builders ────────────────────────────────────────────────────────────


def _pptx(path: Path, slides: list[dict], *, title_layout: bool = True) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for spec in slides:
        layout = prs.slide_layouts[5 if title_layout else 6]   # "Title Only" / "Blank"
        s = prs.slides.add_slide(layout)
        if title_layout and spec.get("title") is not None:
            s.shapes.title.text = spec["title"]
        if spec.get("notes"):
            s.notes_slide.notes_text_frame.text = spec["notes"]
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def _pdf(path: Path, pages: list[list[tuple[float, float, str, float]]]) -> Path:
    """pages: [[(x, y, text, fontsize), …], …] on a 960x540 pt slide."""
    doc = fitz.open()
    for items in pages:
        page = doc.new_page(width=960, height=540)
        for x, y, text, size in items:
            page.insert_text((x, y), text, fontsize=size)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return path


def _pngs(d: Path, n: int, size=(1600, 900), start: int = 1) -> list[Path]:
    d.mkdir(parents=True, exist_ok=True)
    out = []
    for i in range(start, start + n):
        p = d / f"slide_{i:02d}.png"
        Image.new("RGB", size, "white").save(p)
        out.append(p)
    return out


def _fake_renderer(calls: list | None = None, how: str = "fake"):
    """Renders a pptx the way render_images does, without a renderer: one PDF
    page per slide carrying the slide's title, one PNG per page."""

    def render(pptx, out_dir, pdf=None):
        if calls is not None:
            calls.append(Path(pptx))
        titles = []
        for s in Presentation(str(pptx)).slides:
            t = s.shapes.title
            titles.append(t.text_frame.text if t is not None else "untitled")
        _pdf(Path(pdf), [[(60, 80, t, 40), (60, 200, "body words here", 18)] for t in titles])
        return _pngs(Path(out_dir), len(titles)), how

    return render


# ── build_text_map ──────────────────────────────────────────────────────


def test_titles_notes_and_word_boxes(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [
        {"title": "Mitochondria", "notes": "Say: the powerhouse."},
        {"title": "Chloroplasts", "notes": ""},
    ])
    # drawn bottom line FIRST: the map must still read top to bottom
    pdf = _pdf(tmp_path / "deck.pdf", [
        [(60, 300, "second line below", 18), (60, 80, "Mitochondria make energy", 40)],
        [(480, 270, "Chloroplasts", 30)],
    ])
    pngs = _pngs(tmp_path / "png", 2)
    m = DP.build_text_map(pdf, pptx, pngs, renderer="fake")
    assert m["v"] == 1 and m["renderer"] == "fake"
    s1, s2 = m["slides"]
    assert (s1["n"], s1["png"], s1["w"], s1["h"], s1["part"]) == (1, "slide_01.png", 1600, 900, 1)
    assert s2["n"] == 2 and s2["png"] == "slide_02.png"
    assert s1["title"] == "Mitochondria" and s2["title"] == "Chloroplasts"
    assert s1["notes"] == "Say: the powerhouse." and s2["notes"] == ""
    assert s1["text"].split("\n") == ["Mitochondria make energy", "second line below"]
    assert [w["t"] for w in s1["words"]][:4] == ["Mitochondria", "make", "energy", "second"]
    for w in s1["words"] + s2["words"]:
        x0, y0, x1, y1 = w["b"]
        assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1
    # the box is where the word was drawn: x=60/960, baseline 80/540
    x0, y0, x1, y1 = s1["words"][0]["b"]
    assert abs(x0 - 60 / 960) < 0.01 and y0 < 80 / 540 < y1 + 0.02
    # a word drawn at the slide's centre lands near the centre
    cx = (s2["words"][0]["b"][0] + s2["words"][0]["b"][2]) / 2
    assert 0.5 < cx < 0.7
    assert s1["blocks"] and all(0 <= v <= 1 for b in s1["blocks"] for v in b["b"])
    assert any("Mitochondria make energy" in b["t"] for b in s1["blocks"])
    json.dumps(m)   # serialisable as is


def test_title_falls_back_to_the_largest_line_without_a_placeholder(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": None, "notes": "n"}], title_layout=False)
    pdf = _pdf(tmp_path / "deck.pdf", [[(60, 300, "small body", 14), (60, 90, "Plant Cells", 44)]])
    m = DP.build_text_map(pdf, pptx, _pngs(tmp_path / "png", 1))
    assert m["slides"][0]["title"] == "Plant Cells" and m["slides"][0]["notes"] == "n"


def test_caps(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": "Many words", "notes": "x" * 5000}])
    lines = [(10, 12 + 7 * i, " ".join(f"w{i}_{j}" for j in range(40)), 5) for i in range(75)]
    pdf = _pdf(tmp_path / "deck.pdf", [lines])
    s = DP.build_text_map(pdf, pptx, _pngs(tmp_path / "png", 1))["slides"][0]
    assert len(s["words"]) == DP.MAX_WORDS == 1500
    assert len(s["text"]) <= DP.MAX_TEXT == 6000
    assert len(s["notes"]) == DP.MAX_NOTES == 4000


def test_page_and_png_counts_must_agree(tmp_path):
    pdf = _pdf(tmp_path / "deck.pdf", [[(60, 80, "a", 20)], [(60, 80, "b", 20)]])
    with pytest.raises(ValueError, match="2 page"):
        DP.build_text_map(pdf, None, _pngs(tmp_path / "png", 1))


def test_a_pptx_that_disagrees_gives_no_titles_from_it(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": "One"}, {"title": "Two"}, {"title": "Three"}])
    pdf = _pdf(tmp_path / "deck.pdf", [[(60, 80, "Only page", 30)]])
    s = DP.build_text_map(pdf, pptx, _pngs(tmp_path / "png", 1))["slides"][0]
    assert s["title"] == "Only page", "taken from the page, never a different slide's title"


# ── export ──────────────────────────────────────────────────────────────


def test_export_with_an_injected_renderer(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": "A"}, {"title": "B"}])
    calls = []
    res = DP.export_deck_pages(pptx, tmp_path / "work", _fake_renderer(calls))
    assert calls == [pptx] and res["renderer"] == "fake" and len(res["pngs"]) == 2
    assert [s["title"] for s in res["map"]["slides"]] == ["A", "B"]
    assert res["pdf"].exists()


def test_no_pngs_raises_with_the_renderers_reason(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": "A"}])
    with pytest.raises(DP.DeckRenderUnavailable, match="no renderer available"):
        DP.export_deck_pages(pptx, tmp_path / "w", lambda p, o, pdf=None: ([], "no renderer available (x)"))


def test_a_render_over_budget_is_refused(tmp_path):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": "A"}])
    inner = _fake_renderer()

    def slow(p, o, pdf=None):
        time.sleep(0.05)
        return inner(p, o, pdf=pdf)

    with pytest.raises(DP.DeckRenderUnavailable, match="budget"):
        DP.export_deck_pages(pptx, tmp_path / "w", slow, timeout=0)


def test_the_default_renderer_is_qa_render_images_with_the_timeout(monkeypatch, tmp_path):
    from agent5_slides.director import qa

    seen = {}

    def fake_render_images(pptx, out_dir, pdf=None, timeout=600):
        seen["timeout"] = timeout
        return [], "nothing"

    monkeypatch.setattr(qa, "render_images", fake_render_images)
    with pytest.raises(DP.DeckRenderUnavailable):
        DP.export_deck_pages(_pptx(tmp_path / "d.pptx", [{"title": "A"}]), tmp_path / "w", timeout=120)
    assert seen == {"timeout": 120}


def test_the_directors_qa_render_is_reused_only_when_newer_than_the_pptx(tmp_path):
    deck_dir = tmp_path / "deck"
    pptx = _pptx(deck_dir / "deck.pptx", [{"title": "A"}, {"title": "B"}])
    qa1 = deck_dir / "director" / "qa_1"
    _pdf(qa1 / "deck.pdf", [[(60, 80, "A", 30)], [(60, 80, "B", 30)]])
    _pngs(qa1, 2)
    _pdf(deck_dir / "director" / "qa_0" / "deck.pdf", [[(60, 80, "old", 30)]])
    now = time.time()
    os.utime(pptx, (now - 100, now - 100))
    calls = []
    res = DP.export_deck_pages(pptx, tmp_path / "w", _fake_renderer(calls), reuse_qa=True)
    assert calls == [] and "director qa" in res["renderer"] and len(res["pngs"]) == 2
    # a storyboard fallback rewrote the pptx after QA: render afresh
    os.utime(pptx, (now + 100, now + 100))
    res = DP.export_deck_pages(pptx, tmp_path / "w2", _fake_renderer(calls), reuse_qa=True)
    assert calls == [pptx] and res["renderer"] == "fake"


def test_parts_number_on_across_the_deck(tmp_path):
    p1 = _pptx(tmp_path / "in" / "deck.pptx", [{"title": "P1 a"}, {"title": "P1 b"}])
    p2 = _pptx(tmp_path / "in" / "deck_part2.pptx", [{"title": "P2 a"}, {"title": "P2 b"}, {"title": "P2 c"}])
    res = DP.export_deck_parts([(2, p2), (1, p1)], tmp_path / "out", _fake_renderer())
    slides = res["map"]["slides"]
    assert [s["n"] for s in slides] == [1, 2, 3, 4, 5]
    assert [s["part"] for s in slides] == [1, 1, 2, 2, 2]
    assert [s["png"] for s in slides] == [f"slide_{i:02d}.png" for i in range(1, 6)]
    assert [s["title"] for s in slides] == ["P1 a", "P1 b", "P2 a", "P2 b", "P2 c"]
    assert len(res["pngs"]) == 5


def test_deck_part_of():
    assert DP.deck_part_of("u/g/deck.pptx") == 1
    assert DP.deck_part_of("u/g/deck_part3.pptx") == 3
    assert DP.deck_part_of("u/g/other.pptx") is None
    assert DP.deck_part_of("u/g/mydeck.pptx") is None


# ── publish ─────────────────────────────────────────────────────────────


class _Res:
    def __init__(self, data):
        self.data = data


class _EnumError(Exception):
    def __init__(self):
        super().__init__({"code": "22P02", "message": 'invalid input value for enum artifact_kind: "deck_pages_json"'})
        self.code = "22P02"


class _Q:
    def __init__(self, sb, table):
        self.sb, self.table, self.op, self.payload, self.filters = sb, table, "select", None, []

    def select(self, *a, **k):
        self.op = "select"
        return self

    def insert(self, payload):
        self.op, self.payload = "insert", payload
        return self

    def delete(self):
        self.op = "delete"
        return self

    def update(self, payload):
        self.op, self.payload = "update", dict(payload)
        return self

    def eq(self, col, val):
        self.filters.append((col, val))
        return self

    def _match(self, r):
        return all(r.get(c) == v for c, v in self.filters)

    def execute(self):
        rows = self.sb.tables.setdefault(self.table, [])
        kinds = [v for c, v in self.filters if c == "kind"]
        if self.sb.no_0166 and self.table == "artifacts" and DP.KIND_MAP in kinds:
            raise _EnumError()
        if self.op == "select":
            return _Res([dict(r) for r in rows if self._match(r)])
        if self.op == "insert":
            new = self.payload if isinstance(self.payload, list) else [self.payload]
            if self.sb.no_0166 and any(r.get("kind") == DP.KIND_MAP for r in new):
                raise _EnumError()
            if any(r.get("kind") == "boom" for r in new):
                raise RuntimeError("connection reset")
            self.sb.log.append(("insert", self.table, [dict(r) for r in new]))
            if self.sb.stamp_inserts:   # the column default, now()
                new = [{"created_at": db._stamp(datetime.now(timezone.utc)), **r} for r in new]
            rows.extend(dict(r) for r in new)
            return _Res(new)
        if self.op == "delete":
            self.sb.log.append(("delete", self.table, list(self.filters)))
            self.sb.tables[self.table] = [r for r in rows if not self._match(r)]
            return _Res([])
        if self.op == "update":
            self.sb.log.append(("update", self.table, dict(self.payload), list(self.filters)))
            for r in rows:
                if self._match(r):
                    r.update(self.payload)
            return _Res([])
        raise AssertionError(self.op)


class _Bucket:
    def __init__(self, sb):
        self.sb = sb

    def remove(self, paths):
        self.sb.removed.extend(paths)


class _Storage:
    def __init__(self, sb):
        self.sb = sb

    def from_(self, name):
        assert name == "artifacts"
        return _Bucket(self.sb)


class FakeSB:
    def __init__(self, no_0166=False):
        self.tables = {"artifacts": [], "generations": [], "jobs": []}
        self.log, self.removed, self.no_0166 = [], [], no_0166
        self.stamp_inserts = False
        self.storage = _Storage(self)

    def table(self, name):
        return _Q(self, name)


@pytest.fixture
def uploads(monkeypatch):
    got: list[tuple[str, str, bytes]] = []

    def _upload(_sb, local, dest):
        got.append((str(local), dest, Path(local).read_bytes()))
        return dest

    monkeypatch.setattr(db, "upload_artifact", _upload)
    return got


def _result(tmp_path, n=2):
    pptx = _pptx(tmp_path / "deck.pptx", [{"title": f"T{i}"} for i in range(1, n + 1)])
    return DP.export_deck_pages(pptx, tmp_path / "work", _fake_renderer())


def _kinds(sb):
    return sorted((r["kind"], r["storage_path"]) for r in sb.tables["artifacts"])


def test_publish_uploads_and_registers(tmp_path, uploads):
    sb = FakeSB()
    sb.tables["artifacts"] = [
        {"generation_id": "g1", "kind": "deck_pptx", "storage_path": "u1/g1/deck.pptx"},
        {"generation_id": "g1", "kind": "slide_png", "storage_path": "u1/g1/deck_pages/slide_01.png"},
        {"generation_id": "g1", "kind": "slide_png", "storage_path": "u1/g1/deck_pages/slide_02.png"},
        {"generation_id": "g1", "kind": "slide_png", "storage_path": "u1/g1/deck_pages/slide_03.png"},
        {"generation_id": "g1", "kind": "deck_pages_json", "storage_path": "u1/g1/deck_pages/pages.json"},
        {"generation_id": "g2", "kind": "slide_png", "storage_path": "u9/g2/deck_pages/slide_01.png"},
    ]
    out = DP.publish_deck_pages(sb, "g1", "u1", "u1/g1/deck.pptx", _result(tmp_path))
    assert [d for _, d, _ in uploads] == ["u1/g1/deck_pages/slide_01.png", "u1/g1/deck_pages/slide_02.png",
                                         "u1/g1/deck_pages/pages.json"]
    g1 = sorted((r["kind"], r["storage_path"]) for r in sb.tables["artifacts"] if r["generation_id"] == "g1")
    assert g1 == [("deck_pages_json", "u1/g1/deck_pages/pages.json"), ("deck_pptx", "u1/g1/deck.pptx"),
                  ("slide_png", "u1/g1/deck_pages/slide_01.png"), ("slide_png", "u1/g1/deck_pages/slide_02.png")]
    assert {"generation_id": "g2", "kind": "slide_png", "storage_path": "u9/g2/deck_pages/slide_01.png"} \
        in sb.tables["artifacts"], "another generation's pages are untouched"
    # the cleanup was scoped: two deletes, each naming generation AND kind
    deletes = [e[2] for e in sb.log if e[0] == "delete"]
    assert deletes == [[("generation_id", "g1"), ("kind", "slide_png")],
                       [("generation_id", "g1"), ("kind", "deck_pages_json")]]
    assert sb.removed == ["u1/g1/deck_pages/slide_03.png"], "the deck got shorter: the stale PNG goes"
    assert out["map_registered"] is True and out["slides"] == 2
    pages = json.loads(uploads[-1][2])
    assert pages["deck"] == ["u1/g1/deck.pptx"] and [s["png"] for s in pages["slides"]] == \
        ["slide_01.png", "slide_02.png"]


def test_publish_is_idempotent(tmp_path, uploads):
    sb = FakeSB()
    sb.tables["artifacts"] = [{"generation_id": "g1", "kind": "deck_pptx", "storage_path": "u1/g1/deck.pptx"}]
    res = _result(tmp_path)
    DP.publish_deck_pages(sb, "g1", "u1", "u1/g1/deck.pptx", res)
    first = _kinds(sb)
    DP.publish_deck_pages(sb, "g1", "u1", "u1/g1/deck.pptx", res)
    assert _kinds(sb) == first and len(first) == 4 and sb.removed == []


def test_without_0166_the_map_is_uploaded_but_not_registered(tmp_path, uploads, caplog):
    sb = FakeSB(no_0166=True)
    sb.tables["artifacts"] = [{"generation_id": "g1", "kind": "deck_pptx", "storage_path": "u1/g1/deck.pptx"}]
    with caplog.at_level(logging.WARNING):
        out = DP.publish_deck_pages(sb, "g1", "u1", "u1/g1/deck.pptx", _result(tmp_path))
    assert out["map_registered"] is False
    assert _kinds(sb) == [("deck_pptx", "u1/g1/deck.pptx"), ("slide_png", "u1/g1/deck_pages/slide_01.png"),
                          ("slide_png", "u1/g1/deck_pages/slide_02.png")]
    assert "u1/g1/deck_pages/pages.json" in [d for _, d, _ in uploads]
    assert "0166 not applied: text map uploaded but not registered" in caplog.text


def test_any_other_insert_error_is_not_swallowed(tmp_path, uploads, monkeypatch):
    sb = FakeSB()
    monkeypatch.setattr(DP, "KIND_MAP", "boom")
    with pytest.raises(RuntimeError, match="connection reset"):
        DP.publish_deck_pages(sb, "g1", "u1", "u1/g1/deck.pptx", _result(tmp_path))


# ── the hook in the deck build ──────────────────────────────────────────


def _deck_build(monkeypatch, tmp_path):
    from tests.test_deck_generation import ANALYSIS, BOOK, CHAPTER, _StubClient, _reply, _worker_env

    sb, uploads_, process = _worker_env(monkeypatch)
    monkeypatch.setenv("DECK_PAGES", "1")

    def run():
        return process._generate_deck(sb, "job-1", "gen-1", BOOK, CHAPTER, ANALYSIS, _StubClient(_reply()),
                                      {"num_slides": 8}, {}, "en", "ltr", tmp_path, "u1/gen-1", "Cells")

    return sb, run


def test_a_failing_export_never_fails_the_deck(monkeypatch, tmp_path, caplog):
    sb, run = _deck_build(monkeypatch, tmp_path)

    def boom(*a, **k):
        raise DP.DeckRenderUnavailable("no slide images: no renderer available")

    monkeypatch.setattr(DP, "export_deck_pages", boom)
    with caplog.at_level(logging.WARNING):
        title = run()
    assert title.endswith("Slide deck")
    assert [(r["kind"], r["storage_path"]) for r in sb.tables["artifacts"]] == [("deck_pptx", "u1/gen-1/deck.pptx")]
    assert "Board pages not made" in caplog.text
    assert sb.tables["jobs"][0]["progress"] == 96


def test_the_export_runs_after_the_deck_row_with_the_build_budget(monkeypatch, tmp_path):
    sb, run = _deck_build(monkeypatch, tmp_path)
    seen = {}

    def export(pptx, work, renderer=None, **kw):
        seen["export"] = (Path(pptx).suffix, Path(pptx).exists(), kw.get("timeout"), kw.get("reuse_qa"))
        seen["rows_then"] = [r["kind"] for r in sb.tables["artifacts"]]
        return {"pngs": [], "map": {}, "renderer": "fake"}

    def publish(_sb, gen, owner, pptx_path, res):
        seen["publish"] = (gen, owner, pptx_path)

    monkeypatch.setattr(DP, "export_deck_pages", export)
    monkeypatch.setattr(DP, "publish_deck_pages", publish)
    run()
    assert seen["export"] == (".pptx", True, DP.BUILD_RENDER_SECONDS, True), "the local deck just built"
    assert seen["rows_then"] == ["deck_pptx"], "the deck is registered before any page work"
    assert seen["publish"] == ("gen-1", "u1", "u1/gen-1/deck.pptx")


def test_deck_pages_off_switch(monkeypatch, tmp_path):
    sb, run = _deck_build(monkeypatch, tmp_path)
    monkeypatch.setenv("DECK_PAGES", "0")
    monkeypatch.setattr(DP, "export_deck_pages", lambda *a, **k: pytest.fail("exported with DECK_PAGES=0"))
    run()


def test_every_deck_row_site_is_followed_by_the_hook():
    import inspect

    from worker import process

    src = inspect.getsource(process._generate_deck)
    assert src.count('db.add_artifact_row(sb, generation_id, "deck_pptx", dest)') == 3
    assert src.count("_deck_pages_after_build(sb, generation_id, base, deck_path, dest, tmp)") == 3


# ── the deck_pages job ──────────────────────────────────────────────────


def _job_sb(kind="deck", decks=("deck.pptx",)):
    sb = FakeSB()
    sb.tables["generations"] = [{"id": "g1", "kind": kind, "owner_id": "u1", "status": "done"}]
    sb.tables["jobs"] = [{"id": "j1", "type": "deck_pages", "status": "processing", "progress": 1,
                          "generation_id": None, "params": {"generation_id": "g1"}}]
    sb.tables["artifacts"] = [{"generation_id": "g1", "kind": "deck_pptx", "storage_path": f"u1/g1/{d}"}
                              for d in decks]
    sb.tables["artifacts"].append({"generation_id": "g1", "kind": "video_mp4", "storage_path": "u1/g1/lesson.mp4"})
    return sb


@pytest.fixture
def downloads(monkeypatch, tmp_path):
    """download_artifact writes a real deck: deck.pptx has 2 slides, part N
    has N+1, each titled by its file."""
    got = []

    def _download(_sb, path, dest):
        got.append(path)
        part = DP.deck_part_of(path)
        n = 2 if part == 1 else part + 1
        _pptx(Path(dest), [{"title": f"{Path(path).stem} s{i}"} for i in range(1, n + 1)])
        return Path(dest)

    monkeypatch.setattr(db, "download_artifact", _download)
    return got


def test_the_job_renders_and_publishes_a_deck(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb()
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer())
    job = sb.tables["jobs"][0]
    assert job["status"] == "done" and job["error"] is None
    assert sb.tables["generations"][0]["status"] == "done", "the deck's own status is never written"
    assert not [e for e in sb.log if e[1] == "generations"]
    assert ("video_mp4", "u1/g1/lesson.mp4") in _kinds(sb) and ("deck_pptx", "u1/g1/deck.pptx") in _kinds(sb)
    assert [k for k, _ in _kinds(sb)].count("slide_png") == 2
    assert downloads == ["u1/g1/deck.pptx"]


def test_the_job_numbers_a_presentations_parts_on(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb("presentation", decks=("deck_part2.pptx", "deck.pptx"))
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer())
    assert sb.tables["jobs"][0]["status"] == "done"
    pages = json.loads(uploads[-1][2])
    assert [(s["n"], s["part"], s["title"]) for s in pages["slides"]] == [
        (1, 1, "deck s1"), (2, 1, "deck s2"), (3, 2, "deck_part2 s1"), (4, 2, "deck_part2 s2"),
        (5, 2, "deck_part2 s3")]
    assert pages["deck"] == ["u1/g1/deck.pptx", "u1/g1/deck_part2.pptx"]
    assert [k for k, _ in _kinds(sb)].count("slide_png") == 5


@pytest.mark.parametrize("kind,decks,msg", [
    ("worksheet", ("deck.pptx",), "not a deck"),
    ("deck", (), "no deck_pptx artifact"),
])
def test_the_job_refuses_what_is_not_a_deck(uploads, downloads, kind, decks, msg):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb(kind, decks)
    before = list(sb.tables["artifacts"])
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer())
    job = sb.tables["jobs"][0]
    assert job["status"] == "error" and msg in job["error"]
    assert sb.tables["artifacts"] == before and uploads == []
    assert sb.tables["generations"][0]["status"] == "done"


def test_a_render_failure_is_the_jobs_error_not_the_generations(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb()
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=lambda p, o, pdf=None: ([], "no renderer available"))
    job = sb.tables["jobs"][0]
    assert job["status"] == "error" and "DeckRenderUnavailable" in job["error"]
    assert sb.tables["generations"][0]["status"] == "done"


# ── repeat presses: two jobs for one deck render it once ────────────────


def _at(minutes: float) -> str:
    """A PostgREST-shaped stamp ``minutes`` from a fixed moment."""
    base = datetime(2026, 10, 29, 10, 0, tzinfo=timezone.utc)
    return db._stamp(base + timedelta(minutes=minutes))


def _pages(sb, created: str, n: int = 2, gen: str = "g1"):
    sb.tables["artifacts"].extend({"generation_id": gen, "kind": "slide_png", "created_at": created,
                                   "storage_path": f"u1/{gen}/deck_pages/slide_{i:02d}.png"}
                                  for i in range(1, n + 1))


def _assert_skipped(sb, before, calls, uploads, downloads, note):
    job = sb.tables["jobs"][0]
    assert job["status"] == "done" and job["error"] is None
    assert job["stage"]["phase"] == "skipped"
    assert job["stage"]["note"].startswith(note), job["stage"]["note"]
    assert calls == [], "a skipped job never calls the renderer"
    assert uploads == [] and downloads == [], "nor downloads the deck or uploads a page"
    assert sb.tables["artifacts"] == before, "the pages already there are left as they are"
    assert sb.tables["generations"][0]["status"] == "done"
    assert not [e for e in sb.log if e[1] == "generations"], "generations.status is never written"


def test_a_render_newer_than_the_press_is_not_redone(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb()
    sb.tables["jobs"][0]["created_at"] = _at(0)          # pressed at 10:00
    _pages(sb, _at(0.5))                                 # another render finished 10:00:30
    before = [dict(r) for r in sb.tables["artifacts"]]
    calls = []
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    _assert_skipped(sb, before, calls, uploads, downloads, "skipped: rendered at ")
    assert _at(0.5) in sb.tables["jobs"][0]["stage"]["note"]


def test_pages_older_than_the_press_are_rendered_again(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb()
    sb.tables["jobs"][0]["created_at"] = _at(0)
    _pages(sb, _at(-60), n=3)                            # last hour's render: the press asks for a new one
    calls = []
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    job = sb.tables["jobs"][0]
    assert job["status"] == "done" and "stage" not in job
    assert len(calls) == 1 and downloads == ["u1/g1/deck.pptx"]
    assert [k for k, _ in _kinds(sb)].count("slide_png") == 2, "the old three replaced by the new two"


def test_another_job_already_rendering_this_deck_wins(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb()
    sb.tables["jobs"][0].update(created_at=_at(0.1), claimed_at=_at(0.2))
    sb.tables["jobs"].append({"id": "j0", "type": "deck_pages", "status": "processing", "generation_id": None,
                              "params": {"generation_id": "g1"}, "created_at": _at(0), "claimed_at": _at(0.15)})
    before = [dict(r) for r in sb.tables["artifacts"]]
    calls = []
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    _assert_skipped(sb, before, calls, uploads, downloads, "skipped: another render in progress")
    assert sb.tables["jobs"][1]["status"] == "processing", "the other job is not touched"


def test_two_simultaneous_claims_never_both_skip(uploads, downloads):
    """j1 was claimed FIRST: the later claim (j2, still processing) must not
    make it yield, or each would yield to the other and the deck gets none."""
    from worker.deck_pages_job import process_deck_pages, skip_reason

    sb = _job_sb()
    sb.tables["jobs"][0].update(created_at=_at(0), claimed_at=_at(0.1))
    sb.tables["jobs"].append({"id": "j2", "type": "deck_pages", "status": "processing", "generation_id": None,
                              "params": {"generation_id": "g1"}, "created_at": _at(0.05),
                              "claimed_at": _at(0.2)})
    assert skip_reason(sb, sb.tables["jobs"][1]) == "skipped: another render in progress"
    calls = []
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    assert sb.tables["jobs"][0]["status"] == "done" and len(calls) == 1


def test_only_the_same_decks_jobs_and_pages_count(uploads, downloads):
    from worker.deck_pages_job import process_deck_pages

    sb = _job_sb()
    sb.tables["jobs"][0].update(created_at=_at(1), claimed_at=_at(1))
    sb.tables["jobs"] += [
        {"id": "jx", "type": "deck_pages", "status": "processing", "generation_id": None,
         "params": {"generation_id": "other"}, "created_at": _at(0), "claimed_at": _at(0)},
        {"id": "jq", "type": "deck_pages", "status": "queued", "generation_id": None,
         "params": {"generation_id": "g1"}, "created_at": _at(0), "claimed_at": None},
    ]
    _pages(sb, _at(5), n=1, gen="other")                 # another deck's newer pages
    calls = []
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    assert sb.tables["jobs"][0]["status"] == "done" and len(calls) == 1


def test_a_failed_check_renders_anyway(uploads, downloads, monkeypatch):
    import worker.deck_pages_job as J

    sb = _job_sb()
    sb.tables["jobs"][0]["created_at"] = _at(0)

    def broken(*a):
        raise RuntimeError("timeout")

    monkeypatch.setattr(J, "_newest_render", broken)
    calls = []
    J.process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    assert sb.tables["jobs"][0]["status"] == "done" and len(calls) == 1


@pytest.mark.parametrize("raw,iso", [
    ("2026-10-29T10:00:00.12345+00:00", "2026-10-29T10:00:00.123450+00:00"),
    ("2026-10-29T10:00:00Z", "2026-10-29T10:00:00+00:00"),
    ("2026-10-29 18:00:00.1+08:00", "2026-10-29T10:00:00.100000+00:00"),
    ("2026-10-29T10:00:00", "2026-10-29T10:00:00+00:00"),
])
def test_postgrest_stamps_are_read(raw, iso):
    from worker.deck_pages_job import _when

    assert _when(raw).isoformat() == iso
    assert _when(None) is None and _when("") is None and _when("not a time") is None


def test_the_builds_own_render_and_a_later_job_coexist(monkeypatch, tmp_path, uploads, downloads):
    """The deck build exports its pages itself (_deck_pages_after_build). A
    press queued BEFORE that render finished is satisfied by it; a press
    queued after it renders again."""
    from worker import process
    from worker.deck_pages_job import process_deck_pages

    monkeypatch.setenv("DECK_PAGES", "1")
    real_export = DP.export_deck_pages
    monkeypatch.setattr(DP, "export_deck_pages",
                        lambda pptx, work, renderer=None, **kw: real_export(pptx, work, renderer or _fake_renderer(),
                                                                          **kw))
    sb = _job_sb()
    sb.stamp_inserts = True
    sb.tables["jobs"][0]["created_at"] = db._stamp(datetime.now(timezone.utc) - timedelta(minutes=1))
    deck = _pptx(tmp_path / "build" / "deck.pptx", [{"title": "A"}, {"title": "B"}, {"title": "C"}])
    process._deck_pages_after_build(sb, "g1", "u1/g1", deck, "u1/g1/deck.pptx", tmp_path / "build")
    assert [k for k, _ in _kinds(sb)].count("slide_png") == 3, "the build published its pages"
    before = [dict(r) for r in sb.tables["artifacts"]]
    n_uploads = len(uploads)
    calls = []
    process_deck_pages(sb, sb.tables["jobs"][0], renderer=_fake_renderer(calls))
    assert sb.tables["jobs"][0]["stage"]["note"].startswith("skipped: rendered at ")
    assert calls == [] and downloads == [] and len(uploads) == n_uploads
    assert sb.tables["artifacts"] == before

    later = {"id": "j9", "type": "deck_pages", "status": "processing", "generation_id": None,
             "params": {"generation_id": "g1"},
             "created_at": db._stamp(datetime.now(timezone.utc) + timedelta(seconds=1))}
    sb.tables["jobs"].append(later)
    process_deck_pages(sb, later, renderer=_fake_renderer(calls))
    assert sb.tables["jobs"][1]["status"] == "done" and "stage" not in sb.tables["jobs"][1]
    assert len(calls) == 1 and downloads == ["u1/g1/deck.pptx"]


# ── dispatch: an observer, on its own lane, never process_generation ────


def test_deck_pages_is_an_observer_in_both_lists():
    import worker.run as run

    assert "deck_pages" in db.OBSERVER_JOB_TYPES
    assert "deck_pages" in run.OBSERVER_JOB_TYPES
    assert "deck_pages" in run._user_builder_exclusions(), "never on a user lane"
    assert "deck_pages" not in run.CATALOGUE_JOB_TYPES, "not catalogue work: its own lane"
    assert db.generation_to_mirror({"type": "deck_pages", "generation_id": "g1"}) is None


@pytest.fixture
def lanes(monkeypatch):
    import worker.run as run

    monkeypatch.setattr(run, "_claim_catalogue_generation", lambda sb: None)
    run._shutdown.clear()
    with run._inflight_lock:
        run._inflight_jobs.clear()
    yield run
    with run._inflight_lock:
        run._inflight_jobs.clear()


def _guard_sb(*jobs):
    from tests.test_observer_job_guard import _fresh

    return _fresh("done", jobs=jobs)


def _dp_job(job_id="job-dp", created="0"):
    from tests.test_observer_job_guard import GEN

    return {"id": job_id, "type": "deck_pages", "status": "queued", "generation_id": GEN,
            "params": {"generation_id": GEN}, "attempts": 0, "created_at": created, "updated_at": "0"}


def test_run_once_dispatches_deck_pages_to_its_handler(lanes, monkeypatch):
    import worker.deck_pages_job as J

    run = lanes
    handled, rebuilt = [], []
    monkeypatch.setattr(run, "process_generation", lambda *a, **k: rebuilt.append(a))
    monkeypatch.setattr(J, "process_deck_pages", lambda sb, job: (handled.append(job["id"]),
                                                                  db.finish_job(sb, job["id"])))
    sb = _guard_sb(_dp_job())
    assert run.run_once(sb) is True
    assert handled == ["job-dp"] and rebuilt == [], "a deck_pages job must never reach process_generation"
    assert sb.tables["jobs"][0]["status"] == "done"
    assert sb.tables["generations"][0]["status"] == "done", "the deck it read is never relabelled"


def test_a_crashing_handler_never_relabels_the_deck(lanes, monkeypatch):
    import worker.deck_pages_job as J

    run = lanes

    def crash(sb, job):
        raise RuntimeError("unexpected")

    monkeypatch.setattr(J, "process_deck_pages", crash)
    monkeypatch.setattr(run, "process_generation", lambda *a, **k: pytest.fail("rebuilt"))
    sb = _guard_sb(_dp_job())
    run.run_once(sb)
    assert sb.tables["jobs"][0]["status"] == "error"
    assert sb.tables["generations"][0]["status"] == "done"


def test_a_teachers_builder_goes_first(lanes, monkeypatch):
    from tests.test_observer_job_guard import _builder

    import worker.deck_pages_job as J

    run = lanes
    order = []
    monkeypatch.setattr(run, "process_generation", lambda sb, job, gen: (order.append(job["type"]),
                                                                       db.finish_job(sb, job["id"], gen)))
    monkeypatch.setattr(J, "process_deck_pages", lambda sb, job: (order.append(job["type"]),
                                                                  db.finish_job(sb, job["id"])))
    sb = _guard_sb(_dp_job(created="0"), {**_builder("job-p", "presentation"), "created_at": "5"})
    run.run_once(sb)
    run.run_once(sb)
    assert order == ["presentation", "deck_pages"], "the older deck_pages job waits for the user's build"


def test_the_lane_holds_at_its_cap(lanes, monkeypatch):
    run = lanes
    monkeypatch.setattr(run, "MAX_DECK_PAGES_JOBS", 1)
    run._inflight_add({"id": "busy", "type": "deck_pages"})
    sb = _guard_sb(_dp_job())
    assert run._claim_deck_pages(sb) is None
    assert sb.tables["jobs"][0]["status"] == "queued"

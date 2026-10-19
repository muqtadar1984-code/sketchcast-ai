"""The Kit Translator, worker side (catalogue.translate, catalogue.translate_check
and the ``script_override`` seam of worker.process — plan revision 3, T0).

Everything CALLS things: the database is tests/catalogue_fakes.FakeSB, the
model is a fake translator that rewrites each string with an Arabic marker
(and, where a test wants it, Arabic-Indic numerals, so numeric EQUIVALENCE
is what is checked, never identical characters), and every heavy render
stage is the recorder tests/test_catalogue_kit.py already uses. No network,
no model, no live Supabase."""

from __future__ import annotations

import copy
import json
import logging

import pytest

from catalogue import translate as T
from catalogue import translate_check as C
from tests.catalogue_fakes import FakeSB
from tests.test_catalogue_kit import (ARTICLE, FIGURES, KIT, PARAMS, _gen, _job, _presentation_fakes, _sb,
                                      _worker_env)

ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def _arabize(s: str) -> str:
    """The fake translation: an Arabic marker in front, the numerals in
    Arabic-Indic digits, the words otherwise kept (so a test can still read
    them)."""
    return "ترجمة " + s.translate(ARABIC_DIGITS)


class FakeTranslator:
    model = "fake-translator"

    def __init__(self, transform=None):
        self.transform = transform or _arabize
        self.calls: list[str] = []
        self.session_usage = {"calls": 0, "input_tokens": 0, "output_tokens": 0, "cost_usd": 0.0}

    def analyze(self, prompt, system="", max_tokens=0, response_schema=None, **kw):
        self.calls.append(prompt)
        self.session_usage["calls"] += 1
        payload = json.loads(prompt[prompt.rindex("INPUT:") + len("INPUT:"):].strip())
        return {"data": {"items": [self.transform(s) for s in payload]}, "usage": {}}


# ── translate_check ──────────────────────────────────────────────────────


class TestNumerals:
    def test_arabic_indic_and_devanagari_digits_are_the_same_values(self):
        assert C.number_values("P = πr + 2r = 41.12 cm") == ["2", "41.12"]
        assert C.number_values("المحيط ٤١٫١٢ سم و ٢") == ["41.12", "2"]
        assert C.number_values("परिमाप ४१.१२") == ["41.12"]
        ok, missing, extra = C.numbers_equivalent("3.14 and 2,500 and 7", "٣٫١٤ و ٢٥٠٠ و ٧")
        assert ok and not missing and not extra

    def test_a_spoken_sign_or_exponent_is_matched_by_magnitude(self):
        # the Quadratic kit (2026-10-17): "negative three" → -٣, "x squared" → س²
        assert C.english_number_words("x squared minus two x minus eight equals zero") == ["2", "2", "8", "0"]
        en = "Subtracting three from both sides gives x equals negative three. Subtracting five gives x equals negative five."
        ar = "بطرح ٣ من الطرفين نحصل على س = -٣. وبطرح ٥ نحصل على س = -٥."
        ok, missing, extra = C.numbers_equivalent(en, ar)
        assert ok, (missing, extra)
        ok, missing, extra = C.numbers_equivalent("x squared minus two x", "س² − ٢س")
        assert ok, (missing, extra)
        # a digit the English neither writes nor speaks is still invented
        ok, missing, extra = C.numbers_equivalent("x equals negative three", "س = -٣ أو س = ٧")
        assert not ok and extra == ["7"]

    def test_a_spoken_line_may_quote_the_board(self):
        lesson = {"examples": [{"problem": "x^2 - 2x = 8",
                                "steps": [{"before": ["x^2 - 2x = 8"], "after": ["x^2 - 2x - 8 = 0"],
                                           "speech": "We subtract eight from both sides so the right side is zero."}],
                                "final_answer": ["x = 4", "x = -2"]}]}
        assert C.board_numbers(lesson) == ["2", "8", "0", "4", "-2"]
        # the translation restates the equation in digits: every digit is on the board
        ok, missing, extra = C.numbers_equivalent(lesson["examples"][0]["steps"][0]["speech"],
                                                  "نطرح ٨ من الطرفين: س² − ٢س − ٨ = ٠.", board=C.board_numbers(lesson))
        assert ok, (missing, extra)
        # a value on no board and in no words is invented
        ok, missing, extra = C.numbers_equivalent("so the right side is zero", "فيصبح الطرف الأيمن ٠ والجذر ٩", board=["0"])
        assert not ok and extra == ["9"]
        # and the whole-lesson check reads the board for every spoken line
        target = {"examples": [{"problem": "x^2 - 2x = 8",
                                "steps": [{"before": ["x^2 - 2x = 8"], "after": ["x^2 - 2x - 8 = 0"],
                                           "speech": "نطرح ٨: س² − ٢س − ٨ = ٠."}],
                                "final_answer": ["x = 4", "x = -2"]}]}
        assert C.check_maths_lesson(lesson, target)["ok"]

    def test_a_refused_script_quotes_the_target_around_each_invented_number(self):
        samples = C.number_samples("نبدأ من س = -٩ ثم نعود إلى -٩ مرة أخرى، والجذر ٥.", ["-9", "5"])
        assert len(samples) == 2 and "-9" in samples[0] and "5" in samples[1]
        assert C.number_samples("لا أرقام هنا", ["7"]) == []
        src = {"segments": [{"segment_id": "s1", "text": "Factorise and set each factor to zero."}]}
        dst = {"segments": [{"segment_id": "s1", "text": "حلّل ثم ساوِ كل عامل بالصفر، فيكون س = -٩."}]}
        rep = C.check_script(src, dst, language="ar", strict=False)
        f = next(b for b in rep["blocking"] if b["what"] == "numbers differ")
        assert f["extra"] == ["-9"] and f["target_samples"] and "-9" in f["target_samples"][0]

    def test_a_recompiled_board_number_is_a_flag_but_a_spoken_number_still_blocks(self):
        # the Quadratic kit (2026-10-17): today's chart labels the vertex (7/2, -9/4);
        # the English render of 2026-09-28 predates the label
        src = {"segments": [{"segment_id": "s1", "text": "Solve x squared minus seven x plus ten equals zero.",
                             "scene": {"elements": [{"type": "text", "text": "x = 2, x = 5"}]}}]}
        dst = {"segments": [{"segment_id": "s1", "text": "حلّ س² − ٧س + ١٠ = ٠.",
                             "scene": {"elements": [{"type": "text", "text": "x = 2, x = 5 رأسه عند (7/2, -9/4)"}]}}]}
        rep = C.check_script(src, dst, language="ar", strict=False)
        assert rep["ok"], rep["blocking"]
        f = next(x for x in rep["flags"] if x["what"].startswith("board numbers differ"))
        assert "-9" in f["extra"] and f["target_samples"]
        # the same new value SPOKEN in a FRAMING segment (the translator's words) is still invented
        dst2 = {"segments": [{"segment_id": "s1", "text": "حلّ س² − ٧س + ١٠ = ٠، والرأس عند -٩/٤.",
                              "scene": {"elements": [{"type": "text", "text": "x = 2, x = 5"}]}}]}
        rep2 = C.check_script(src, dst2, language="ar", strict=False)
        assert not rep2["ok"] and rep2["blocking"][0]["extra"] == ["-9", "4"]
        # … but spoken by a COMPILED board segment (scene.compiled, the compiler's own speech
        # from the verified record) it is the compiler's: a flag, not a block
        src3 = {"segments": [{"segment_id": "s9", "text": "The curve cuts the x-axis at x equals two and five.",
                              "scene": {"compiled": True, "elements": [{"type": "text", "text": "x = 2, x = 5"}]}}]}
        dst3 = {"segments": [{"segment_id": "s9", "text": "يقطع المنحنى المحور السيني عند س = ٢ و ٥. رأسه عند (7/2, -9/4).",
                              "scene": {"compiled": True, "elements": [{"type": "text", "text": "x = 2, x = 5 (7/2, -9/4)"}]}}]}
        rep3 = C.check_script(src3, dst3, language="ar", strict=False)
        assert rep3["ok"], rep3["blocking"]
        assert any(f["what"].startswith("board numbers differ") for f in rep3["flags"])

    def test_a_number_before_a_full_stop_is_a_number(self):
        assert C.number_values("So x equals 6.") == ["6"]
        assert C.number_values("إذن س تساوي ٦.") == ["6"]
        assert C.number_values("the ratio is 2.5. Then 3") == ["2.5", "3"]

    def test_a_french_decimal_comma_is_a_decimal_mark(self):
        assert C.number_values("1,5 cm") == ["1.5"]
        assert C.number_values("41,120 people") == ["41120"]

    def test_a_changed_number_is_caught_either_way(self):
        ok, missing, extra = C.numbers_equivalent("x = 5", "س = ٦")
        assert not ok and missing == ["5"] and extra == ["6"]

    def test_english_number_words_parse_to_values(self):
        # the Angles kit's spoken lines (2026-10-17): the voice reads words
        assert C.english_number_words("they must sum to three hundred and sixty degrees") == ["360"]
        assert C.english_number_words("angle AOC is one hundred and thirty-five degrees, and BOD is three z") == ["135", "3"]
        assert C.english_number_words("ninety-five plus one hundred and thirty-five plus two x equals three hundred and sixty") == ["95", "135", "2", "360"]
        assert C.english_number_words("two x minus ten equals one hundred and ten") == ["2", "10", "110"]
        assert C.english_number_words("x equals forty-five; two point five metres; twelve thousand and six") == ["45", "2.5", "12006"]
        assert C.english_number_words("a hundred ways, and hundreds of them") == []
        assert C.english_number_words("one of the angles and the point O") == ["1"]

    def test_a_number_the_english_spells_out_may_be_a_digit_in_the_translation(self):
        en = "Adding the given angles gives two hundred and seventy. Subtracting two hundred and seventy from three hundred and sixty, we find y equals ninety."
        ar = "بجمع الزاويتين المعطاتين نحصل على ٢٧٠. وبطرح ٢٧٠ من ٣٦٠ نجد أن y تساوي ٩٠."
        ok, missing, extra = C.numbers_equivalent(en, ar)
        assert ok and not missing and not extra
        # … but it is never REQUIRED there: "one of the angles" need not become a digit
        ok, missing, extra = C.numbers_equivalent("one of the angles is 40 degrees", "إحدى الزوايا تساوي ٤٠ درجة")
        assert ok and not missing and not extra
        # and a digit the English never spoke is still an invented number
        ok, missing, extra = C.numbers_equivalent("three z equals one hundred and thirty-five", "٣ ع = ١٣٥ أي ع = ٤٥")
        assert not ok and extra == ["45"] and not missing
        # a spelt number absorbs ONE digit, not two
        ok, missing, extra = C.numbers_equivalent("two x", "٢ س و ٢ ص")
        assert not ok and extra == ["2"]

    def test_curriculum_codes_are_invariant_tokens(self):
        assert C.code_tokens("covers 7Gg.12 and cbse:7:GEO:01") == ["7Gg.12", "cbse:7:GEO:01"]


class TestMathsSpans:
    def test_a_hyphenated_word_and_a_bare_dash_are_not_maths(self):
        assert C.maths_spans("Match each circle-related term to its correct description.") == []
        assert C.maths_spans("Label the x-axis and the well-known centre.") == []
        ok, lost = C.prose_maths_equivalent("Match each circle-related term to its correct description.", "طابق كل مصطلح متعلق بالدائرة مع وصفه الصحيح.")
        assert ok and lost == []

    def test_a_number_hyphenated_to_a_word_is_an_adjective_not_a_subtraction(self):
        """Trigonometric Ratios, 2026-10-19: five bank items naming an angle
        were refused in Arabic and in Malay — "the 90-degree angle" read as
        the maths "90-", and "sudut 90 darjah" lost it."""
        stem = "In a right-angled triangle, which side is always located directly opposite the 90-degree angle?"
        assert C.maths_spans(stem) == ["90"]
        ok, lost = C.prose_maths_equivalent(stem, "Dalam segi tiga bersudut tegak, sisi manakah yang sentiasa bertentangan dengan sudut 90 darjah?")
        assert ok and lost == []
        assert C.maths_spans("Find the side opposite the 38-degree angle.") == ["38"]
        assert C.maths_spans("the 3-step method") == ["3"]

    def test_a_number_before_a_function_or_a_variable_keeps_its_minus(self):
        assert C.maths_spans("simplify 1-cos x") == ["1-"] or "1-" in C.maths_spans("simplify 1-cos x")[0]
        assert "2-sin" not in C.maths_spans("2-sin θ")[0] or C.maths_spans("2-sin θ")[0].startswith("2-")
        assert C.maths_spans("Solve 2-x = 5") == ["2-x = 5"]
        assert C.maths_spans("10-4 = 6") == ["10-4 = 6"]

    def test_real_maths_still_counts(self):
        assert C.maths_spans("Use π = 3.14 here") == ["π = 3.14"]
        assert C.maths_spans("Solve 2x + 3 = 11 for x.") == ["2x + 3 = 11", "x"] or "2x + 3 = 11" in C.maths_spans("Solve 2x + 3 = 11 for x.")
        assert C.maths_spans("the value of x - 2 when x = 5") == ["x - 2", "x = 5"]
        assert "x-2" in C.maths_spans("find x-2 if x = 5")[0]
        assert C.maths_spans("A circle has a radius of 7 cm.") == ["7"]


class TestEnglishLeftover:
    def test_latin_letters_in_an_arabic_target_are_flagged_but_protected_maths_is_not(self):
        assert C.english_leftover("The perimeter of the semicircle is", "ar") is not None
        assert C.english_leftover("محيط نصف الدائرة هو P = πr + 2r", "ar", protected=["P = πr + 2r"]) is None
        assert C.english_leftover("محيط نصف الدائرة", "ar") is None

    def test_english_function_words_in_a_latin_target_are_flagged(self):
        assert C.english_leftover("The cell is the smallest unit of life and this is that from which all are made", "fr")
        assert C.english_leftover("La cellule est la plus petite unité du vivant, et tout organisme en est fait.", "fr") is None


def _seg(sid, text, *, who=("teacher", "student"), points=("one", "two"), expr=None, label="Nucleus"):
    lines = [{"who": w, "line": f"{text} {w}"} for w in who]
    scene = {"elements": [{"id": "t1", "type": "text", "text": label, "at": [0, 0]}]}
    if expr:
        scene["elements"].append({"id": "m1", "type": "math", "expr": expr, "at": [0, 1]})
    return {"segment_id": sid, "type": "explore", "text": " ".join(d["line"] for d in lines),
            "elevenlabs_text": text, "slide_heading": f"Heading {text}", "slide_points": list(points),
            "dialogue": lines, "scene": scene, "estimated_duration_seconds": 20}


SRC = {"segments": [_seg("s001", "The cell has 2 parts and 41.12 cm", expr="P = πr + 2r"),
                    _seg("s002", "Nucleus holds DNA")]}


def _translated(src, transform=_arabize):
    dst = copy.deepcopy(src)
    for seg in dst["segments"]:
        for d in seg["dialogue"]:
            d["line"] = transform(d["line"])
        seg["text"] = " ".join(d["line"] for d in seg["dialogue"])
        seg["slide_heading"] = transform(seg["slide_heading"])
        seg["slide_points"] = [transform(p) for p in seg["slide_points"]]
        for e in seg["scene"]["elements"]:
            if e["type"] == "text":
                e["text"] = transform(e["text"])
    return dst


class TestCheckScript:
    def test_a_faithful_translation_passes_with_only_the_leftover_flag(self):
        r = C.check_script(SRC, _translated(SRC), language="ar", glossary=[("nucleus", "النواة")])
        assert r["ok"] and r["blocking"] == []
        whats = {f["what"] for f in r["flags"]}
        # the fake keeps the English words, so the leftover and the glossary
        # flags fire — flags, not blocks (D3: the reviewer decides)
        assert "still reads as English" in whats and "glossary term not used" in whats

    def test_a_dropped_speaker_turn_blocks(self):
        dst = _translated(SRC)
        dst["segments"][1]["dialogue"].pop()
        r = C.check_script(SRC, dst, language="ar")
        assert not r["ok"] and any(f["what"] == "speaker turns differ" and f["where"] == "s002" for f in r["blocking"])

    def test_a_changed_number_blocks(self):
        dst = _translated(SRC, lambda s: _arabize(s.replace("41.12", "41.21")))
        r = C.check_script(SRC, dst, language="ar")
        assert any(f["what"] == "numbers differ" and "41.12" in f["missing"] for f in r["blocking"])

    def test_an_altered_math_expression_blocks(self):
        dst = _translated(SRC)
        dst["segments"][0]["scene"]["elements"][1]["expr"] = "P = 2πr"
        r = C.check_script(SRC, dst, language="ar")
        assert any(f["what"] == "protected expression altered" for f in r["blocking"])

    def test_reordered_segments_block_in_strict_mode_only(self):
        dst = _translated(SRC)
        dst["segments"].reverse()
        assert not C.check_script(SRC, dst, language="ar")["ok"]
        assert C.check_script(SRC, dst, language="ar", strict=False)["ok"]

    def test_the_lenient_mode_flags_a_retypeset_board_expression_instead_of_blocking(self):
        # the Circle pilot (2026-10-16): the Arabic board was recompiled by a
        # newer compiler and typeset the same maths a little differently —
        # the record itself was intact, so this is the reviewer's to judge
        dst = _translated(SRC)
        dst["segments"][0]["scene"]["elements"][1]["expr"] = "P = π·r + 2·r"
        strict = C.check_script(SRC, dst, language="ar")
        assert any(f["what"] == "protected expression altered" for f in strict["blocking"])
        lenient = C.check_script(SRC, dst, language="ar", strict=False)
        assert lenient["ok"]
        assert any(f["what"].startswith("board expressions typeset differently") and f["examples"] == ["P = πr + 2r"]
                   for f in lenient["flags"])
        # a caller-protected expression (the record's own problem) still blocks
        refused = C.check_script(SRC, dst, language="ar", strict=False, protected={"s001": ["P = πr + 2r"]})
        assert not refused["ok"]

    def test_the_lenient_mode_still_catches_a_lost_number(self):
        dst = _translated(SRC, lambda s: _arabize(s.replace("41.12", "")))
        r = C.check_script(SRC, dst, language="ar", strict=False)
        assert any(f["what"] == "numbers differ" for f in r["blocking"])


class TestQuestionProse:
    def test_maths_spans_keep_the_maths_and_free_the_words(self):
        assert C.maths_spans("Find the area of a circle with a radius of 6 cm. (Use π = 3.14)") == ["6", "π = 3.14"]
        assert C.maths_spans("Solve 2x + 3 = 11 for x.") == ["2x + 3 = 11"]
        assert C.maths_spans("A semicircle has a radius of 8 cm. Find its total perimeter.") == ["8"]
        assert C.maths_spans("Identify which of these shapes represents a circle") == []

    def test_a_translated_question_passes_when_its_maths_survives_and_blocks_when_not(self):
        src = "Find the area of a circle with a radius of 6 cm. (Use π = 3.14)"
        ok, lost = C.prose_maths_equivalent(src, "أوجد مساحة دائرة نصف قطرها 6 سم. (استخدم π = 3.14)")
        assert ok and lost == []
        ok, lost = C.prose_maths_equivalent(src, "أوجد مساحة دائرة نصف قطرها ٦ سم. (استخدم π = ٣٫١٤)")
        assert ok, "Arabic-Indic numerals are the same values"
        ok, lost = C.prose_maths_equivalent(src, "أوجد مساحة دائرة نصف قطرها 6 سم. (استخدم باي = 3.14)")
        assert not ok and lost == ["π = 3.14"], "the symbol was spelled out — the board would lose the expression"
        ok, lost = C.prose_maths_equivalent("Solve 2x + 3 = 11 for x.", "حل 2x + 3 = 12 بالنسبة إلى x.")
        assert not ok and "2x + 3 = 11" in lost

    def test_check_maths_lesson_judges_the_question_by_its_spans(self):
        src = {"examples": [{"label": "Example 1", "problem": "Find the area of a circle with a radius of 6 cm. (Use π = 3.14)",
                             "steps": [], "final_answer": ["A = 113.04"]}]}
        good = C.check_maths_lesson(src, {"examples": [{**src["examples"][0], "label": "مثال 1",
                                                        "problem": "أوجد مساحة دائرة نصف قطرها 6 سم. (استخدم π = 3.14)"}]})
        assert good["ok"], good
        bad = C.check_maths_lesson(src, {"examples": [{**src["examples"][0], "problem": "أوجد مساحة دائرة نصف قطرها 7 سم."}]})
        assert any(f["what"] == "maths in the question altered" for f in bad["blocking"])


class TestCheckMathsLesson:
    LESSON = {"topic": "Linear equations", "hook": [{"who": "teacher", "line": "Double it and add 5 gives 21."}],
              "examples": [{"label": "Example 1", "problem": "2x + 3 = 11", "steps": [
                  {"operation": "subtract 3", "before": ["2x + 3 = 11"], "after": ["2x = 8"],
                   "speech": "Take 3 from both sides: 2x equals 8."}], "final_answer": ["x = 4"]}]}

    def test_the_maths_is_untouched_and_the_speech_equivalent(self):
        dst = T.apply_strings(self.LESSON, [(p, _arabize(t)) for p, t in T.lesson_text_items(self.LESSON)])
        r = C.check_maths_lesson(self.LESSON, dst)
        assert r["ok"], r
        assert dst["examples"][0]["steps"][0]["after"] == ["2x = 8"]
        assert dst["examples"][0]["steps"][0]["speech"].startswith("ترجمة")

    def test_an_altered_after_line_or_a_lost_value_blocks(self):
        dst = copy.deepcopy(self.LESSON)
        dst["examples"][0]["steps"][0]["after"] = ["2x = 9"]
        assert any(f["what"] == "protected field altered" for f in C.check_maths_lesson(self.LESSON, dst)["blocking"])
        dst = copy.deepcopy(self.LESSON)
        dst["hook"][0]["line"] = "ضعفه زائد خمسة"
        assert any(f["what"] == "numbers differ" for f in C.check_maths_lesson(self.LESSON, dst)["blocking"])


class TestItems:
    SRC = {"id": "q-1", "item_type": "mcq", "answer_mode": "objective", "marks": 1, "difficulty": 2, "cognitive_level": "recall",
           "objective_ref": "o1", "stem": "A circle has a radius of 7 cm. Its diameter is?",
           "options": [{"key": "A", "text": "14 cm"}, {"key": "B", "text": "7 cm"}, {"key": "C", "text": "3.5 cm"}, {"key": "D", "text": "22 cm"}],
           "answer": {"key": "A"}, "marking_scheme": [], "distractor_rationale": {"B": {"why_wrong": "that is the radius"}},
           "explanation": "The diameter is twice the radius."}

    def test_item_text_items_take_the_words_and_leave_the_keys(self):
        paths = [p for p, _ in T.item_text_items(self.SRC)]
        assert "stem" in paths and "options[0].text" in paths and "distractor_rationale.B.why_wrong" in paths and "explanation" in paths
        assert not any(p.endswith(".key") for p in paths)
        dst = T.translated_item(self.SRC, [(p, _arabize(t)) for p, t in T.item_text_items(self.SRC)])
        assert dst["options"][0]["key"] == "A" and dst["options"][0]["text"].startswith("ترجمة") and dst["answer"] == {"key": "A"}
        assert C.check_item(self.SRC, dst)["ok"]

    def test_check_item_blocks_a_changed_number_a_changed_key_and_collapsed_options(self):
        dst = T.translated_item(self.SRC, [(p, _arabize(t)) for p, t in T.item_text_items(self.SRC)])
        dst["options"][0]["text"] = "ترجمة 15 cm"
        assert any(f["what"] == "numbers differ" for f in C.check_item(self.SRC, dst)["blocking"])
        dst = T.translated_item(self.SRC, [(p, _arabize(t)) for p, t in T.item_text_items(self.SRC)])
        dst["answer"] = {"key": "B"}
        assert any(f["what"] == "answer key changed" for f in C.check_item(self.SRC, dst)["blocking"])
        dst = T.translated_item(self.SRC, [(p, _arabize(t)) for p, t in T.item_text_items(self.SRC)])
        dst["options"][1]["text"] = dst["options"][0]["text"]
        r = C.check_item(self.SRC, dst)
        assert any(f["what"] == "translated options are no longer distinct" for f in r["blocking"])

    def test_a_translated_items_hash_tells_arabic_stems_apart_and_is_scoped_to_the_kit(self):
        # catalogue.key.canonical_key folds non-Latin text to "" — the stock
        # content_hash would give every Arabic item of the bank one hash
        a = {"item_type": "mcq", "stem": "ما هو نصف قطر الدائرة؟"}
        b = {"item_type": "mcq", "stem": "ما هي مساحة الدائرة؟"}
        assert T._item_hash(a, "kit-1") != T._item_hash(b, "kit-1")
        assert T._item_hash(a, "kit-1") == T._item_hash({"item_type": "mcq", "stem": "  ما هو نصف قطر الدائرة  "}, "kit-1")
        assert T._item_hash(a, "kit-1") != T._item_hash(a, "kit-2"), "a re-translation for a new kit must not collide with the retired kit's items"

    def test_a_match_items_pairs_read_the_same_in_options_and_answer(self):
        pairs = [{"left": "Radius", "right": "centre to the boundary"}, {"left": "Chord", "right": "two boundary points"}]
        src = {"id": "m-1", "item_type": "match", "answer_mode": "objective", "marks": 2, "difficulty": 1,
               "cognitive_level": "recall", "objective_ref": "o1", "stem": "Match each term to its description.",
               "options": {"pairs": [dict(p) for p in pairs]}, "answer": {"pairs": [dict(p) for p in pairs]},
               "marking_scheme": [], "distractor_rationale": None, "explanation": None}
        seen: dict[str, int] = {}

        def numbered(s: str) -> str:            # a different rendering every time it is asked (no digits: those are checked)
            seen[s] = seen.get(s, 0) + 1
            return f"ترجمة {s} {'ـ' * seen[s]}"

        c = FakeTranslator(numbered)
        words = T.item_text_items(src)
        got = T.translate_strings(c, [w for _, w in words], "ar", context=T.ITEM_RULE)
        dst = T.translated_item(src, list(zip([p for p, _ in words], got)))
        assert dst["options"]["pairs"] == dst["answer"]["pairs"], "the same term must read the same twice"
        assert all(n == 1 for n in seen.values()), "each distinct string is asked once"
        assert C.check_item(src, dst)["ok"]

    def test_a_refused_item_is_translated_once_more_with_the_reason_and_a_second_refusal_drops_it(self):
        sb = _job_sb()
        sb.tables["topic_questions"] = [_item(1), _item(2), _item(3)]
        sb.tables["topic_kits"][0]["doc_generation_ids"] = {"deck": "gen-deck"}   # no bank piece: the floor does not bind
        asked: list[str] = []

        def flaky(s: str) -> str:
            # item 1's stem gains a number the FIRST time it is asked (the
            # retry gets it right); item 2's stem gains one EVERY time
            asked.append(s)
            if s.startswith("Which part of a cell is the nucleus") and asked.count(s) == 1:
                return _arabize(s) + " 42"
            if s.startswith("The membrane of a cell"):
                return _arabize(s) + " 7"
            return _arabize(s)

        c = FakeTranslator(flaky)
        kit = {"id": "kit-ar", "topic_id": "t-cell"}
        mapping, summary = T.bank_stage(sb, c, {"kit": sb.tables["topic_kits"][0]}, kit, "ar", [], "art-ar")
        assert len(c.calls) == 2, "one call for the bank, one retry call"
        assert "A previous translation of these strings was refused" in c.calls[1] and "numbers differ" in c.calls[1]
        assert "q-1" in mapping, "the retry fixed item 1"
        assert summary["retried"] >= 1 and summary["staged"] == len(mapping)
        if "q-2" not in mapping:
            d = next(d for d in summary["dropped"] if d["source_item_id"] == "q-2")
            assert any(f["what"] == "numbers differ" for f in d["findings"])
            flags = T.bank_flags(summary)
            assert flags and flags[0]["what"] == "bank item not staged" and "numbers differ" in flags[0]["detail"]

    def test_a_translation_that_fails_before_any_piece_is_removed_with_its_items_and_article(self):
        sb = FakeSB()
        sb.tables["topic_kits"] = [{"id": "kit-ar", "status": "generating", "translation_status": "staged", "language": "ar",
                                    "article_id": "art-ar", "presentation_generation_id": None, "doc_generation_ids": {}},
                                   {"id": "kit-other", "status": "approved", "translation_status": "approved", "language": "ar",
                                    "article_id": "art-other", "presentation_generation_id": "g", "doc_generation_ids": {}}]
        sb.tables["topic_articles"] = [{"id": "art-ar", "status": "staged", "language": "ar"},
                                       {"id": "art-other", "status": "approved", "language": "ar"}]
        sb.tables["topic_questions"] = [
            {"id": "a", "translation_kit_id": "kit-ar", "status": "staged"},
            {"id": "b", "translation_kit_id": "kit-ar", "status": "retired"},
            {"id": "c", "translation_kit_id": "kit-other", "status": "approved"},
        ]
        order: list[str] = []
        real_table = sb.table

        def spy(name):
            order.append(name)
            return real_table(name)

        sb.table = spy
        T._fail_translation(sb, {"id": "kit-ar"}, "translation failed: x", report={"ok": False})
        sb.table = real_table
        deletes = [n for n in order if n in ("topic_kits", "topic_articles")]
        assert deletes.index("topic_kits") < deletes.index("topic_articles") or deletes[-2:] == ["topic_kits", "topic_articles"],             "the kit (which references the article) is deleted BEFORE the article — the live FK refused the other order"
        assert [k["id"] for k in sb.tables["topic_kits"]] == ["kit-other"]
        assert [a["id"] for a in sb.tables["topic_articles"]] == ["art-other"]
        assert [q["id"] for q in sb.tables["topic_questions"]] == ["c"]

    def test_a_translation_that_fails_with_pieces_is_kept_as_failed_for_the_retry(self):
        sb = FakeSB()
        sb.tables["topic_kits"] = [{"id": "kit-ar", "status": "generating", "translation_status": "staged", "language": "ar",
                                    "article_id": "art-ar", "presentation_generation_id": "gen-p", "doc_generation_ids": {"deck": "gen-d"}}]
        sb.tables["topic_articles"] = [{"id": "art-ar", "status": "staged", "language": "ar"}]
        sb.tables["topic_questions"] = [{"id": "a", "translation_kit_id": "kit-ar", "status": "staged"}]
        T._fail_translation(sb, {"id": "kit-ar"}, "translation failed: x")
        k = sb.tables["topic_kits"][0]
        assert k["status"] == "failed" and k["translation_status"] == "failed"
        assert sb.tables["topic_questions"][0]["status"] == "retired" and sb.tables["topic_articles"]

    def test_fidelity_scores_a_faithful_back_translation_high_and_flags_a_stray_one(self):
        rep = C.fidelity_report([("s1", "The cell is the smallest unit of life", "The cell is the smallest unit of life."),
                                 ("s2", "Plants have a cell wall and chloroplasts", "The weather was fine today")])
        assert rep["ok"] and rep["samples"] == 2 and [f["where"] for f in rep["flags"]] == ["s2"]

    def test_the_composer_accepts_staged_items_of_the_kit_only(self):
        from catalogue import composer as CO
        sb = FakeSB()
        sb.tables["topic_questions"] = [
            {"id": "a", "topic_id": "t", "language": "ar", "status": "approved"},
            {"id": "b", "topic_id": "t", "language": "ar", "status": "staged", "translation_kit_id": "kit-ar"},
            {"id": "c", "topic_id": "t", "language": "ar", "status": "staged", "translation_kit_id": "kit-other"},
        ]
        assert [r["id"] for r in CO.load_items_by_id(sb, ["a", "b"], topic_id="t", language="ar", staged_for="kit-ar")] == ["a", "b"]
        with pytest.raises(RuntimeError, match="no longer approved"):
            CO.load_items_by_id(sb, ["a", "b"], topic_id="t", language="ar")
        with pytest.raises(RuntimeError, match="no longer approved"):
            CO.load_items_by_id(sb, ["c"], topic_id="t", language="ar", staged_for="kit-ar")
        assert {r["id"] for r in CO.load_bank(sb, "t", "ar", staged_for="kit-ar")} == {"a", "b"}
        assert {r["id"] for r in CO.load_bank(sb, "t", "ar")} == {"a"}


# ── the worksheet (T1.5): translated from its verified questions ─────────

PNG_1X1 = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAA"
           "AABJRU5ErkJggg==")
WS_PATH = "sys/gen-w/worksheet_questions.json"
# the English maths worksheet's questions_json (docgen.questions.write_worksheet):
# a fact fill-in, a true/false, a matching item, two ladder problems with their
# working (one with accepted spellings), a figure question with its picture
WS_QUESTIONS = {
    "title": "Worksheet — Circles", "instructions": "Answer every question.",
    "questions": [
        {"id": "q1", "type": "fill_blank", "prompt": "The longest chord of a circle is called the ____.",
         "answer": "diameter", "marks": 1},
        {"id": "q2", "type": "true_false", "prompt": "A radius of 7 cm gives a diameter of 14 cm.", "answer": True, "marks": 1},
        {"id": "q3", "type": "match", "prompt": "Match each term to its meaning.", "marks": 2,
         "pairs": [{"left": "radius", "right": "centre to the edge"},
                   {"left": "chord", "right": "a line joining two points on the circle"}]},
        {"id": "q4", "type": "fill_blank", "prompt": "Solve: 3x + 5 = 20", "answer": "x = 5", "accept": ["5"], "marks": 1,
         "working": ["3x = 15   (subtract 5 from both sides)", "x = 5", "Check: 3(5) + 5 = 20", "Answer: x = 5"]},
        {"id": "q5", "type": "fill_blank", "prompt": "Find the circumference of a circle with a radius of 6 cm. (Use π = 3.14)",
         "answer": "37.68 cm", "accept": ["37.68"], "marks": 1, "working": ["C = 2 × 3.14 × 6", "Answer: 37.68 cm"]},
        {"id": "q6", "type": "fill_blank", "prompt": "Find x in the diagram. Give the number only.", "answer": "40",
         "marks": 2, "image": PNG_1X1},
    ],
}
WS_WORDS = {text for q in WS_QUESTIONS["questions"] for _, text in T.question_text_items(q)}


def _ws_job_sb():
    """The job fixture with a MATHS worksheet: the source worksheet generation
    carries a questions_json artifact (the verifier's English questions)."""
    sb = _job_sb()
    sb.tables["artifacts"].append({"id": "a-q", "generation_id": KIT["doc_generation_ids"]["worksheet"],
                                   "kind": "questions_json", "storage_path": WS_PATH})
    sb.files[("artifacts", WS_PATH)] = json.dumps(WS_QUESTIONS, ensure_ascii=False).encode("utf-8")
    return sb


def _ws_run(sb, monkeypatch, tmp_path, client):
    uploads = _uploads(monkeypatch)
    monkeypatch.setattr("spike.scene_engine.raster_assets.cache_dir_for", lambda key, cache_dir=None: tmp_path / "nocache" / key)
    res = T.run_translate_job(sb, sb.tables["jobs"][0], client=client)
    assert res is not None, sb.tables["jobs"][0].get("error")
    k = next(k for k in sb.tables["topic_kits"] if k["language"] == "ar")
    staged = [body for dest, body in uploads if dest.endswith("/translated/worksheet_questions.json")]
    return k, staged, uploads


class TestBankBatching:
    """The bank stage's calls run side by side in whole-item batches; a batch
    whose reply cannot be used falls back to one call per item for that batch
    only; the kit-scoped twin rule and the "bank item not staged" flags are
    exactly what they were."""

    def _sb(self, n=12):
        sb = _job_sb()
        sb.tables["topic_questions"] = [_item(i) for i in range(1, n + 1)]
        sb.tables["topic_kits"][0]["doc_generation_ids"] = {"deck": "gen-deck"}   # no bank piece: the floor does not bind
        return sb

    def test_batch_items_never_splits_an_item_and_counts_a_shared_string_once(self):
        items = [_item(i) for i in range(1, 13)]
        batches = T.batch_items(items, 10)
        assert [it["id"] for b in batches for it in b] == [it["id"] for it in items], "every item, once, in order"
        for b in batches[:-1]:
            distinct = {t for it in b for _, t in T.item_text_items(it)}
            assert len(distinct) <= 10, "a batch holds at most the chunk's distinct strings"
        assert len(batches) >= 3, "the shared option texts and explanation are counted once per batch, so a batch holds several items"
        huge = {**_item(1), "stem": "x", "options": [{"key": k, "text": f"option {k}"} for k in "ABCDEFGHIJKLMNOP"]}
        assert T.batch_items([huge, _item(2)], 5) == [[huge], [_item(2)]], "an item wider than the chunk is a batch of its own"
        assert T.batch_items([], 10) == []

    def test_the_knobs_are_read_from_the_environment_and_clamped(self, monkeypatch):
        monkeypatch.delenv(T.BANK_CHUNK_ENV, raising=False)
        monkeypatch.delenv(T.BANK_WORKERS_ENV, raising=False)
        assert T.bank_chunk() == T.BANK_CHUNK_DEFAULT and T.bank_workers() == T.BANK_WORKERS_DEFAULT
        monkeypatch.setenv(T.BANK_CHUNK_ENV, "1000")
        monkeypatch.setenv(T.BANK_WORKERS_ENV, "0")
        assert T.bank_chunk() == T.CHUNK and T.bank_workers() == 1
        monkeypatch.setenv(T.BANK_CHUNK_ENV, "abc")
        monkeypatch.setenv(T.BANK_WORKERS_ENV, "99")
        assert T.bank_chunk() == T.BANK_CHUNK_DEFAULT and T.bank_workers() == T.BANK_WORKERS_MAX

    def test_the_bank_is_translated_in_concurrent_batches_and_the_tally_is_recorded(self, monkeypatch):
        import threading
        monkeypatch.setenv(T.BANK_CHUNK_ENV, "10")
        monkeypatch.setenv(T.BANK_WORKERS_ENV, "3")
        sb = self._sb()
        threads: list[str] = []

        class Threaded(FakeTranslator):
            def analyze(self, prompt, **k):
                threads.append(threading.current_thread().name)
                return super().analyze(prompt, **k)

        c = Threaded()
        kit = {"id": "kit-ar", "topic_id": "t-cell"}
        expected = len(T.batch_items([_item(i) for i in range(1, 13)], 10))   # before the run stages 12 more rows
        mapping, summary = T.bank_stage(sb, c, {"kit": sb.tables["topic_kits"][0]}, kit, "ar", [], "art-ar")
        assert expected >= 3 and len(mapping) == 12 and summary["staged"] == 12 and summary["dropped"] == []
        assert len(c.calls) == expected, "one call per batch, no retry"
        assert summary["batching"] == {"chunk": 10, "workers": 3, "batches": expected, "batch_calls": expected,
                                       "fallback_items": 0, "fallback_failed": 0, "model_calls": expected}
        assert any(t.startswith("bank-translate") for t in threads), "the batches ran on the pool"
        staged = [r for r in sb.tables["topic_questions"] if r.get("translation_kit_id") == "kit-ar"]
        assert len(staged) == 12 and all(r["status"] == "staged" and r["stem"].startswith("ترجمة") for r in staged)
        assert len({r["content_hash"] for r in staged}) == 12, "the kit-scoped hash tells every item apart"

    def test_a_batch_whose_reply_is_unusable_falls_back_to_one_call_per_item_for_that_batch_only(self, monkeypatch):
        monkeypatch.setenv(T.BANK_CHUNK_ENV, "10")
        monkeypatch.setenv(T.BANK_WORKERS_ENV, "2")
        sb = self._sb()
        bad_stem = _item(6)["stem"]                     # "The cytoplasm of a cell is ____."
        batches = T.batch_items(sb.tables["topic_questions"], 10)
        bad_batch = next(b for b in batches if any(it["id"] == "q-6" for it in b))
        assert len(bad_batch) >= 2, "the failing batch has company"

        class Garbled(FakeTranslator):
            def analyze(self, prompt, **k):
                # any reply that would carry item 6 comes back as prose the
                # client could not read — the whole batch, and item 6 alone
                if bad_stem in prompt:
                    self.calls.append(prompt)
                    self.session_usage["calls"] += 1
                    return {"data": {"raw_text": "عذراً، لا يمكنني"}, "usage": {}}
                return super().analyze(prompt, **k)

        c = Garbled()
        kit = {"id": "kit-ar", "topic_id": "t-cell"}
        mapping, summary = T.bank_stage(sb, c, {"kit": sb.tables["topic_kits"][0]}, kit, "ar", [], "art-ar")
        # every item but the garbled one is staged — the rest of its batch by its own call
        assert set(mapping) == {f"q-{i}" for i in range(1, 13)} - {"q-6"}
        for it in bad_batch:
            if it["id"] != "q-6":
                own = [p for p in c.calls if it["stem"] in p and bad_stem not in p]
                assert len(own) == 1, f"{it['id']} was translated once on its own after the batch failed"
        for it in batches[0]:
            if it not in bad_batch:
                assert sum(1 for p in c.calls if it["stem"] in p) == 1, "a batch that worked was asked exactly once"
        # item 6: dropped with the reason, after the one retry the stage gives a refused item
        assert summary["retried"] == 1 and summary["batching"]["fallback_items"] == len(bad_batch) + 1
        assert summary["batching"]["fallback_failed"] == 2, "once in the first pass, once in the retry"
        d = next(d for d in summary["dropped"] if d["source_item_id"] == "q-6")
        assert [f["what"] for f in d["findings"]] == [T.UNUSABLE_REPLY, T.UNUSABLE_REPLY]
        assert d["findings"][1]["where"].endswith("(retry)")
        flags = T.bank_flags(summary)
        assert flags == [{"where": "item q-6", "what": "bank item not staged",
                          "detail": f"{T.UNUSABLE_REPLY}; {T.UNUSABLE_REPLY}"}], "the flag is the one the reviewer already knows"
        # the tally: one call per batch, plus the fallback and retry calls
        assert summary["batching"]["batches"] == len(batches) + 1 and summary["batching"]["batch_calls"] == len(batches) + 1
        assert summary["batching"]["model_calls"] == len(c.calls)

    def test_two_items_that_translate_to_one_sentence_are_still_one_item_with_the_twin_flag(self, monkeypatch):
        monkeypatch.setenv(T.BANK_CHUNK_ENV, "10")
        sb = self._sb()
        same = _arabize(_item(1)["stem"])

        def collapse(s: str) -> str:
            # item 3's stem comes back as item 1's sentence (both mcq): one
            # Arabic question, two English sources
            return same if s == _item(3)["stem"] else _arabize(s)

        c = FakeTranslator(collapse)
        kit = {"id": "kit-ar", "topic_id": "t-cell"}
        mapping, summary = T.bank_stage(sb, c, {"kit": sb.tables["topic_kits"][0]}, kit, "ar", [], "art-ar")
        assert "q-1" in mapping and "q-3" not in mapping and len(mapping) == 11
        d = next(d for d in summary["dropped"] if d["source_item_id"] == "q-3")
        assert d["findings"] == [{"where": "q-3", "what": "translated stem duplicates another item"}]
        assert T.bank_flags(summary) == [{"where": "item q-3", "what": "bank item not staged",
                                          "detail": "translated stem duplicates another item"}]
        assert summary["retried"] == 0, "a twin is a database fact, not a translation fault — no retry"


class TestQuestionCheck:
    def test_is_maths_tells_notation_from_words(self):
        for s in ("x = 5", "9π cm", "3x + 5 = 20", "7", "37.68", "C = 2 × 3.14 × 6", ""):
            assert C.is_maths(s), s
        for s in ("diameter", "Solve: 3x + 5 = 20", "Answer: x = 5", "3x = 15   (subtract 5 from both sides)",
                  "centre to the edge", "cm"):
            assert not C.is_maths(s), s

    def test_question_text_items_take_the_words_and_leave_the_notation(self):
        paths = dict(T.question_text_items(WS_QUESTIONS["questions"][3]))
        assert set(paths) == {"prompt", "working[0]", "working[2]", "working[3]"}, "x = 5 and the accepted 5 are copied"
        assert dict(T.question_text_items(WS_QUESTIONS["questions"][2])) == {
            "prompt": "Match each term to its meaning.", "pairs[0].left": "radius", "pairs[0].right": "centre to the edge",
            "pairs[1].left": "chord", "pairs[1].right": "a line joining two points on the circle"}
        assert "answer" not in dict(T.question_text_items(WS_QUESTIONS["questions"][4])), "37.68 cm is notation"
        assert dict(T.question_text_items(WS_QUESTIONS["questions"][0]))["answer"] == "diameter"

    def test_check_question_blocks_a_changed_number_a_lost_span_and_a_changed_picture(self):
        src = WS_QUESTIONS["questions"][3]
        ok = T.apply_strings(src, [(p, _arabize(t)) for p, t in T.question_text_items(src)])
        assert C.check_question(src, ok)["ok"] is True
        bad = {**ok, "prompt": "ترجمة Solve: ٣x + ٥ = ٢١"}
        whats = {f["what"] for f in C.check_question(src, bad)["blocking"]}
        assert {"numbers differ", "maths in the question altered"} <= whats
        lost = {**ok, "working": [*ok["working"][:2], "ترجمة Check: ٣(٥) + ٥", ok["working"][3]]}
        assert any(f["what"] == "maths in the working altered" for f in C.check_question(src, lost)["blocking"])
        fig = WS_QUESTIONS["questions"][5]
        swapped = {**fig, "prompt": _arabize(fig["prompt"]), "image": "data:image/png;base64,AAAA"}
        assert any(f["what"] == "image changed" for f in C.check_question(fig, swapped)["blocking"])
        tf = WS_QUESTIONS["questions"][1]
        assert any(f["what"] == "answer changed" for f in C.check_question(tf, {**tf, "answer": False})["blocking"])


class TestTranslatedWorksheet:
    def test_a_verified_maths_worksheet_is_translated_from_its_questions_and_rendered_as_a_kit_piece(self, monkeypatch, tmp_path):
        sb = _ws_job_sb()
        c = FakeTranslator()
        k, staged, uploads = _ws_run(sb, monkeypatch, tmp_path, c)

        assert len(staged) == 1
        body = staged[0]
        qs = {q["id"]: q for q in body["questions"]}
        assert list(qs) == ["q1", "q2", "q3", "q4", "q5", "q6"] and body["language"] == "ar"
        art = next(a for a in sb.tables["topic_articles"] if a["language"] == "ar")
        assert body["title"].endswith(art["title"]) and body["title"] != art["title"], "the sheet's title: the document word, the translated topic"
        assert body["source"] == {"kit_id": "kit-1", "language": "en", "title": "Worksheet — Circles", "questions": 6, "dropped": []}
        # the words translated, every number equivalent, the notation copied
        assert all(q["prompt"].startswith("ترجمة") for q in qs.values())
        assert qs["q1"]["answer"].startswith("ترجمة"), "a word answer is translated"
        assert qs["q4"]["answer"] == "x = 5" and qs["q4"]["accept"] == ["5"], "notation is copied"
        assert qs["q5"]["answer"] == "37.68 cm" and qs["q2"]["answer"] is True and qs["q6"]["image"] == PNG_1X1
        assert qs["q4"]["working"][1] == "x = 5" and qs["q5"]["working"][0] == "C = 2 × 3.14 × 6", "pure working stays verbatim"
        assert qs["q4"]["working"][0].startswith("ترجمة") and qs["q4"]["working"][3].startswith("ترجمة"), "a working line with words is translated"
        for src, dst in zip(WS_QUESTIONS["questions"], body["questions"]):
            assert C.numbers_equivalent(" ".join(C._q_strings(src)), " ".join(C._q_strings(dst)))[0], src["id"]
            assert (src["id"], src["type"], src["marks"]) == (dst["id"], dst["type"], dst["marks"])
        assert [p["left"] for p in qs["q3"]["pairs"]] == ["ترجمة radius", "ترجمة chord"]
        # the protected expressions travelled with the call
        assert any("3x + 5 = 20" in call and "UNCHANGED" in call for call in c.calls)
        # the piece: a worksheet generation of the kit, rendered from the staged file, no flag
        assert "worksheet" in k["doc_generation_ids"]
        g = next(g for g in sb.tables["generations"] if g["id"] == k["doc_generation_ids"]["worksheet"])
        assert g["kind"] == "worksheet" and g["status"] == "queued" and g["params"]["kit_id"] == k["id"]
        assert g["params"]["translation"]["questions"] == f"sys/{k['id']}/translated/worksheet_questions.json"
        assert g["params"]["translation"]["source_kit_id"] == "kit-1" and "question_set_id" not in g["params"]
        assert not [f for f in k["precheck"]["flags"] if f["what"] == T.DOCUMENT_FLAG]
        assert k["precheck"]["ok"] is True
        assert k["translation"]["stage"]["worksheet"] == {"questions": 6, "kept": 6, "dropped": [],
                                                           "path": g["params"]["translation"]["questions"]}
        assert set(k["doc_generation_ids"]) == {"deck", "activity", "case_study", "worksheet", "question_bank", "quiz"}

    def test_a_question_whose_translation_changes_a_number_is_dropped_with_a_flag_and_the_rest_rendered(self, monkeypatch, tmp_path):
        sb = _ws_job_sb()
        c = FakeTranslator(lambda s: _arabize(s.replace("20", "21")) if s == "Solve: 3x + 5 = 20" else _arabize(s))
        k, staged, _ = _ws_run(sb, monkeypatch, tmp_path, c)

        body = staged[0]
        assert [q["id"] for q in body["questions"]] == ["q1", "q2", "q3", "q5", "q6"] and body["source"]["dropped"] == ["q4"]
        assert "worksheet" in k["doc_generation_ids"], "the sheet is still made"
        flags = [f for f in k["precheck"]["flags"] if f["what"] == T.DOCUMENT_FLAG]
        assert len(flags) == 1 and flags[0]["where"] == "worksheet q4"
        assert "numbers differ" in flags[0]["detail"] and "5 of 6" in flags[0]["detail"]
        assert k["precheck"]["ok"] is True, "a dropped question is a flag, never a block"
        assert k["translation"]["stage"]["worksheet"]["dropped"] == ["q4"]
        # asked once more with the reason before it was dropped
        assert any("refused" in call and "numbers differ" in call and "3x + 5 = 20" in call for call in c.calls)

    def test_a_worksheet_with_no_surviving_question_is_not_made_but_flagged(self, monkeypatch, tmp_path):
        sb = _ws_job_sb()
        c = FakeTranslator(lambda s: _arabize(s + " 99") if s in WS_WORDS else _arabize(s))
        k, staged, _ = _ws_run(sb, monkeypatch, tmp_path, c)

        assert staged == [] and "worksheet" not in k["doc_generation_ids"]
        gens = [g for g in sb.tables["generations"] if (g["params"].get("translation") or {}).get("kit_id") == k["id"]]
        assert not any(g["kind"] == "worksheet" and not g["params"].get("question_set_id") for g in gens), "no rebuilt worksheet"
        flag = next(f for f in k["precheck"]["flags"] if f["where"] == "worksheet")
        assert flag["what"] == T.DOCUMENT_FLAG and "quiz" in flag["detail"] and "6 tried" in flag["detail"]
        assert len([f for f in k["precheck"]["flags"] if f["what"] == T.DOCUMENT_FLAG]) == 7, "one per question, one for the sheet"
        assert k["precheck"]["ok"] is True and "quiz" in k["doc_generation_ids"] and "question_bank" in k["doc_generation_ids"]

    def test_a_science_worksheet_without_questions_json_is_still_rebuilt(self, monkeypatch, tmp_path):
        sb = _job_sb()
        k, staged, _ = _ws_run(sb, monkeypatch, tmp_path, FakeTranslator())
        assert staged == [] and k["translation"]["stage"]["worksheet"] is None
        g = next(g for g in sb.tables["generations"] if g["id"] == k["doc_generation_ids"]["worksheet"])
        assert "questions" not in g["params"]["translation"], "rebuilt from the article, as before"


class TestTranslatedScriptValidation:
    def test_a_textbook_figure_slide_survives_the_models_round_trip(self):
        body = _maths_body()["script"]
        seg = body["segments"][0]
        seg["slide_visual"] = {"kind": "figure", "src": "books/x/fig3.png", "attribution": "p. 12", "caption": "خلية نباتية"}
        body["segments"][1]["slide_visual"] = {"kind": "icons", "items": [{"icon": "cell", "label": "خلية"}]}
        model, dump = T.validate_translated_script(body)
        assert dump["segments"][0]["slide_visual"] == seg["slide_visual"], "the figure is kept verbatim"
        assert dump["segments"][1]["slide_visual"]["kind"] == "icons", "a model kind goes through the model"
        assert model.segments[0].slide_visual is None, "the model itself never saw the figure"
        assert body["segments"][0]["slide_visual"]["kind"] == "figure", "the input is not mutated"

    def test_a_broken_field_still_fails_validation(self):
        body = _maths_body()["script"]
        body["segments"][0]["type"] = "not_a_type"
        with pytest.raises(Exception):
            T.validate_translated_script(body)


class TestPostRender:
    def test_drift_and_silence_flag_and_a_silent_segment_blocks(self):
        script = {"segments": [{"segment_id": "s1", "text": "a"}, {"segment_id": "s2", "text": "b"},
                               {"segment_id": "s3", "text": "c"}]}
        video = {"segments": [{"segment_id": "s1", "audio_duration_seconds": 10.0, "clip_duration_seconds": 10.5},
                              {"segment_id": "s2", "audio_duration_seconds": 10.0, "clip_duration_seconds": 18.0},
                              {"segment_id": "s3", "audio_duration_seconds": 0.0, "clip_duration_seconds": 4.0}]}
        r = C.post_render(script, video, overlaps=[{"segment_id": "s1", "pair": ["a", "b"]}])
        assert not r["ok"]
        assert {f["what"] for f in r["blocking"]} == {"no audio for a spoken segment"}
        assert {(f["where"], f["what"]) for f in r["flags"]} >= {
            ("s2", "clip and narration drift"), ("s2", "long silence after the narration"),
            ("s1", "text overlap reported by acceptance")}

    def test_merge_reports_unions_the_findings(self):
        m = C.merge_reports({"check": "A", "ok": True, "blocking": [], "flags": [{"what": "x"}]},
                            {"check": "C", "ok": False, "blocking": [{"what": "y"}], "flags": []})
        assert m["ok"] is False and [f["what"] for f in m["blocking"]] == ["y"] and len(m["checks"]) == 2


# ── the words in, the words out ─────────────────────────────────────────


class TestFieldSelection:
    def test_a_science_segment_round_trips_and_the_narration_follows_the_dialogue(self):
        seg = SRC["segments"][0]
        items = T.segment_text_items(seg)
        paths = [p for p, _ in items]
        assert "dialogue[0].line" in paths and "slide_points[1]" in paths and "scene.elements[0].text" in paths
        assert not any(".expr" in p for p in paths), "a math element is never sent to the translator"
        out = T._finish_segment(T.apply_strings(seg, [(p, _arabize(t)) for p, t in items]))
        assert out["text"] == " ".join(d["line"] for d in out["dialogue"]) and out["text"].startswith("ترجمة")
        assert out["elevenlabs_text"] == out["text"]
        assert out["scene"]["elements"][1]["expr"] == "P = πr + 2r"

    def test_lesson_text_items_take_the_spoken_fields_and_leave_the_maths(self):
        from tests.test_maths_lesson import LESSON
        paths = [p for p, _ in T.lesson_text_items(LESSON)]
        assert "hook[0].line" in paths and "method.steps[0]" in paths and "closing" in paths
        assert "examples[0].steps[0].speech" in paths and "try_it.speech" in paths
        assert any(p.endswith("common_mistake.why_wrong") for p in paths) and any(p.endswith("common_mistake.speech") for p in paths)
        # a QUESTION with words is translated (the founder's review of the
        # first Arabic video); a bare equation ("3x + 5 = 20", this fixture)
        # is typeset and stays; the working stays maths
        assert not any(p.endswith(".problem") for p in paths), "the fixture's questions are bare equations"
        import copy as _copy
        prose = _copy.deepcopy(LESSON)
        prose["examples"][0]["problem"] = "A coin has a radius of 4 cm. Find its circumference. (Use π = 3.14)"
        assert "examples[0].problem" in [p for p, _ in T.lesson_text_items(prose)]
        assert not any(p.endswith((".before[0]", ".after[0]", ".final_answer[0]", ".givens[0]")) for p in paths)
        # a step's operation is the board's muted note: translated (founder, 2026-10-17)
        assert any(p.endswith(".operation") for p in paths)

    def test_article_text_items_leave_ids_codes_and_figure_keys(self):
        paths = [p for p, _ in T.article_text_items(ARTICLE)]
        assert "title" in paths and "sections[0].heading" in paths and "sections[0].body_md" in paths
        assert "glossary[0].term" in paths and "glossary[0].definition" in paths and "claims[0].text" in paths
        assert not any("figure_keys" in p or "covers" in p or p.endswith(".id") for p in paths)

    def test_translate_strings_keeps_order_skips_empties_and_chunks(self, monkeypatch):
        monkeypatch.setattr(T, "CHUNK", 3)
        c = FakeTranslator(lambda s: s.upper())
        out = T.translate_strings(c, ["a", "", "b", "c", "d", "e"], "fr")
        assert out == ["A", "", "B", "C", "D", "E"] and len(c.calls) == 2

    def test_a_reply_with_the_wrong_count_is_asked_again_then_refused(self):
        class Short(FakeTranslator):
            def analyze(self, prompt, **k):
                self.calls.append(prompt)
                return {"data": {"items": ["only one"]}, "usage": {}}
        with pytest.raises(RuntimeError, match="did not return 2 items"):
            T.translate_strings(Short(), ["a", "b"], "fr")


# ── one part: the maths path recompiles the board ───────────────────────


def _maths_body():
    from maths.lesson import to_episode_script
    from maths.schema import Lesson
    from tests.test_maths_lesson import LESSON
    ep = to_episode_script(Lesson.model_validate(LESSON), {"status": "verified"}, book_id="catalogue-t-lin",
                           chapter_num=-1, episode_num=1, title="Linear Equations in One Variable",
                           avatars={"teacher": "avatar_teacher_female__face_f1", "student": "avatar_student_male__face_m1"},
                           language="en").model_dump()
    return {"part": 1, "of": 1, "language": "en", "voice": "g-en-f", "student_voice": "g-en-student-m",
            "avatars": ep["avatars"], "script": ep, "video": {"segments": []}}


class TestTranslateEpisodeMaths:
    def test_the_board_is_recompiled_in_the_language_and_the_maths_is_identical(self):
        body = _maths_body()
        c = FakeTranslator()
        dst, report = T.translate_episode(c, body, "ar", [("equation", "معادلة")], context="")
        src = body["script"]
        assert report["ok"], report["blocking"]
        assert dst["maths"]["language"] == "ar"
        assert dst["maths"]["lesson"]["examples"][0]["steps"][0]["after"] == src["maths"]["lesson"]["examples"][0]["steps"][0]["after"]
        assert dst["maths"]["lesson"]["try_it"]["problem"] == src["maths"]["lesson"]["try_it"]["problem"]
        assert dst["maths"]["lesson"]["hook"][0]["line"].startswith("ترجمة")
        assert dst["avatars"] == src["avatars"], "the same faces in every language"
        # the board's own words come from maths.i18n, not the translator:
        # the try-it board says "try it" in Arabic, never in English
        board = " ".join(e.get("text") or "" for s in dst["segments"]
                         for e in ((s.get("scene") or {}).get("elements") or []) if e.get("type") == "text")
        assert "جرّب" in board and "Try it" not in board and "Pause the video" not in board
        assert abs(len(dst["segments"]) - len(src["segments"])) <= 2
        # the structural model accepts what the compiler produced
        from agent3_scripts.models import EpisodeScript
        EpisodeScript.model_validate(dst)

    def test_every_example_of_an_extended_record_reaches_the_recompiled_board(self):
        """The English kit's record may hold more than the schema's four
        examples (the length floor appends verified ones after validation);
        the recompile must keep them all — the Circle pilot lost four."""
        import copy as _copy
        from maths.schema import Lesson
        from tests.test_maths_lesson import LESSON
        rec = _copy.deepcopy(LESSON)
        base = rec["examples"][0]
        rec["examples"] = [{**_copy.deepcopy(base), "label": f"Example {i}", "problem": f"{i}x + 1 = {i * 3 + 1}",
                            "final_answer": ["x = 3"]} for i in range(1, 7)]
        assert len(Lesson.model_validate(rec).examples) == 4, "the schema caps a model reply at four"
        built = T.lesson_from_record(rec)
        assert [e.label for e in built.examples] == [f"Example {i}" for i in range(1, 7)]
        body = {"part": 1, "of": 1, "language": "en", "avatars": {}, "video": {},
                "script": {"script_id": "sc", "book_id": "b", "chapter_num": -1, "episode_num": 1,
                           "episode_title": "Linear", "generated_at": "now", "narrator_persona": "x",
                           "segments": [], "visual_plan": None, "avatars": None,
                           "maths": {"lesson": rec, "verification": {}, "language": "en"},
                           "total_estimated_duration_seconds": 1, "question_hook_count": 0}}
        dst = T.recompile_maths(body["script"], rec, "ar", "خطي")
        board = " ".join(e.get("expr") or "" for s in dst["segments"]
                         for e in ((s.get("scene") or {}).get("elements") or []) if e.get("type") == "math")
        assert all(f"{i}x + 1 = {i * 3 + 1}" in board for i in range(1, 7)), board

    def test_a_translation_that_changes_a_spoken_number_is_retried_then_refused(self):
        body = _maths_body()
        c = FakeTranslator(lambda s: _arabize(s.replace("five", "5").replace("x equals 5", "x equals 6")) if "equals" in s else _arabize(s))
        with pytest.raises(T.TranslateRefused, match="pre-check A"):
            T.translate_episode(c, body, "ar", [], context="")
        assert c.session_usage["calls"] >= 2, "the part is asked again once before the job gives up"


class TestTranslateEpisodeScience:
    def test_segments_keep_ids_turns_and_scene_plan(self):
        body = {"part": 1, "of": 1, "language": "en", "avatars": {}, "video": {},
                "script": {"script_id": "sc", "book_id": "b", "chapter_num": -1, "episode_num": 1,
                           "episode_title": "Cells", "generated_at": "now", "narrator_persona": "x",
                           "segments": copy.deepcopy(SRC["segments"]), "visual_plan": None, "avatars": None,
                           "total_estimated_duration_seconds": 40, "question_hook_count": 0}}
        dst, report = T.translate_episode(FakeTranslator(), body, "ar", [], context="")
        assert report["ok"], report["blocking"]
        assert [s["segment_id"] for s in dst["segments"]] == ["s001", "s002"]
        assert dst["segments"][0]["scene"]["elements"][1]["expr"] == "P = πr + 2r"
        assert dst["segments"][0]["dialogue"][1]["who"] == "student"
        assert dst["episode_title"].startswith("ترجمة")


# ── the job ──────────────────────────────────────────────────────────────

SCRIPT_PATH = "sys/gen-p/script.json"


def _job_sb(*, kit_status="approved", with_assets=True):
    sb = _sb(kit_row={**KIT, "status": kit_status, "source_kit_id": None, "voice_pair": {"teacher": "g-en-f", "student": "g-en-student-m"}})
    sb.tables["artifacts"] = [{"id": "a1", "generation_id": "gen-p", "kind": "video_mp4", "storage_path": "sys/gen-p/lesson.mp4"},
                              {"id": "a2", "generation_id": "gen-p", "kind": "script_json", "storage_path": SCRIPT_PATH}]
    body = _maths_body()
    sb.files[("artifacts", SCRIPT_PATH)] = json.dumps(body).encode("utf-8")
    sb.tables["topic_term_translations"] = []
    # the roster rows as the library keeps them: asset_key = the ROSTER key,
    # canonical_key = the LIBRARY normaliser's spelling (avatar_female_teacher,
    # not the topic-key one's), the face key's suffix = the row id's prefix
    from shared.asset_keys import canonical_key as _ck
    assert _ck("avatar_teacher_female") == "avatar_female_teacher"
    sb.tables["visual_assets"] = ([{"id": "f1a2b3c4-0000-4000-8000-000000000001", "asset_key": "avatar_teacher_female",
                                    "canonical_key": _ck("avatar_teacher_female"), "storage_path": "lib/f1.png", "status": "approved"},
                                   {"id": "m1a2b3c4-0000-4000-8000-000000000002", "asset_key": "avatar_student_male",
                                    "canonical_key": _ck("avatar_student_male"), "storage_path": "lib/m1.png", "status": "approved"}]
                                  if with_assets else [])
    sb.tables["jobs"] = [{"id": "job-t", "type": "kit_translate", "status": "processing", "generation_id": None,
                          "params": {"source_kit_id": "kit-1", "language": "ar"}}]
    # the English bank: 12 approved items (two kinds), and the source kit's
    # composed quiz — a worksheet with a question set of 4 of them
    sb.tables["topic_questions"] = [_item(i) for i in range(1, 13)]
    sb.tables["topic_kits"][0]["doc_generation_ids"] = {**KIT["doc_generation_ids"], "quiz": "gen-quiz", "question_bank": "gen-bank"}
    sb.tables["question_sets"] = [{"id": "set-en", "blueprint_id": "bp-1", "topic_ids": ["t-cell"], "language": "en",
                                   "question_ids": ["q-3", "q-1", "q-8", "q-5"], "seed": 42}]
    sb.tables["generations"] += [
        {"id": "gen-quiz", "kind": "worksheet", "status": "done", "owner_id": "sys", "book_id": None, "chapter_ref": None,
         "params": {"catalogue": True, "topic_id": "t-cell", "article_id": "art-1", "language": "en",
                    "question_set_id": "set-en", "blueprint_id": "bp-1", "seed": 42}},
        {"id": "gen-bank", "kind": "question_bank", "status": "done", "owner_id": "sys", "book_id": None, "chapter_ref": None,
         "params": {**PARAMS}},
    ]
    return sb


_PARTS = ["nucleus", "membrane", "wall", "vacuole", "ribosome", "cytoplasm",
          "chloroplast", "mitochondrion", "lysosome", "nucleolus", "vesicle", "centriole"]


def _item(i: int) -> dict:
    # stems differ in WORDS, not only in a number: content_hash canonicalises
    # the stem, so two stems that differ only by a digit are one item
    mcq = i % 2 == 1
    part = _PARTS[(i - 1) % len(_PARTS)]
    return {"id": f"q-{i}", "topic_id": "t-cell", "article_id": "art-1", "language": "en", "status": "approved",
            "item_type": "mcq" if mcq else "fill_blank", "answer_mode": "objective", "difficulty": 1 + i % 3,
            "cognitive_level": "recall", "marks": 1, "est_seconds": 60, "objective_ref": "o1", "claim_ref": None,
            "stem": f"Which part of a cell is the {part} and what does it do?" if mcq else f"The {part} of a cell is ____.",
            "options": ([{"key": "A", "text": "a nucleus"}, {"key": "B", "text": "a wall"},
                         {"key": "C", "text": "a membrane"}, {"key": "D", "text": "a vacuole"}] if mcq else None),
            "answer": {"key": "A"} if mcq else {"text": "a nucleus"},
            "marking_scheme": [], "distractor_rationale": ({"B": {"why_wrong": "walls are plant only"}} if mcq else None),
            "explanation": "Because the nucleus holds the DNA.", "tags": [], "content_hash": f"h{i}"}


def _uploads(monkeypatch):
    uploads = []
    from worker import client as db

    def _upload(_sb, local_path, dest):
        uploads.append((dest, json.loads(open(local_path, encoding="utf-8").read())))
        return dest
    monkeypatch.setattr(db, "upload_artifact", _upload)
    return uploads


class TestTranslateJob:
    def test_a_kit_is_translated_end_to_end_without_a_single_render_call(self, monkeypatch, tmp_path):
        sb = _job_sb()
        uploads = _uploads(monkeypatch)
        monkeypatch.setattr("spike.scene_engine.raster_assets.cache_dir_for", lambda key, cache_dir=None: tmp_path / "nocache" / key)
        c = FakeTranslator()

        summary = T.run_translate_job(sb, sb.tables["jobs"][0], client=c)

        assert summary and summary["step"] == "done" and summary["language"] == "ar" and summary["precheck_ok"] is True
        job = sb.tables["jobs"][0]
        assert job["status"] == "done" and job["error"] is None
        # the translated kit row: lineage, staged, the language's voices
        kits = [k for k in sb.tables["topic_kits"] if k.get("source_kit_id") == "kit-1"]
        assert len(kits) == 1
        k = kits[0]
        assert (k["language"], k["status"], k["translation_status"]) == ("ar", "generating", "staged")
        assert k["voice_pair"] == {"teacher": "g-ar-f", "student": "g-ar-student-m"}
        # the article: staged, the source named, the figures copied
        arts = [a for a in sb.tables["topic_articles"] if a["language"] == "ar"]
        assert len(arts) == 1 and arts[0]["status"] == "staged" and arts[0]["source_article_id"] == "art-1"
        assert arts[0]["title"].startswith("ترجمة") and arts[0]["sections"][0]["id"] == "s1"
        assert arts[0]["sections"][1]["figure_keys"] == ["animal_cell", "plant_cell"]
        assert k["article_id"] == arts[0]["id"]
        copied = [f for f in sb.tables["article_figures"] if f["article_id"] == arts[0]["id"]]
        assert len(copied) == len(FIGURES) and all(f["caption"].startswith("ترجمة") for f in copied if f["caption"])
        # the glossary, pinned first
        terms = sb.tables["topic_term_translations"]
        assert {t["term_en"] for t in terms} == {"cell", "organelle"} and all(t["language"] == "ar" and t["source"] == "model" for t in terms)
        assert "cell → ترجمة cell" in c.calls[-1] or any("cell → ترجمة cell" in p for p in c.calls)
        # the bank (D4): every English item staged for THIS kit, by source id
        staged = [q for q in sb.tables["topic_questions"] if q.get("translation_kit_id") == k["id"]]
        assert len(staged) == 12 and all(q["status"] == "staged" and q["language"] == "ar" and q["article_id"] == arts[0]["id"] for q in staged)
        by_src = {q["source_item_id"]: q for q in staged}
        assert by_src["q-1"]["stem"].startswith("ترجمة") and by_src["q-1"]["answer"] == {"key": "A"}
        assert [o["key"] for o in by_src["q-1"]["options"]] == ["A", "B", "C", "D"] and by_src["q-1"]["options"][0]["text"].startswith("ترجمة")
        assert by_src["q-2"]["answer"]["text"].startswith("ترجمة"), "a text answer is translated"
        assert by_src["q-1"]["content_hash"] != "h1", "the hash is the translated stem's"
        # the quiz: the SAME four questions, by id, in the source set's order
        sets = [s for s in sb.tables["question_sets"] if s["language"] == "ar"]
        assert len(sets) == 1 and sets[0]["question_ids"] == [by_src[i]["id"] for i in ("q-3", "q-1", "q-8", "q-5")]
        assert sets[0]["blueprint_id"] == "bp-1" and sets[0]["seed"] == 42
        # the generations: every piece of the source kit
        gens = [g for g in sb.tables["generations"] if (g["params"].get("translation") or {}).get("kit_id") == k["id"]]
        labels = {("quiz" if g["params"].get("question_set_id") else g["kind"]): g for g in gens}
        assert set(labels) == {"presentation", "deck", "activity", "case_study", "worksheet", "question_bank", "quiz"}
        p = labels["presentation"]["params"]
        assert p["catalogue"] is True and p["language"] == "ar" and p["tts_voice"] == "g-ar-f" and p["article_id"] == arts[0]["id"]
        assert p["translation"]["source_kit_id"] == "kit-1" and p["translation"]["scripts"] == [f"sys/{k['id']}/translated/script.json"]
        assert "scripts" not in labels["deck"]["params"]["translation"]
        assert labels["presentation"]["owner_id"] == "sys" and labels["presentation"]["book_id"] is None
        q = labels["quiz"]["params"]
        assert "kit_id" not in q and q["question_set_id"] == sets[0]["id"] and q["translation"]["kit_id"] == k["id"], "a composed set has no kit_id; the marker ties it to the kit"
        assert k["presentation_generation_id"] == labels["presentation"]["id"]
        assert set(k["doc_generation_ids"]) == {"deck", "activity", "case_study", "worksheet", "question_bank", "quiz"}
        # pre-check B ran on a sample and flagged nothing (the fake keeps the words)
        bank = dict(k["precheck"]["bank"])
        batching = bank.pop("batching")
        assert bank == {"source_items": 12, "staged": 12, "reused": 0, "retried": 0}
        assert batching["batches"] >= 1 and batching["batch_calls"] == batching["batches"] and batching["fallback_items"] == 0
        assert k["translation"]["stage"]["bank"]["batching"] == batching, "the stage record carries the batch tally"
        fid = next(c for c in k["precheck"]["checks"] if c.get("check") == "B")
        assert fid["samples"] >= 4 and fid["mean_score"] is not None and fid["mean_score"] > 0.9
        # the staged script: the translated lesson, Arabic board words
        assert [d for d, _ in uploads] == [f"sys/{k['id']}/translated/script.json"]
        staged = uploads[0][1]
        assert staged["language"] == "ar" and staged["script"]["maths"]["language"] == "ar"
        # provenance and the pre-check on the kit
        tr = k["translation"]
        assert tr["source_scripts"][0]["path"] == SCRIPT_PATH and tr["assets"] and tr["cost"]["translation"]["usage"]["calls"] >= 3
        assert k["precheck"]["ok"] is True and k["precheck"]["parts"] == [{"part": 1, "ok": True, "blocking": 0, "flags": k["precheck"]["parts"][0]["flags"]}]
        # nothing from the English pipeline was asked for
        assert ("select", "books") not in sb.calls
        assert sb.tables["topics"][0]["status"] == TOPIC_STATUS_UNCHANGED(sb)

    def test_a_missing_picture_is_redrawn_by_the_render_and_flagged(self, monkeypatch, tmp_path):
        """Founder, 2026-10-17: a picture that exists nowhere but inside the
        English video is drawn afresh by the translated render, and the
        reviewer sees a flag per picture; the presentation carries the list
        the renderer may draw (params.translation.redraw)."""
        sb = _job_sb(with_assets=False)
        _uploads(monkeypatch)
        monkeypatch.delenv(T.REDRAW_ENV, raising=False)
        monkeypatch.setattr("spike.scene_engine.raster_assets.cache_dir_for", lambda key, cache_dir=None: tmp_path / "nocache" / key)

        summary = T.run_translate_job(sb, sb.tables["jobs"][0], client=FakeTranslator())
        assert summary is not None and sb.tables["jobs"][0]["status"] == "done"
        k = next(k for k in sb.tables["topic_kits"] if k.get("source_kit_id") == "kit-1")
        redrawn = [f for f in k["precheck"]["flags"] if f["what"] == T.REDRAW_FLAG]
        assert redrawn and all("English video" in f["detail"] for f in redrawn)
        assert k["precheck"]["ok"] is True, "a redrawn picture is a flag, never a block"
        pres = next(g for g in sb.tables["generations"] if g["kind"] == "presentation" and (g["params"].get("translation") or {}).get("kit_id") == k["id"])
        assert sorted(pres["params"]["translation"]["redraw"]) == sorted(f["where"] for f in redrawn)
        assert k["precheck"]["bank"]["staged"] == 12 and len(redrawn) == len(pres["params"]["translation"]["redraw"]),             "the bank stage's write did not replace the assets stage's flags"
        assert "avatar_teacher_female__face_f1" in pres["params"]["translation"]["redraw"]
        assert k["translation"]["stage"]["assets"]["redraw"] == pres["params"]["translation"]["redraw"]
        assert k["translation"]["stage"]["assets"]["shelf"] == [], "nothing came from the colour shelf"

    def test_the_refusal_is_one_switch_away(self, monkeypatch, tmp_path):
        sb = _job_sb(with_assets=False)
        _uploads(monkeypatch)
        monkeypatch.setenv(T.REDRAW_ENV, "0")
        monkeypatch.setattr("spike.scene_engine.raster_assets.cache_dir_for", lambda key, cache_dir=None: tmp_path / "nocache" / key)
        assert T.run_translate_job(sb, sb.tables["jobs"][0], client=FakeTranslator()) is None
        job = sb.tables["jobs"][0]
        assert job["status"] == "error" and "cannot be found" in job["error"]
        # a translation that failed before any piece is not kept: no kit row,
        # no staged article, no staged items — the job's error is the record
        assert not [k for k in sb.tables["topic_kits"] if k.get("source_kit_id") == "kit-1"]
        assert not [a for a in sb.tables["topic_articles"] if a.get("language") == "ar"]
        assert not [g for g in sb.tables["generations"] if (g["params"].get("translation") or {}).get("source_kit_id") == "kit-1"], "nothing was queued to render"

    def test_the_renderer_draws_only_the_listed_keys_under_the_switch(self, monkeypatch):
        from spike.scene_engine import colour as COL
        from spike.scene_engine import raster_assets as RA
        monkeypatch.setattr(COL, "pictures_enabled", lambda: False)
        token = RA._GENERATION_VAR.set("gen-x")
        try:
            RA.set_generation_forbidden(True, "gen-x", allow_keys=["sk_whale", "plant_upgrades" + COL.COLOUR_KEY_SUFFIX])
            b = RA._STATE["gen-x"]
            assert b["forbid_generate"] is True and b["redraw_keys"] == {"sk_whale", "plant_upgrades"}, "kept by the script's plain key"
            assert RA.generation_forbidden("sk_whale") is False
            assert RA.generation_forbidden("sk_whale" + COL.COLOUR_KEY_SUFFIX) is False, "the colour key is the same picture"
            assert RA.generation_forbidden("plant_upgrades") is False
            assert RA.generation_forbidden("eukaryotic_factory") is True, "a board key off the list stays forbidden"
            assert RA.generation_forbidden("avatar_teacher__face_1") is True
            assert RA.generation_forbidden() is True
            # the board-colour pictures switch: the ink library cannot serve a
            # colour picture, so every BOARD key may be drawn; avatars never
            monkeypatch.setattr(COL, "pictures_enabled", lambda: True)
            assert RA.generation_forbidden("eukaryotic_factory") is False
            assert RA.generation_forbidden("eukaryotic_factory" + COL.COLOUR_KEY_SUFFIX) is False
            assert RA.generation_forbidden("avatar_teacher__face_1") is True
        finally:
            RA.set_generation_forbidden(False, "gen-x")
            RA._GENERATION_VAR.reset(token)
            assert RA._STATE["gen-x"]["redraw_keys"] == set()

    def test_under_colour_pictures_every_board_picture_is_a_redraw_and_a_flag(self, monkeypatch, tmp_path):
        from spike.scene_engine import colour as COL
        sb = _job_sb(with_assets=True)
        # a board picture the ink LIBRARY holds — found when pictures are ink,
        # a redraw when they are coloured (the render cannot use the ink row)
        from shared.asset_keys import canonical_key as asset_key_of
        sb.tables["visual_assets"].append({"id": "b0a0b0c0-0000-4000-8000-00000000000b", "asset_key": "sk_whale",
                                           "canonical_key": asset_key_of("sk_whale"), "storage_path": "lib/whale.png", "status": "approved"})
        body = json.loads(sb.files[("artifacts", SCRIPT_PATH)].decode("utf-8"))
        body["script"]["segments"][0].setdefault("scene", {"id": "x", "elements": [], "actions": []})
        body["script"]["segments"][0]["scene"]["elements"].append({"id": "pic", "type": "illustration", "asset": "sk_whale", "at": [600, 380], "scale": 1.0})
        sb.files[("artifacts", SCRIPT_PATH)] = json.dumps(body).encode("utf-8")
        monkeypatch.setattr("spike.scene_engine.raster_assets.cache_dir_for", lambda key, cache_dir=None: tmp_path / "nocache" / key)
        monkeypatch.setattr(COL, "pictures_enabled", lambda: False)
        found, missing = T.resolve_assets(sb, ["sk_whale", "avatar_teacher_female__face_f1"])
        assert [f["key"] for f in found] == ["sk_whale", "avatar_teacher_female__face_f1"] and missing == []
        monkeypatch.setattr(COL, "pictures_enabled", lambda: True)
        found, missing = T.resolve_assets(sb, ["sk_whale", "avatar_teacher_female__face_f1"])
        assert [f["key"] for f in found] == ["avatar_teacher_female__face_f1"] and missing == ["sk_whale"], "the avatar still comes from the roster"
        # the colour SHELF (2026-10-17): the library row keyed by the exact
        # colour key is the picture the render will hydrate — found, not a
        # redraw; the ink row is still not it
        import shared.visual_library as VL
        monkeypatch.delenv(VL.COLOUR_SHELF_ENV, raising=False)
        sb.tables["visual_assets"].append({"id": "c0a0b0c0-0000-4000-8000-00000000000c", "asset_key": "sk_whale" + COL.COLOUR_KEY_SUFFIX,
                                           "canonical_key": asset_key_of("sk_whale" + COL.COLOUR_KEY_SUFFIX),
                                           "storage_path": "lib/whale_colour.png", "status": "approved", "asset_format": "png"})
        found, missing = T.resolve_assets(sb, ["sk_whale", "avatar_teacher_female__face_f1"])
        assert missing == [] and [f["key"] for f in found] == ["sk_whale", "avatar_teacher_female__face_f1"]
        whale = found[0]
        assert whale["source"] == "library" and whale["colour"] is True and whale["cache_key"] == "sk_whale" + COL.COLOUR_KEY_SUFFIX
        assert whale["id"] == "c0a0b0c0-0000-4000-8000-00000000000c" and whale["storage_path"] == "lib/whale_colour.png"
        assert "colour" not in found[1], "the avatar is the roster's, not the shelf's"
        # a retired shelf row, or one without bytes, is a miss
        sb.tables["visual_assets"][-1]["status"] = "retired"
        assert T.resolve_assets(sb, ["sk_whale"])[1] == ["sk_whale"]
        sb.tables["visual_assets"][-1]["status"] = "approved"
        # the shelf is one variable away: off, every board picture is a redraw again
        monkeypatch.setenv(VL.COLOUR_SHELF_ENV, "0")
        assert T.resolve_assets(sb, ["sk_whale"])[1] == ["sk_whale"]
        monkeypatch.delenv(VL.COLOUR_SHELF_ENV, raising=False)
        # with the pictures switch OFF the shelf row is not consulted: the
        # render asks for the ink key, and only the ink row answers
        monkeypatch.setattr(COL, "pictures_enabled", lambda: False)
        found, _ = T.resolve_assets(sb, ["sk_whale"])
        assert found[0]["id"] == "b0a0b0c0-0000-4000-8000-00000000000b" and "colour" not in found[0]
        # (the job-level flag and redraw list for a missing key are pinned by
        # test_a_missing_picture_is_redrawn_by_the_render_and_flagged; this
        # fixture is a maths kit whose scenes are recompiled, so a board
        # picture injected here would not survive to the assets stage)

    def test_a_pre_check_refusal_names_its_findings_and_removes_the_kit(self, monkeypatch, tmp_path, caplog):
        sb = _job_sb()
        _uploads(monkeypatch)
        monkeypatch.setattr("spike.scene_engine.raster_assets.cache_dir_for", lambda key, cache_dir=None: tmp_path / "nocache" / key)
        c = FakeTranslator(lambda s: _arabize(s.replace("five", "5").replace("x equals 5", "x equals 6")) if "equals" in s else _arabize(s))

        with caplog.at_level(logging.WARNING, logger="worker.translate"):
            assert T.run_translate_job(sb, sb.tables["jobs"][0], client=c) is None

        # the kit is gone (a failed translation is not stored); the job's error
        # and the log carry the findings
        assert not [k for k in sb.tables["topic_kits"] if k.get("source_kit_id") == "kit-1"]
        assert "pre-check A" in sb.tables["jobs"][0]["error"] and "numbers differ" in sb.tables["jobs"][0]["error"]
        assert "refused" in caplog.text and "numbers differ" in caplog.text

    @pytest.mark.parametrize("params, why", [
        ({"source_kit_id": "kit-1", "language": "en"}, "already English"),
        ({"source_kit_id": "kit-1", "language": "xx"}, "not a lesson language"),
        ({"source_kit_id": "nope", "language": "ar"}, "not found"),
        ({"language": "ar"}, "source_kit_id"),
    ])
    def test_refusals_before_any_call(self, params, why, monkeypatch):
        sb = _job_sb()
        sb.tables["jobs"][0]["params"] = params
        c = FakeTranslator()
        assert T.run_translate_job(sb, sb.tables["jobs"][0], client=c) is None
        assert why in sb.tables["jobs"][0]["error"] and c.calls == []

    def test_an_unapproved_source_kit_is_refused(self):
        sb = _job_sb(kit_status="in_review")
        c = FakeTranslator()
        assert T.run_translate_job(sb, sb.tables["jobs"][0], client=c) is None
        assert "not approved" in sb.tables["jobs"][0]["error"] and c.calls == []

    def test_a_source_without_a_script_json_is_refused(self):
        sb = _job_sb()
        sb.tables["artifacts"] = [a for a in sb.tables["artifacts"] if a["kind"] != "script_json"]
        assert T.run_translate_job(sb, sb.tables["jobs"][0], client=FakeTranslator()) is None
        assert "script_json" in sb.tables["jobs"][0]["error"]


def TOPIC_STATUS_UNCHANGED(sb):
    # the source topic is wherever it was (published, say): a translation
    # never moves it — the fixture's topic says "generating" and so must it
    return "generating"


# ── the render: script_override in worker.process ───────────────────────

TRANSLATED_SCRIPT_PATH = "sys/kit-ar/translated/script.json"


def _override_script():
    heads = ["ما هي الخلية", "تراكيب مشتركة", "الخلايا النباتية"]
    segs = [{"segment_id": f"s1{i:02d}", "type": "explore", "text": f"المعلم {i}. يسأل الطالب.",
             "elevenlabs_text": f"المعلم {i}. يسأل الطالب.", "slide_heading": h, "slide_points": [],
             "dialogue": [{"who": "teacher", "line": f"المعلم {i}."}, {"who": "student", "line": "يسأل الطالب."}],
             "estimated_duration_seconds": 50} for i, h in enumerate(heads, start=1)]
    return {"script_id": "sc-ar", "book_id": "catalogue-t-cell", "chapter_num": -1, "episode_num": 1,
            "episode_title": "الخلايا", "generated_at": "now", "narrator_persona": "conversational",
            "segments": segs, "visual_plan": None,
            "avatars": {"teacher": "avatar_teacher_female__face_f1", "student": "avatar_student_male__face_m1"},
            "total_estimated_duration_seconds": 150, "question_hook_count": 0}


def _render_sb():
    sb = _sb(statuses={"gen-p": "done"})
    sb.tables["topic_articles"].append({**ARTICLE, "id": "art-ar", "language": "ar", "status": "staged",
                                        "source_article_id": "art-1", "title": "الخلايا"})
    sb.tables["topic_kits"][0]["status"] = "approved"
    sb.tables["topic_kits"].append({**KIT, "id": "kit-ar", "article_id": "art-ar", "language": "ar", "status": "generating",
                                    "source_kit_id": "kit-1", "translation_status": "staged",
                                    "presentation_generation_id": "gen-par", "doc_generation_ids": {"deck": "gen-dar"},
                                    "translation": {"source_kit_id": "kit-1"}, "precheck": {"ok": True, "blocking": [], "flags": []}})
    tr = {"source_kit_id": "kit-1", "source_generation_id": "gen-p", "source_language": "en", "scripts": [TRANSLATED_SCRIPT_PATH]}
    params = {**PARAMS, "kit_id": "kit-ar", "article_id": "art-ar", "language": "ar", "tts_voice": "g-ar-f",
              "student_voice": "g-ar-student-m", "translation": tr}
    sb.tables["generations"].append({"id": "gen-par", "kind": "presentation", "status": "processing", "owner_id": "sys",
                                     "book_id": None, "chapter_ref": None, "params": params})
    sb.tables["generations"].append({"id": "gen-dar", "kind": "deck", "status": "queued", "owner_id": "sys",
                                     "book_id": None, "chapter_ref": None,
                                     "params": {**params, "translation": {k: v for k, v in tr.items() if k != "scripts"}}})
    sb.files[("artifacts", TRANSLATED_SCRIPT_PATH)] = json.dumps(
        {"part": 1, "of": 1, "language": "ar", "script": _override_script()}).encode("utf-8")
    job = _job("job-par", "gen-par", "presentation")
    job["params"]["kit_id"] = "kit-ar"
    sb.tables["jobs"] = [job]
    return sb


class TestScriptOverrideRender:
    def test_a_translated_presentation_renders_the_staged_script_with_no_analysis_and_no_script_call(self, monkeypatch, tmp_path):
        sb = _render_sb()
        process, uploads, analysis_calls = _worker_env(monkeypatch, sb)
        fakes = _presentation_fakes(monkeypatch, tmp_path)
        forbidden: list[bool] = []
        monkeypatch.setattr("spike.scene_engine.raster_assets.set_generation_forbidden",
                            lambda on, generation_id=None, **kw: forbidden.append(bool(on)))

        process.process_generation(sb, sb.tables["jobs"][0], "gen-par")

        assert analysis_calls == [], "a translated build never analyses"
        assert fakes["scripts"] == [], "a translated build never writes a script"
        assert forbidden == [True, False], "image generation forbidden for the build, released after"
        # the slides took the translated segments, the composer the Arabic voices
        assert fakes["slide_inputs"] == [["s101", "s102", "s103"]]
        assert fakes["compose_calls"] == [{"student_voice": "g-ar-student-m", "tts_voice": "g-ar-f", "lang": "ar"}]
        # the part's own script_json: Arabic, the source kit's faces
        body = json.loads(next(raw for _, d, raw in uploads if d.endswith("gen-par/script.json")).decode("utf-8"))
        assert body["language"] == "ar" and body["voice"] == "g-ar-f"
        assert body["script"]["segments"][0]["slide_heading"] == "ما هي الخلية"
        assert body["avatars"] == _override_script()["avatars"]
        # the kit lifecycle ran: timestamps, a lesson plan in the language
        # carrying the translation marker without the scripts
        k = next(k for k in sb.tables["topic_kits"] if k["id"] == "kit-ar")
        assert [c["part"] for c in k["chapters"]] == [1] and k["chapters"][0]["chapters"][0]["label"] == "ما هي الخلية"
        lp = next(g for g in sb.tables["generations"] if g["kind"] == "lesson_plan")
        assert lp["params"]["language"] == "ar" and lp["params"]["article_id"] == "art-ar"
        assert lp["params"]["translation"] == {"source_kit_id": "kit-1", "source_generation_id": "gen-p", "source_language": "en"}
        assert k["doc_generation_ids"]["lesson_plan"] == lp["id"]
        # pre-check C recorded on the kit; the English kit and its topic untouched
        post = k["precheck"]["post_render"]
        assert post["ok"] is True and post["checks"][0]["check"] == "C" and k["precheck"]["ok"] is True
        assert sb.tables["topic_kits"][0]["status"] == "approved"
        gen = next(g for g in sb.tables["generations"] if g["id"] == "gen-par")
        assert gen["status"] == "done"
        assert ("select", "books") not in sb.calls

    def test_a_translated_kit_in_review_carries_translation_status_in_review(self):
        """approve_topic_kit's translation branch accepts in_review only, so
        the lifecycle column follows the kit's status; an English kit (no
        translation_status) is untouched."""
        from catalogue import kit as K
        sb = _render_sb()
        k = sb.tables["topic_kits"][1]
        k["doc_generation_ids"] = {"deck": "gen-dar", "lesson_plan": "gen-lar"}
        for gid in ("gen-par", "gen-dar"):
            next(g for g in sb.tables["generations"] if g["id"] == gid)["status"] = "done"
        sb.tables["generations"].append({"id": "gen-lar", "kind": "lesson_plan", "status": "done", "owner_id": "sys",
                                         "book_id": None, "chapter_ref": None, "params": {}})
        assert K.sync_kit_completion(sb, k) is True
        row = next(r for r in sb.tables["topic_kits"] if r["id"] == "kit-ar")
        assert (row["status"], row["translation_status"]) == ("in_review", "in_review")
        assert sb.tables["topics"][0]["status"] == "generating", "the topic is not a translation's to move"
        en = sb.tables["topic_kits"][0]
        assert "translation_status" not in en or en["translation_status"] is None

    def test_a_staged_article_is_accepted_only_for_a_translation(self):
        from catalogue import kit as K
        sb = _render_sb()
        gen = next(g for g in sb.tables["generations"] if g["id"] == "gen-par")
        prepared = K.prepare(sb, gen)
        assert prepared.translation["scripts"] == [TRANSLATED_SCRIPT_PATH] and prepared.language == "ar"
        # the same staged article behind a kit that is NOT a translation is refused
        plain = {**gen, "params": {k: v for k, v in gen["params"].items() if k != "translation"}}
        sb.tables["topic_kits"][1]["source_kit_id"] = None
        with pytest.raises(K.CatalogueRefused, match="not approved"):
            K.prepare(sb, plain)


# ── the render: a translated worksheet in worker.process ────────────────

STAGED_QUESTIONS_PATH = "sys/kit-ar/translated/worksheet_questions.json"


def _staged_arabic_questions() -> dict:
    kept, dropped = T.translate_questions(FakeTranslator(), WS_QUESTIONS["questions"], "ar", [], context="")
    assert len(kept) == 6 and dropped == []
    return {"title": "ورقة عمل — الدوائر", "instructions": "أجب عن كل سؤال.", "language": "ar", "questions": kept,
            "source": {"kit_id": "kit-1", "language": "en", "title": "Worksheet — Circles", "questions": 6, "dropped": []}}


def _render_ws_sb():
    sb = _render_sb()
    pres = next(g for g in sb.tables["generations"] if g["id"] == "gen-par")
    tr = {k: v for k, v in pres["params"]["translation"].items() if k != "scripts"}
    sb.tables["generations"].append({"id": "gen-war", "kind": "worksheet", "status": "processing", "owner_id": "sys",
                                     "book_id": None, "chapter_ref": None,
                                     "params": {**pres["params"], "translation": {**tr, "questions": STAGED_QUESTIONS_PATH}}})
    sb.tables["topic_kits"][1]["doc_generation_ids"]["worksheet"] = "gen-war"
    sb.files[("artifacts", STAGED_QUESTIONS_PATH)] = json.dumps(_staged_arabic_questions(), ensure_ascii=False).encode("utf-8")
    job = _job("job-war", "gen-war", "worksheet")
    job["params"]["kit_id"] = "kit-ar"
    sb.tables["jobs"] = [job]
    return sb


def _docx_text(raw: bytes) -> str:
    import io
    from docx import Document
    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


class TestTranslatedWorksheetRender:
    def test_the_sheet_and_its_key_are_drawn_from_the_staged_questions_with_no_model_call(self, monkeypatch):
        import io
        import zipfile
        sb = _render_ws_sb()
        process, uploads, analysis_calls = _worker_env(monkeypatch, sb)

        process.process_generation(sb, sb.tables["jobs"][0], "gen-war")

        assert analysis_calls == [], "a translated sheet is never analysed"
        dests = [d for _, d, _ in uploads]
        assert dests == ["sys/gen-war/worksheet.docx", "sys/gen-war/answer_key.docx", "sys/gen-war/worksheet_questions.json"]
        rows = [a for a in sb.tables["artifacts"] if a["generation_id"] == "gen-war"]
        assert [a["kind"] for a in rows] == ["docx", "answer_key_docx", "questions_json"]
        sheet = _docx_text(uploads[0][2])
        key = _docx_text(uploads[1][2])
        # the student sheet: every question in Arabic, the maths intact, the key not on it
        assert "ترجمة Solve: ٣x + ٥ = ٢٠" in sheet and "ترجمة The longest chord" in sheet and "ترجمة radius" in sheet
        assert "x = 5" not in sheet and "ترجمة diameter" not in sheet, "the answers are not on the student sheet"
        # the key: the working lines and the answers
        assert "ترجمة ٣x = ١٥" in key and "x = 5" in key and "C = 2 × 3.14 × 6" in key, "working with words translated, pure working verbatim"
        assert "ترجمة diameter" in key and "ترجمة Answer: ٣٧.٦٨ cm" in key and "40" in key
        # the figure question carries its picture
        from docx import Document
        assert len(Document(io.BytesIO(uploads[0][2])).inline_shapes) == 1
        # RTL: bidi paragraphs and <w:rtl/> on the runs
        xml = zipfile.ZipFile(io.BytesIO(uploads[0][2])).read("word/document.xml").decode("utf-8")
        assert "<w:rtl/>" in xml and "<w:bidi/>" in xml
        # the quiz player's file: the staged questions, answers included
        quiz = json.loads(uploads[2][2].decode("utf-8"))
        assert [q["id"] for q in quiz["questions"]] == ["q1", "q2", "q3", "q4", "q5", "q6"] and quiz["questions"][3]["answer"] == "x = 5"
        gen = next(g for g in sb.tables["generations"] if g["id"] == "gen-war")
        assert gen["status"] == "done" and gen["title"] == "ورقة عمل — الدوائر"
        assert ("select", "books") not in sb.calls

    def test_a_staged_file_with_no_question_fails_the_piece_before_any_upload(self, monkeypatch):
        sb = _render_ws_sb()
        sb.files[("artifacts", STAGED_QUESTIONS_PATH)] = json.dumps({"title": "x", "questions": []}).encode("utf-8")
        process, uploads, _ = _worker_env(monkeypatch, sb)
        with pytest.raises(Exception, match="no question"):
            process.process_generation(sb, sb.tables["jobs"][0], "gen-war")
        assert uploads == []
        # the lifecycle's drop-with-flag (PR #251): the piece leaves the kit, the reason stays for the reviewer
        k = next(k for k in sb.tables["topic_kits"] if k["id"] == "kit-ar")
        assert "worksheet" not in k["doc_generation_ids"] and "deck" in k["doc_generation_ids"]
        flag = next(f for f in k["precheck"]["flags"] if f["where"] == "worksheet")
        assert flag["what"] == T.DOCUMENT_FLAG and "no question" in flag["detail"]
        assert k["status"] == "generating", "one document is not the kit"


class TestOneTranslationAtATime:
    """T2 (decision D7): the observer lane claims the next kit_translate only
    when no translation is in flight — none staging, none rendering."""

    def _sb(self, *, queued=True, processing=False, rendering=None):
        sb = FakeSB()
        if queued:
            sb.tables["jobs"].append({"id": "jt-q", "type": "kit_translate", "status": "queued", "params": {}})
        if processing:
            sb.tables["jobs"].append({"id": "jt-p", "type": "kit_translate", "status": "processing", "params": {}})
        sb.tables["topic_kits"] += [{"id": "kit-en", "language": "en", "status": "generating"}]
        if rendering:
            sb.tables["topic_kits"].append({"id": "kit-ar", "language": rendering, "status": "generating"})
        return sb

    def test_the_two_reads(self):
        assert T.translations_queued(self._sb()) and not T.translations_queued(self._sb(queued=False))
        assert T.translation_in_flight(self._sb()) is None, "an English kit rendering is not a translation"
        assert "staging" in T.translation_in_flight(self._sb(processing=True))
        assert "rendering" in T.translation_in_flight(self._sb(rendering="ar"))

    def test_the_lane_drops_kit_translate_while_one_is_in_flight(self, monkeypatch):
        from worker import run as R
        assert "kit_translate" in R._catalogue_observer_types(self._sb()), "nothing in flight: the full list"
        assert "kit_translate" not in R._catalogue_observer_types(self._sb(processing=True))
        assert "kit_translate" not in R._catalogue_observer_types(self._sb(rendering="ar"))
        assert "topic_article" in R._catalogue_observer_types(self._sb(rendering="ar")), "the other observers still run"
        assert "kit_translate" in R._catalogue_observer_types(self._sb(queued=False, rendering="ar")), "no sweep queued: no gate"

        class Boom:
            def table(self, *a, **k):
                raise RuntimeError("down")
        assert "kit_translate" in R._catalogue_observer_types(Boom()), "a failed gate fails open"


class TestStepOperationIsTranslated:
    """The muted note beside each step of the working is `{operation}
    {explanation}` (maths/board.py). The Substitution Arabic video showed it
    in English — "Substitute y and a" — because `operation` was an invariant
    field. Now it is prose with maths inside, like the problem: translated,
    its maths spans and numbers checked."""

    def test_a_translated_operation_passes_and_a_lost_number_blocks(self):
        src = {"examples": [{"problem": "5y - 2a", "steps": [
            {"kind": "transform", "operation": "Substitute y and a", "before": ["5y - 2a"], "after": ["5(3) - 2(-4)"],
             "explanation": "", "speech": "We substitute."},
            {"kind": "transform", "operation": "subtract 3 from both sides", "before": ["x + 3 = 7"], "after": ["x = 4"],
             "explanation": "", "speech": "Subtract three."}]}]}
        good = {"examples": [{"problem": "5y - 2a", "steps": [
            {**src["examples"][0]["steps"][0], "operation": "عوّض y و a", "speech": "نعوّض."},
            {**src["examples"][0]["steps"][1], "operation": "اطرح 3 من الطرفين", "speech": "نطرح ثلاثة."}]}]}
        rep = C.check_maths_lesson(src, good)
        assert rep["ok"], rep["blocking"]
        bad = {"examples": [{"problem": "5y - 2a", "steps": [
            good["examples"][0]["steps"][0],
            {**good["examples"][0]["steps"][1], "operation": "اطرح من الطرفين"}]}]}
        rep = C.check_maths_lesson(src, bad)
        assert not rep["ok"] and rep["blocking"][0]["what"] == "maths in the question altered"
        assert rep["blocking"][0]["where"].endswith("steps[1].operation")


class TestAnOrphanedTranslationResumes:
    """2026-10-18: a translate job cut off after creating its kit row (a
    redeploy, then the reaper) must not be blocked by that row."""

    def test_its_own_kit_does_not_count_but_anyone_elses_does(self):
        sb = FakeSB()
        sb.tables["jobs"].append({"id": "jt", "type": "kit_translate", "status": "queued",
                                  "params": {"source_kit_id": "src-1", "language": "ar"}})
        sb.tables["topic_kits"].append({"id": "orphan", "language": "ar", "status": "generating", "source_kit_id": "src-1"})
        assert T.translation_in_flight(sb) is None, "the job's own resumable kit is not in flight"
        sb.tables["topic_kits"].append({"id": "other", "language": "ar", "status": "generating", "source_kit_id": "src-2"})
        assert "other" in T.translation_in_flight(sb), "another kit still rendering blocks the lane"

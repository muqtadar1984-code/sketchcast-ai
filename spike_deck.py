"""Build the spike deck from the LIVE Cells figures, offline, with no model call.

Five slides, every word of which is a real PowerPoint text object:

  1  title
  2  animal cell, fully annotated       <- the thing the founder asked for
  3  the nucleus, focus view            <- same PNG, cropped
  4  label-the-diagram                  <- same PNG, badges + answer blanks
  5  animal vs plant                    <- same two PNGs, difference labelled

Slides 3-5 spend ZERO extra image generations. Run:

    python spike_deck.py        # reads spike_out/fixture.json
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from agent5_slides import annotated_figure as af  # noqa: E402
from agent5_slides.theme import (FAINT, GRAPHITE, INK, LINE, MIST,  # noqa: E402
                                 TEAL_DK, WHITE)

IN = af.EMU_IN
OUT = ROOT / "spike_out"


def _rgb(t):
    from pptx.dml.color import RGBColor
    return RGBColor(*t)


def _load():
    fix = json.loads((OUT / "fixture.json").read_text(encoding="utf-8"))
    by_key = {}
    for f in fix:
        fig, asset = f["figure"], f["asset"]
        vision = asset.get("vision") or {}
        by_key[fig["figure_key"]] = {
            "png": OUT / f["png"],
            "caption": fig.get("caption") or "",
            "parts": (fig.get("spec") or {}).get("parts") or [],
            "notes": (fig.get("spec") or {}).get("notes") or "",
            "regions": vision.get("regions") or {},
            "w": float(vision.get("w") or 0),
            "h": float(vision.get("h") or 0),
        }
    return by_key


def _new(prs, bg=WHITE):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = _rgb(bg)
    return s


def _title_block(s, kicker, title, sub=""):
    from pptx.enum.text import PP_ALIGN
    af._textbox(s, 0.62 * IN, 0.42 * IN, 11.0 * IN, 0.26 * IN,
                kicker.upper(), 11, _rgb(TEAL_DK), bold=True)
    af._textbox(s, 0.62 * IN, 0.72 * IN, 11.0 * IN, 0.52 * IN,
                title, 27, _rgb(INK), bold=True)
    if sub:
        af._textbox(s, 0.62 * IN, 6.72 * IN, 11.0 * IN, 0.3 * IN,
                    sub, 11, _rgb(FAINT), align=PP_ALIGN.LEFT)



def build(dst: Path) -> tuple[Path, list[str]]:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN

    data = _load()
    animal, plant = data["animal_cell"], data["plant_cell"]
    faults: list[str] = []

    prs = Presentation()
    prs.slide_width, prs.slide_height = af.SLIDE_W, af.SLIDE_H

    # ── 1  title ────────────────────────────────────────────────────
    s = _new(prs, INK)
    af._textbox(s, 1.0 * IN, 2.4 * IN, 11.3 * IN, 0.9 * IN,
                "Cells", 44, _rgb(WHITE), bold=True)
    af._textbox(s, 1.0 * IN, 3.35 * IN, 11.3 * IN, 0.4 * IN,
                "Structures of the animal and plant cell", 18, _rgb(TEAL_DK))
    af._textbox(s, 1.0 * IN, 6.7 * IN, 11.3 * IN, 0.3 * IN,
                "SketchCast AI  ·  every label on every slide is editable text",
                11, _rgb(FAINT))

    # ── 2  the annotated diagram ────────────────────────────────────
    s = _new(prs)
    _title_block(s, "Structure", "The animal cell", animal["caption"])
    geo = af.figure_geometry(animal["parts"], animal["regions"],
                             animal["w"], animal["h"])
    missing = af.unresolved_parts(animal["parts"], animal["regions"])
    rep = af.add_annotated_figure(s, animal["png"], geo, animal["w"], animal["h"],
                                  frame=(3.45 * IN, 1.42 * IN, 6.45 * IN, 5.15 * IN))
    faults += [f"[animal] {f}" for f in af.validate(rep)]
    faults += [f"[animal] no region for {m!r}" for m in missing]
    s.notes_slide.notes_text_frame.text = animal["notes"]

    # ── 3  focus view — the SAME png, cropped ───────────────────────
    s = _new(prs)
    _title_block(s, "Zoom in", "Inside the nucleus",
                 "Same artwork as the previous slide — cropped, not redrawn.")
    nuc = next(g for g in geo if g["part"] == "nucleus")
    _pic, win = af.add_focus_view(s, animal["png"], nuc["union"],
                                  animal["w"], animal["h"],
                                  frame=(0.62 * IN, 1.45 * IN, 6.1 * IN, 5.1 * IN))
    tf = af._textbox(s, 7.1 * IN, 1.7 * IN, 5.6 * IN, 0.42 * IN,
                     "nucleus", 22, _rgb(INK), bold=True)
    body = af._textbox(s, 7.1 * IN, 2.3 * IN, 5.6 * IN, 3.4 * IN,
                       "The nucleus holds the cell's DNA and directs everything "
                       "the cell does. The darker sphere inside it is the "
                       "nucleolus, where ribosomes are assembled.",
                       15, _rgb(GRAPHITE))
    af._textbox(s, 7.1 * IN, 5.5 * IN, 5.6 * IN, 0.9 * IN,
                "Ask the class:  what would happen to a cell that lost its "
                "nucleus?", 14, _rgb(TEAL_DK), bold=True)
    s.notes_slide.notes_text_frame.text = (
        f"Crop window in the artwork's own pixels: "
        f"({win[0]:.0f}, {win[1]:.0f}) to ({win[2]:.0f}, {win[3]:.0f}) "
        f"of {animal['w']:.0f}x{animal['h']:.0f}. No new image was generated.")

    # ── 4  label-the-diagram (the unlabelled question view) ─────────
    s = _new(prs)
    _title_block(s, "Check", "Name each structure",
                 "Answers are in the speaker notes.")
    # Same figure, same ladder — only the label text changes. The numbers are
    # re-issued over the parts actually ASKED, so the worksheet reads 1..6
    # rather than inheriting gaps from the parts this view leaves out.
    quiz = [dict(g) for g in geo if not g["encloser"]][:6]
    for n, g in enumerate(quiz, 1):
        g["n"] = n
    rep4 = af.add_annotated_figure(s, animal["png"], quiz, animal["w"], animal["h"],
                                   frame=(1.6 * IN, 1.45 * IN, 5.2 * IN, 5.1 * IN),
                                   label_pt=14, label_text=af.numbers_only,
                                   gutter=0.42 * IN)
    faults += [f"[quiz] {f}" for f in af.validate(rep4, label_pt=14)]
    y = 1.75 * IN
    for n, _g in enumerate(quiz, 1):
        af._textbox(s, 7.9 * IN, y, 4.8 * IN, 0.34 * IN,
                    f"{n}.  ______________________", 16, _rgb(INK))
        y += 0.62 * IN
    s.notes_slide.notes_text_frame.text = "Answers:  " + "   ".join(
        f"{n}. {g['part']}" for n, g in enumerate(quiz, 1))

    # ── 5  comparison — difference computed, not authored ───────────
    s = _new(prs)
    _title_block(s, "Compare", "Animal cell vs plant cell")
    only_plant = [p for p in plant["parts"]
                  if p.lower() not in {q.lower() for q in animal["parts"]}]
    half = (0.62 * IN, 1.7 * IN, 5.7 * IN, 4.1 * IN)
    ax, ay, aw, ah = af._fit(animal["w"], animal["h"], half)
    s.shapes.add_picture(str(animal["png"]), int(ax), int(ay), int(aw), int(ah))
    af._textbox(s, 0.62 * IN, 1.32 * IN, 5.7 * IN, 0.3 * IN,
                "ANIMAL", 12, _rgb(GRAPHITE), bold=True, align=PP_ALIGN.CENTER)
    half2 = (6.9 * IN, 1.7 * IN, 5.7 * IN, 4.1 * IN)
    bx, by, bw, bh = af._fit(plant["w"], plant["h"], half2)
    s.shapes.add_picture(str(plant["png"]), int(bx), int(by), int(bw), int(bh))
    af._textbox(s, 6.9 * IN, 1.32 * IN, 5.7 * IN, 0.3 * IN,
                "PLANT", 12, _rgb(GRAPHITE), bold=True, align=PP_ALIGN.CENTER)
    af._textbox(s, 0.62 * IN, 6.02 * IN, 12.1 * IN, 0.34 * IN,
                "Only the plant cell has:  " + ",  ".join(only_plant),
                16, _rgb(TEAL_DK), bold=True)
    s.notes_slide.notes_text_frame.text = (
        "The contrast is a set difference over the two figures' declared "
        "parts — nothing about cells is written into the code.")

    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    return dst, faults


if __name__ == "__main__":
    path, faults = build(OUT / "SketchCast_Cells_annotated_spike.pptx")
    print(f"saved {path.name}  ({path.stat().st_size / 1024:.0f} KB)")
    if faults:
        print("GEOMETRY FAULTS:")
        for f in faults:
            print("  -", f)
    else:
        print("validator: clean")

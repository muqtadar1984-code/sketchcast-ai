"""Spike: an annotated diagram slide whose every word is a real PowerPoint object.

WHY THIS EXISTS. ``_build_designed_deck`` embeds one 1280x720 PNG per slide,
full-bleed. That PNG is a *rendering of bullet points*, so the deck the teacher
downloads is (a) bullets and (b) uneditable — a heading they cannot retype sits
inside a picture. Both complaints are the same defect: the text was rasterised
before it reached the file.

Nothing new has to be generated to fix it. ``article_figures.spec.parts``
already names what the picture contains, ``visual_assets.vision.regions``
already records WHERE each of those parts is in the artwork's own pixels, and
``vision.w``/``h`` give the frame those pixels were measured in. A region box
plus that frame is a pure linear map onto the slide, so the label can be a text
box, the leader a connector, and the artwork the only raster on the slide.

WHAT IT DELIBERATELY IS NOT. There is no archetype framework here and no model
call. One function draws one figure. The teaching views below are the same
image cropped different ways — ``Picture.crop_*`` is a settable property in
python-pptx 1.0.2 — so four views cost exactly one image generation. That is
the point: image capacity is ~1/min and shared, so a view that costs artwork is
a view we cannot afford to offer.
"""

from __future__ import annotations

import logging
import math
from pathlib import Path

logger = logging.getLogger(__name__)

EMU_IN = 914400
SLIDE_W = 12192000            # 13.333in — 16:9, same canvas as the designed deck
SLIDE_H = 6858000             # 7.5in

# A region covering most of the frame is not a blob you can point at, it is the
# thing everything else sits inside. `cell membrane` is [12,14,690,598] of a
# 704x613 picture: its centroid is the middle of the cell, so a leader drawn to
# that centroid lands on the NUCLEUS and teaches the wrong structure. Measured
# on the live Cells rows every encloser (membrane, cytoplasm, cell wall) sits
# above 0.55 area-fraction and every pointable organelle far below it.
_ENCLOSER_AREA = 0.55

# Below this the part is a scatter of dots (ribosome: 6 boxes, ~0.0002 each) and
# a textbook labels ONE of them, not all six.
_SPECK_AREA = 0.004


# ─────────────────────────── geometry ────────────────────────────

# Name matching lives in `shared.lesson_model`, which the storyboard uses
# too. Two copies of this rule is what lost `climate_impact` its slide.
from shared.lesson_model import display_part, part_boxes as _boxes  # noqa: E402


def _area_frac(box, w: float, h: float) -> float:
    x0, y0, x1, y1 = box
    return abs((x1 - x0) * (y1 - y0)) / max(1.0, w * h)


def _union(boxes):
    return (min(b[0] for b in boxes), min(b[1] for b in boxes),
            max(b[2] for b in boxes), max(b[3] for b in boxes))


def _pick_instance(boxes, side: str):
    """Which of N stored instances this label points at.

    A scatter part (ribosome) gets the instance nearest the gutter the label
    lives in, so the leader is short and crosses the least artwork. A part with
    a few real instances (5 mitochondria) is treated the same way — pointing at
    the nearest one is what a textbook does, and it keeps the leader out of the
    middle of the drawing. Ties break toward the bigger instance: a reader
    hunting for "a mitochondrion" should be sent to one they can actually see.
    """
    if len(boxes) == 1:
        return boxes[0]
    sign = 1 if side == "left" else -1
    return sorted(boxes, key=lambda b: (sign * (b[0] + b[2]) / 2,
                                        -(b[2] - b[0]) * (b[3] - b[1])))[0]


def _distribute(desired: list[float], heights: list[float],
                top: float, bottom: float, gap: float) -> list[float]:
    """Place label boxes at their wanted heights without overlapping.

    Order in == order out, and that IS the non-crossing guarantee: the labels
    on one side are sorted by their anchor's y before they get here, so a
    solver that never reorders them cannot produce two leaders that cross.
    """
    n = len(desired)
    if not n:
        return []
    y = list(desired)
    y[0] = max(y[0], top)
    for i in range(1, n):                       # push down
        y[i] = max(y[i], y[i - 1] + heights[i - 1] + gap)
    overflow = (y[-1] + heights[-1]) - bottom
    if overflow > 0:                            # ...then relax back up
        for i in range(n):
            y[i] -= overflow
        for i in range(n - 2, -1, -1):
            y[i] = min(y[i], y[i + 1] - heights[i] - gap)
        y[0] = max(y[0], top)
        for i in range(1, n):
            y[i] = max(y[i], y[i - 1] + heights[i - 1] + gap)
    return y


def _anchor_ladder(spans: list[tuple[float, float]],
                   wanted: list[float]) -> list[float]:
    """Where each leader touches its part — provably without crossing.

    Each anchor must land inside its own part's vertical span, and the anchors
    must run down the page in the same order as their labels. A greedy running
    floor is not enough, because parts CONTAIN one another: the nucleolus sits
    inside the nucleus, so their spans overlap while font pitch forces their
    labels apart, and greedy pushed the nucleus's anchor below the nucleolus's
    box altogether. That crossing was ~0.0007in — invisible, and still a lie
    about which structure the line points to.

    So bound first, place second. ``lo`` is the lowest each anchor may sit
    given everything above it, ``hi`` the highest given everything below. Both
    sequences come out non-decreasing by construction, ``wanted`` (the label
    heights) already is, and clamping a non-decreasing target into
    non-decreasing bounds is non-decreasing — min and max preserve the order.
    Crossing therefore cannot be produced, only reported: if ``lo > hi`` the
    spans genuinely cannot be ordered this way, and the caller's validator says
    so rather than this function guessing.
    """
    n = len(spans)
    if not n:
        return []
    lo = [0.0] * n
    hi = [0.0] * n
    run = -1e18
    for i, (y0, _y1) in enumerate(spans):
        run = max(y0, run)
        lo[i] = run
    run = 1e18
    for i in range(n - 1, -1, -1):
        run = min(spans[i][1], run)
        hi[i] = run
    return [min(max(wanted[i], lo[i]), max(hi[i], lo[i])) for i in range(n)]


def _text_height(label: str, width_emu: float, pt: float) -> float:
    """Enough room for the wrapped label, through the shared script-aware
    metric — a second Latin-only estimate here is how two Telugu labels end
    up stacked while the validator, using the same wrong number, agrees."""
    from .metrics import text_height_in
    return text_height_in(label, width_emu / EMU_IN, pt) * EMU_IN + 0.06 * EMU_IN


def _assign_sides(geo: list[dict], w: float) -> None:
    """Left or right gutter, balanced.

    Centroid decides, then the parts nearest the midline move across until the
    two columns differ by at most one. An unbalanced split is not merely ugly:
    nine labels in one gutter forces a font small enough to be unreadable
    projected, while the other half of the slide sits empty.
    """
    mid = w / 2
    for g in geo:
        g["side"] = "left" if g["cx"] < mid else "right"
        g["pull"] = abs(g["cx"] - mid)
    while True:
        left = [g for g in geo if g["side"] == "left"]
        right = [g for g in geo if g["side"] == "right"]
        if abs(len(left) - len(right)) <= 1:
            return
        heavy = left if len(left) > len(right) else right
        mover = min(heavy, key=lambda g: g["pull"])
        mover["side"] = "right" if mover["side"] == "left" else "left"


def figure_geometry(parts: list[str], regions: dict, w: float, h: float) -> list[dict]:
    """Resolve declared parts against stored regions. Subject-blind by
    construction: nothing here reads a part NAME, only its box."""
    geo = []
    for part in parts:
        boxes = _boxes(regions, part)
        if not boxes:
            continue
        u = _union(boxes)
        # ENCLOSURE IS A PROPERTY OF ONE INSTANCE, NOT OF THE SET. The live
        # animal cell has 5 mitochondria scattered from (57,·) to (·,544); their
        # UNION covers 60% of the frame, so a union-based test called
        # `mitochondrion` an encloser and pointed the leader at the middle of
        # the cell — at the nucleus. Any part with instances near opposite
        # corners would have failed the same way, in any subject.
        biggest = max(boxes, key=lambda b: (b[2] - b[0]) * (b[3] - b[1]))
        geo.append({
            "part": part, "boxes": boxes, "union": u,
            "cx": (u[0] + u[2]) / 2, "cy": (u[1] + u[3]) / 2,
            "encloser": _area_frac(biggest, w, h) >= _ENCLOSER_AREA,
            "speck": _area_frac(biggest, w, h) <= _SPECK_AREA,
        })
    _assign_sides(geo, w)
    for i, g in enumerate(geo, 1):
        g["n"] = i
        g["target"] = g["union"] if g["encloser"] else _pick_instance(g["boxes"], g["side"])
    return geo


def unresolved_parts(parts: list[str], regions: dict) -> list[str]:
    """Declared parts the artwork cannot locate. A caller that refuses to build
    a labelled slide unless this is empty gets the fail-over-degrade rule for
    free — an unlabelled organelle is a diagram that teaches a hole."""
    return [p for p in parts if not _boxes(regions, p)]


# ─────────────────────────── rendering ────────────────────────────

def _target_cy(g: dict) -> float:
    """Mid-height of the box the leader lands on."""
    return (g["target"][1] + g["target"][3]) / 2


def _ink_mask(png: str | Path):
    """The artwork as a row-addressable boolean of "there is a stroke here".

    Composited onto white first: these PNGs are RGBA and a bare ``convert("L")``
    turns transparent pixels black, which would make the very first column
    "ink" and every scan below return the frame edge.
    """
    from PIL import Image

    im = Image.open(str(png)).convert("RGBA")
    flat = Image.new("RGB", im.size, (255, 255, 255))
    flat.paste(im, mask=im.split()[3])
    g = flat.convert("L")
    return g.load(), g.width, g.height


def _ellipse_x(box, row: float, side: str) -> float:
    """Where the ellipse inscribed in `box` crosses `row`, on `side`.

    Almost every organelle in this library is drawn as a blob inside its box,
    so the box's near EDGE is only the right answer at the box's widest row.
    At three-quarter height the nucleus's outline has moved ~100px inward and
    the leader ends in blank cytoplasm. Geometry is used rather than the pixels
    here on purpose: an inner structure overlaps its neighbours' boxes, so a
    pixel scan would happily snap the nucleus's leader onto the ER.
    """
    x0, y0, x1, y1 = box
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    a, b = (x1 - x0) / 2, (y1 - y0) / 2
    if b <= 0:
        return x0 if side == "left" else x1
    t = min(1.0, abs(row - cy) / b)
    dx = a * math.sqrt(max(0.0, 1.0 - t * t))
    return cx - dx if side == "left" else cx + dx


def _outer_ink_x(px, iw: int, ih: int, box, row: float, side: str):
    """First stroke met travelling INWARD from outside the picture, on `row`.

    Only ever asked of an encloser, and for an encloser it is exact: the part
    that contains everything else is the one thing reachable from outside
    without crossing something else first. That is what makes it right for an
    irregular cell outline, where no ellipse fits and the bounding box's edge
    is the widest point rather than the point at this height.
    """
    r = int(round(row))
    if not (0 <= r < ih):
        return None
    x0, x1 = int(max(0, box[0])), int(min(iw - 1, box[2]))
    xs = range(x0, x1 + 1) if side == "left" else range(x1, x0 - 1, -1)
    for x in xs:
        for dy in (0, -2, 2, -4, 4):
            ry = r + dy
            if 0 <= ry < ih and px[x, ry] < 128:
                return float(x)
    return None


def _theme():
    from pptx.dml.color import RGBColor

    from .theme import FAINT, GRAPHITE, INK, TEAL_DK, WHITE
    return (RGBColor(*INK), RGBColor(*WHITE), RGBColor(*TEAL_DK),
            RGBColor(*GRAPHITE), RGBColor(*FAINT))


def _textbox(slide, l, t, w, h, text, pt, color, bold=False, align=None, anchor=None):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    tf = slide.shapes.add_textbox(int(l), int(t), int(w), int(h)).text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor or MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align or PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.name = "Calibri"
    r.font.color.rgb = color
    return tf


def _leader(slide, x1, y1, x2, y2, color):
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Pt

    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(0.75)
    d = 0.055 * EMU_IN                       # the dot that says "this exact spot"
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, int(x2 - d / 2), int(y2 - d / 2),
                               int(d), int(d))
    o.fill.solid()
    o.fill.fore_color.rgb = color
    o.line.fill.background()
    o.shadow.inherit = False
    return c


def _fit(img_w: float, img_h: float, box) -> tuple[float, float, float, float]:
    """Largest same-aspect rectangle inside `box` (l, t, w, h), centred."""
    l, t, bw, bh = box
    s = min(bw / img_w, bh / img_h)
    w, h = img_w * s, img_h * s
    return l + (bw - w) / 2, t + (bh - h) / 2, w, h


def numbers_only(g: dict) -> str:
    """Label text for the question view: the number, not the answer.

    Badges stamped straight onto each region's centroid look obvious and are
    not: the nucleolus sits AT the nucleus's centroid, so its badge covered the
    nucleus's completely and the first question silently vanished from the
    worksheet. Numbering through the gutter reuses the ladder that already
    guarantees no overlap and no crossing, and costs nothing to maintain.
    """
    return f"{g['n']}."


def add_annotated_figure(slide, png: str | Path, geo: list[dict], img_w: float,
                         img_h: float, frame, label_pt: float = 12.0,
                         label_text=None, gutter: float | None = None) -> dict:
    """Draw the artwork inside `frame` and label it with native text + leaders.

    `frame` is (l, t, w, h) in EMU — the area the picture may use, gutters
    excluded. Returns a geometry report the validator reads; nothing about it
    is figure- or subject-specific.
    """
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    ink, _white, _teal, grey, _faint = _theme()
    fl, ft, fw, fh = frame
    px, py, pw, ph = _fit(img_w, img_h, (fl, ft, fw, fh))
    slide.shapes.add_picture(str(png), int(px), int(py), int(pw), int(ph))
    sx, sy = pw / img_w, ph / img_h
    try:
        pix, iw, ih = _ink_mask(png)
    except Exception as exc:                # noqa: BLE001 — geometry still works
        logger.warning("no ink mask for %s (%s); enclosers fall back to geometry",
                       png, exc)
        pix = iw = ih = None

    gut_w = gutter or max(1.6 * EMU_IN, min(2.55 * EMU_IN, px - 0.9 * EMU_IN))
    placed = []
    for side in ("left", "right"):
        # Order by the box the leader actually POINTS AT, not by the union of
        # every instance. `ribosome`'s union spans most of the cell while the
        # instance chosen for the left gutter sits near the top; ordering on the
        # union put that label below neighbours it had to reach above, and the
        # validator called the crossing before this comment existed.
        col = sorted([g for g in geo if g["side"] == side], key=_target_cy)
        if not col:
            continue
        # Parts are matched as declared (`thirty_year_calendar`) and SHOWN as
        # words: the identifier went up verbatim on the Weather deck, and a
        # teacher should not have to read snake_case off a projected slide.
        texts = [(label_text or (lambda g: display_part(g["part"])))(g) for g in col]
        heights = [_text_height(t, gut_w, label_pt) for t in texts]
        desired = [py + _target_cy(g) * sy - hh / 2 for g, hh in zip(col, heights)]
        ys = _distribute(desired, heights, ft, ft + fh, 0.07 * EMU_IN)
        gx = (px - 0.30 * EMU_IN - gut_w) if side == "left" else (px + pw + 0.30 * EMU_IN)
        spans = [(py + g["target"][1] * sy, py + g["target"][3] * sy) for g in col]
        anchors = _anchor_ladder(spans, [ly + hh / 2 for hh, ly in zip(heights, ys)])
        for i, (g, hh, ly, ay, text) in enumerate(zip(col, heights, ys, anchors, texts)):
            _textbox(slide, gx, ly, gut_w, hh, text, label_pt, ink,
                     align=PP_ALIGN.RIGHT if side == "left" else PP_ALIGN.LEFT,
                     anchor=MSO_ANCHOR.MIDDLE)
            ly_mid = ly + hh / 2
            row = (ay - py) / sy                    # back into artwork pixels
            ax_px = None
            if g["encloser"] and pix is not None:
                ax_px = _outer_ink_x(pix, iw, ih, g["target"], row, side)
            if ax_px is None:
                ax_px = _ellipse_x(g["target"], row, side)
            ax = px + ax_px * sx
            lx = gx + gut_w + 0.05 * EMU_IN if side == "left" else gx - 0.05 * EMU_IN
            _leader(slide, lx, ly_mid, ax, ay, grey)
            placed.append({"part": g["part"], "side": side, "order": i,
                           "label": (gx, ly, gut_w, hh), "anchor": (ax, ay)})
    return {"picture": (px, py, pw, ph), "labels": placed, "gutter": gut_w}


def add_focus_view(slide, png, region, img_w, img_h, frame, pad: float = 0.18):
    """The SAME image, cropped to one region — no second generation.

    ``Picture.crop_left/right/top/bottom`` are settable in python-pptx 1.0.2 and
    are fractions of the ORIGINAL image, so the crop window is computed in the
    artwork's own pixel space — the space the regions were measured in. The
    window is grown to the frame's aspect ratio FIRST, or PowerPoint stretches
    the visible pixels to fill the shape and the organelle comes out oval.
    """
    fl, ft, fw, fh = frame
    x0, y0, x1, y1 = region
    mx, my = (x1 - x0) * pad, (y1 - y0) * pad
    x0, y0, x1, y1 = x0 - mx, y0 - my, x1 + mx, y1 + my
    want = fw / fh
    cw, ch = x1 - x0, y1 - y0
    if cw / ch < want:                       # too tall for the frame — widen
        need = ch * want
        cx = (x0 + x1) / 2
        x0, x1 = cx - need / 2, cx + need / 2
    else:                                    # too wide — heighten
        need = cw / want
        cy = (y0 + y1) / 2
        y0, y1 = cy - need / 2, cy + need / 2
    # Never crop past the edge of the artwork; slide the window back inside.
    if x0 < 0:
        x1, x0 = x1 - x0, 0.0
    if y0 < 0:
        y1, y0 = y1 - y0, 0.0
    if x1 > img_w:
        x0, x1 = x0 - (x1 - img_w), float(img_w)
    if y1 > img_h:
        y0, y1 = y0 - (y1 - img_h), float(img_h)
    x0, y0 = max(0.0, x0), max(0.0, y0)
    x1, y1 = min(float(img_w), x1), min(float(img_h), y1)

    pic = slide.shapes.add_picture(str(png), int(fl), int(ft), int(fw), int(fh))
    pic.crop_left = x0 / img_w
    pic.crop_right = 1.0 - x1 / img_w
    pic.crop_top = y0 / img_h
    pic.crop_bottom = 1.0 - y1 / img_h
    return pic, (x0, y0, x1, y1)


# ─────────────────────────── validator ────────────────────────────

def validate(report: dict, min_pt: float = 11.0, label_pt: float = 12.0) -> list[str]:
    """Geometry faults a reader would see. Empty list == the slide is legible.

    This is what makes the approach safe to ship. A labelled diagram fails
    VISIBLY — a label over the artwork, two labels stacked, a crossed leader —
    and unlike a video nobody watches a deck fail; they open it in front of a
    class. So the check runs on the geometry we just wrote, before the save.
    """
    faults = []
    if label_pt < min_pt:
        faults.append(f"label font {label_pt}pt is below the {min_pt}pt floor")
    labels = report.get("labels") or []
    for i, a in enumerate(labels):
        for b in labels[i + 1:]:
            ax, ay, aw, ah = a["label"]
            bx, by, bw, bh = b["label"]
            if ax < bx + bw and bx < ax + aw and ay < by + bh and by < ay + ah:
                faults.append(f"labels overlap: {a['part']} / {b['part']}")
    px, py, pw, ph = report.get("picture", (0, 0, 0, 0))
    for a in labels:
        ax, ay, aw, ah = a["label"]
        if ax < px + pw and px < ax + aw and ay < py + ph and py < ay + ah:
            faults.append(f"label sits on the artwork: {a['part']}")
    for side in ("left", "right"):
        col = sorted([a for a in labels if a["side"] == side],
                     key=lambda a: a["order"])
        ys = [a["anchor"][1] for a in col]
        if any(ys[i] > ys[i + 1] for i in range(len(ys) - 1)):
            faults.append(f"{side} leaders cross")
    return faults
